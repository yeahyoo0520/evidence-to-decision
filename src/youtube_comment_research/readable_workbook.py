from __future__ import annotations

import math
import re
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .evidence import ResearchContext
from .decision_support import DecisionBriefRecord
from .exporters import EXPORT_FIELDS
from .human_review import HumanReviewBundle
from .models import CommentRecord
from .insight_synthesis import InsightRecord, insight_section, mark_insight_evidence
from .profiles import DEFAULT_PROFILE_ID, ResearchProfile, load_profile
from .review_sampling import apply_review_sampling, summarize_review_workload
from .reporting import (
    CATEGORY_LABELS,
    DATA_LIMITATIONS,
    build_video_comparison,
)
from .semantic_workflow import build_review_sections, prepare_report_rows
from .theme_discovery import ThemeInsightBundle
from .verification import resolve_effective_insights

SHEET_NAMES = [
    "01_项目总览",
    "02_对话视图",
    "03_语义标注",
    "04_主题洞察",
    "05_视频对比",
    "06_原始数据",
    "07_字段说明",
    "08_主题发现",
    "09_洞察草稿",
    "10_人工审核",
    "11_决策支持",
]

GENERIC_SHEET_NAMES = [
    *SHEET_NAMES[:3],
    "04_分类概览",
    "05_来源对比",
    *SHEET_NAMES[5:9],
    "10_可选核验",
    SHEET_NAMES[10],
]

GENERIC_RAW_FIELDS = (
    "evidence_id",
    "text",
    "source_type",
    "source_name",
    "created_at",
    "language",
    "market",
    "region",
    "metadata",
)

HEADER_FILL = PatternFill("solid", fgColor="FF1F4E78")
HEADER_FONT = Font(color="FFFFFFFF", bold=True)
MAIN_COMMENT_FILL = PatternFill("solid", fgColor="FFFFF2CC")
REPLY_FILL = PatternFill("solid", fgColor="FFDDEBF7")
SECTION_FILL = PatternFill("solid", fgColor="FFD9EAF7")
THIN_BORDER = Border(
    left=Side(style="thin", color="FFD9E2F3"),
    right=Side(style="thin", color="FFD9E2F3"),
    top=Side(style="thin", color="FFD9E2F3"),
    bottom=Side(style="thin", color="FFD9E2F3"),
)

CONVERSATION_COLUMNS: tuple[tuple[str, str], ...] = (
    ("thread", "对话组 Thread"),
    ("type", "类型 Type"),
    ("comment_text", "评论正文 Comment"),
    ("reply_to", "回复对象（主评论） Reply to"),
    ("comment_like_count", "点赞 Likes"),
    ("reply_count_display", "回复数 Replies"),
    ("author_display_name", "作者 Author"),
    ("comment_published_at", "发布时间 Published at"),
    ("video_title", "视频 Video"),
    ("channel_title", "频道 Channel"),
    ("comment_id", "评论 ID comment_id"),
    ("parent_comment_id", "父评论 ID parent_comment_id"),
)

CLASSIFICATION_COLUMNS: tuple[tuple[str, str], ...] = (
    ("comment_id", "评论 ID comment_id"),
    ("comment_text", "评论正文 Comment"),
    ("codex_subject", "评论主体 Subject"),
    ("codex_primary_category", "Codex 主分类 Primary category"),
    ("codex_secondary_categories", "Codex 次分类 Secondary categories"),
    ("codex_subtopics", "Codex 子话题 Subtopics"),
    ("codex_lifecycle_stage", "阶段 Stage"),
    ("codex_confidence", "Codex 置信度 Confidence"),
    ("codex_review_risk", "Review Risk"),
    ("codex_manual_review", "Mandatory Manual Review"),
    ("codex_review_triggers", "Review Triggers"),
    ("qa_sample", "QA Sample"),
    ("qa_status", "QA Status"),
    ("evidence_review_required", "Evidence Review Required"),
    ("evidence_review_status", "Evidence Review Status"),
    ("review_status", "Review Status"),
    ("reviewer_note", "Reviewer Note"),
    ("codex_review_reason", "复核原因 Review reason"),
    ("codex_evidence", "判断依据 Evidence"),
    ("codex_mentioned_brands", "提及品牌 Mentioned brands"),
    ("codex_mentioned_products", "提及产品 Mentioned products"),
    ("human_confirmed_primary_category", "人工确认主分类 Human primary"),
    ("human_confirmed_secondary_categories", "人工确认次分类 Human secondary"),
    ("human_confirmed_subtopics", "人工确认子话题 Human subtopics"),
    ("human_confirmed_lifecycle_stage", "人工确认生命周期 Human lifecycle"),
    ("reviewed_at", "复核时间 Reviewed at"),
    ("classification_source", "分类来源 Classification source"),
    ("classifier_version", "分类器版本 Classifier version"),
    ("rubric_version", "规则手册版本 Rubric version"),
    ("classified_at", "分类时间 Classified at"),
    ("classification_history", "历史分类 Classification history"),
    ("research_profile", "Research Profile"),
    ("qa_sample_reason", "QA Sample Reason"),
    ("qa_sample_seed", "QA Sample Seed"),
    ("qa_sampled_at", "QA Sampled At"),
    ("video_id", "视频 ID video_id"),
    ("is_reply", "是否回复 Is reply"),
    ("comment_like_count", "点赞 Likes"),
    ("rule_classification_label", "规则定位 Rule status"),
    ("rule_candidate_categories", "规则候选分类 Rule-based candidate"),
    ("rule_candidate_subtopics", "规则候选子话题 Rule subtopics"),
    ("detected_entities", "规则检测实体 Detected entities"),
    ("detected_comparison_phrases", "规则比较短语 Comparison phrases"),
    ("detected_lifecycle_signals", "规则生命周期信号 Lifecycle signals"),
    ("rule_evidence", "规则依据 Rule evidence"),
    # Phase 2A.1 compatibility columns remain available but are visually secondary.
    ("comment_id", "comment_id"),
    ("comment_text", "comment_text"),
    ("primary_category", "primary_category"),
    ("primary_category_cn", "primary_category_cn"),
    ("category_cn", "category_cn"),
    ("secondary_categories", "secondary_categories"),
    ("secondary_categories_cn", "secondary_categories_cn"),
    ("subtopics", "subtopics"),
    ("subtopic", "subtopic"),
    ("lifecycle_stage", "lifecycle_stage"),
    ("mentioned_brands", "mentioned_brands"),
    ("mentioned_products", "mentioned_products"),
    ("category_confidence", "category_confidence"),
    ("category_evidence", "category_evidence"),
    ("manual_review", "manual_review"),
    ("review_reason", "review_reason"),
    ("matched_rules", "matched_rules"),
)

GENERIC_CLASSIFICATION_COLUMNS: tuple[tuple[str, str], ...] = (
    ("comment_id", "Evidence ID"),
    ("comment_text", "Evidence Text"),
    ("codex_subject", "Subject"),
    ("codex_primary_category", "Primary Category"),
    ("codex_secondary_categories", "Secondary Categories"),
    ("codex_subtopics", "Subtopics"),
    ("codex_lifecycle_stage", "Stage"),
    ("codex_confidence", "Confidence"),
    ("codex_review_risk", "Review Risk"),
    ("codex_manual_review", "Review Recommended"),
    ("codex_review_triggers", "Review Triggers"),
    ("codex_review_reason", "Review Reasons"),
    ("codex_evidence", "Evidence"),
    ("classification_source", "Classification Source"),
    ("classifier_version", "Classifier Version"),
    ("rubric_version", "Rubric Version"),
    ("classified_at", "Classified At"),
    ("research_profile", "Research Profile"),
)

TOPIC_FIELDS: tuple[tuple[str, str], ...] = (
    ("report_section", "报告分区 Report section"),
    ("theme", "Theme"),
    ("evidence_count", "Evidence Count"),
    ("evidence_videos", "Evidence Videos"),
    ("representative_comment", "Representative Comment"),
    ("representative_comment_id", "Representative Comment ID"),
    ("insight", "Insight"),
    ("action_opportunity", "Action Opportunity"),
    ("priority", "Priority"),
    ("priority_reason", "Priority Reason"),
    ("status", "Status"),
    ("evidence_review_status", "Evidence Review Status"),
    ("needs_human_confirmation", "Needs Human Confirmation"),
)

RAW_FIELD_DESCRIPTIONS = {
    "project_id": "项目或批次标识。",
    "video_id": "YouTube 视频 ID。",
    "video_url": "公开视频链接。",
    "video_title": "采集时的视频标题。",
    "channel_title": "采集时的频道名称。",
    "video_published_at": "视频发布时间。",
    "video_view_count": "采集时的视频观看次数。",
    "video_like_count": "采集时的视频点赞次数。",
    "video_comment_count": "采集时 API 返回的视频评论总数。",
    "comment_id": "YouTube 评论唯一 ID，用于结果追溯。",
    "parent_comment_id": "回复所对应的顶级评论 ID；顶级评论为空。",
    "is_reply": "是否为回复。",
    "comment_text": "API 返回的公开评论正文。",
    "comment_like_count": "采集时的评论点赞数。",
    "reply_count": "顶级评论的 API 回复数；回复行为 0。",
    "comment_published_at": "评论发布时间。",
    "comment_updated_at": "评论最后更新时间。",
    "author_display_name": "作者显示名；启用匿名化后为稳定匿名标签。",
    "author_channel_id": "作者频道 ID；启用匿名化后为空。",
    "collected_at": "本批次采集时间。",
    "collection_order": "采集顺序。",
}

REPORT_FIELD_DESCRIPTIONS = {
    "primary_category": "规则分类得到的英文主分类。",
    "primary_category_cn": "主分类中文名称。",
    "category_cn": "兼容旧版输出的主分类中文名称别名。",
    "secondary_categories": "保留评论中次要但有证据支持的英文分类。",
    "secondary_categories_cn": "次分类中文名称。",
    "subtopics": "独立于主分类识别的一个或多个可追溯子话题。",
    "subtopic": "兼容旧版输出的首个子话题别名。",
    "lifecycle_stage": "购买生命周期阶段；不会直接覆盖主分类。",
    "mentioned_brands": "评论中明确提及或由产品型号指向的品牌。",
    "mentioned_products": "评论中识别出的受保护产品型号。",
    "category_confidence": "综合匹配总分、分差、强证据、冲突和主体歧义校准的置信度。",
    "category_evidence": "主分类与次分类的匹配短语、规则、得分和目标实体。",
    "matched_rules": "完整可追溯规则匹配记录，JSON 格式。",
    "manual_review": "是否需要人工复核。",
    "review_reason": "触发人工复核的具体原因。",
    "Thread": "对话组编号，用于把顶级评论与其回复组织在一起。",
    "reply_to": "回复所对应顶级评论的文本摘要。",
}

REPORT_FIELD_DESCRIPTIONS.update(
    {
        "rule_classification_label": "固定标识为 Rule-based candidate / 规则候选。",
        "rule_candidate_categories": "仅供参考的规则候选分类，不代表确认结论。",
        "rule_candidate_subtopics": "规则匹配出的候选子话题。",
        "detected_entities": "规则层检测到的品牌与产品实体。",
        "detected_comparison_phrases": "规则层检测到的明确比较关系短语。",
        "detected_lifecycle_signals": "规则层检测到的购买生命周期信号。",
        "rule_evidence": "规则候选的可追溯匹配依据。",
        "codex_subject": "Codex 先识别的评论主体。",
        "codex_primary_category": "Codex 辅助语义判断的主分类；在人工确认前属于初步结果。",
        "codex_secondary_categories": "Codex 辅助语义判断保留的真实次要意图。",
        "codex_subtopics": "Codex 辅助语义判断的子话题。",
        "codex_lifecycle_stage": "Codex 辅助语义判断的购买生命周期。",
        "codex_mentioned_brands": "Codex 从评论中明确识别出的品牌。",
        "codex_mentioned_products": "Codex 从评论中明确识别出的产品。",
        "codex_confidence": "Codex 对主分类判断把握的内部标尺；不是统计概率，也不等于无需人工复核。",
        "codex_evidence": "直接来自原评论的短证据片段。",
        "codex_review_risk": "独立于置信度的复核风险：low、medium 或 high。",
        "codex_review_triggers": "触发复核风险的简短、可解释原因数组。",
        "codex_manual_review": "Codex 是否建议进入人工复核队列。",
        "codex_review_reason": "建议人工复核的具体原因。",
        "human_confirmed_primary_category": "人工确认或纠正后的主分类。",
        "human_confirmed_secondary_categories": "人工确认或纠正后的次分类。",
        "human_confirmed_subtopics": "人工确认或纠正后的子话题。",
        "human_confirmed_lifecycle_stage": "人工确认或纠正后的生命周期。",
        "review_status": "人工复核状态：pending、confirmed、corrected 或 rejected。",
        "reviewed_at": "人工复核时间。",
        "reviewer_note": "复核者记录；自动流程不得覆盖。",
        "classification_source": "当前 Codex 结果来源，固定为 codex-assisted。",
        "classifier_version": "Codex 分类器或提示版本。",
        "rubric_version": "分类标准版本。",
        "classified_at": "Codex 分类结果生成时间。",
        "classification_history": "重新分类前的 Codex 结果历史，JSON 格式。",
        "research_profile": "当前研究项目使用的 Research Profile。",
        "qa_sample": "是否被固定 seed 选入 low/medium risk 质量抽样。",
        "qa_sample_reason": "QA 抽样原因、风险层和比例说明。",
        "qa_status": "QA 状态：not_selected、pending、confirmed 或 corrected。",
        "evidence_review_required": "评论作为核心证据或代表评论时是否必须人工确认。",
        "evidence_review_status": "证据复核状态：not_required、pending、confirmed 或 rejected。",
    }
)


def _record_row(record: CommentRecord | dict[str, Any]) -> dict[str, Any]:
    if isinstance(record, CommentRecord):
        return record.export_dict()
    return dict(record)


def _as_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value).strip().casefold() in {"1", "true", "yes"}


def _as_int(value: Any) -> int:
    try:
        return int(value or 0)
    except (TypeError, ValueError):
        return 0


def _is_generic_evidence(rows: list[dict[str, Any]]) -> bool:
    return any(str(row.get("evidence_id") or "").strip() for row in rows)


def _set_header(sheet: Worksheet, headers: Iterable[str]) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    sheet.row_dimensions[1].height = 32
    sheet.freeze_panes = "A2"


def _finish_table(sheet: Worksheet) -> None:
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _set_widths(sheet: Worksheet, widths: dict[int, float]) -> None:
    for index, width in widths.items():
        sheet.column_dimensions[get_column_letter(index)].width = width


def estimate_row_height(
    values: Iterable[tuple[Any, float]],
    *,
    minimum: float = 30.0,
    maximum: float = 150.0,
) -> float:
    """Estimate wrapped Excel row height without shortening cell contents."""
    visual_lines = 1
    for value, column_width in values:
        segments = str(value or "").splitlines() or [""]
        wrapped_lines = 0
        usable_width = max(8.0, column_width - 2.0)
        for segment in segments:
            units = 0.0
            for character in segment:
                if character == "\t":
                    units += 4.0
                elif ord(character) > 0xFFFF:
                    units += 2.0
                elif unicodedata.east_asian_width(character) in {"W", "F", "A"}:
                    units += 2.0
                else:
                    units += 1.0
            wrapped_lines += max(1, math.ceil(units / usable_width))
        visual_lines = max(visual_lines, wrapped_lines)
    return max(minimum, min(maximum, 15.0 * visual_lines + 8.0))


def _summary(text: Any, limit: int = 90) -> str:
    normalized = re.sub(r"\s+", " ", str(text or "")).strip()
    return normalized if len(normalized) <= limit else f"{normalized[: limit - 1]}…"


def _ordered_conversations(rows: list[dict[str, Any]]) -> list[tuple[str, dict[str, Any]]]:
    indexed = list(enumerate(rows))
    sort_key = lambda item: (_as_int(item[1].get("collection_order")), item[0])
    top_rows = sorted((item for item in indexed if not _as_bool(item[1].get("is_reply"))), key=sort_key)
    replies: dict[tuple[str, str], list[tuple[int, dict[str, Any]]]] = defaultdict(list)
    for item in indexed:
        row = item[1]
        if _as_bool(row.get("is_reply")):
            key = (str(row.get("video_id") or ""), str(row.get("parent_comment_id") or ""))
            replies[key].append(item)
    for values in replies.values():
        values.sort(key=sort_key)

    ordered: list[tuple[str, dict[str, Any]]] = []
    consumed: set[int] = set()
    for thread_number, (source_index, top) in enumerate(top_rows, start=1):
        thread = f"Thread {thread_number:03d}"
        ordered.append((thread, top))
        consumed.add(source_index)
        key = (str(top.get("video_id") or ""), str(top.get("comment_id") or ""))
        for reply_index, reply in replies.get(key, []):
            ordered.append((thread, reply))
            consumed.add(reply_index)

    for source_index, row in sorted(indexed, key=sort_key):
        if source_index not in consumed:
            thread = f"Thread orphan-{source_index + 1:03d}"
            ordered.append((thread, row))
    return ordered


def _add_overview(
    workbook: Workbook,
    rows: list[dict[str, Any]],
    classified: list[dict[str, Any]],
    *,
    profile: ResearchProfile,
    research_goal: str,
    research_questions: list[str],
    research_context: ResearchContext,
    theme_bundle: ThemeInsightBundle | None,
    human_review_bundle: HumanReviewBundle | None,
    decision_briefs: list[DecisionBriefRecord],
) -> None:
    sheet = workbook.create_sheet(SHEET_NAMES[0])
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:B1")
    generic_evidence = _is_generic_evidence(rows)
    sheet["A1"] = (
        "研究证据项目总览 Research Evidence Overview"
        if generic_evidence
        else "YouTube 评论研究项目总览"
    )
    sheet["A1"].fill = HEADER_FILL
    sheet["A1"].font = Font(color="FFFFFFFF", bold=True, size=16)
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.row_dimensions[1].height = 28

    top_rows = [row for row in rows if not _as_bool(row.get("is_reply"))]
    reply_rows = [row for row in rows if _as_bool(row.get("is_reply"))]
    top_rows_with_replies = [row for row in top_rows if _as_int(row.get("reply_count")) > 0]
    projects = sorted({str(row.get("project_id")) for row in rows if row.get("project_id")})
    collected = sorted({str(row.get("collected_at")) for row in rows if row.get("collected_at")})
    codex_rows = [row for row in classified if row.get("codex_primary_category")]
    human_rows = [
        row
        for row in classified
        if str(row.get("review_status") or "").casefold()
        in {"confirmed", "corrected"}
        and row.get("human_confirmed_primary_category")
    ]
    workload = summarize_review_workload(classified)
    videos = sorted(
        {
            f"{row.get('video_title') or row.get('video_id')} [{row.get('video_id')}]"
            for row in rows
            if row.get("video_id")
        }
    )
    sources = sorted(
        {
            f"{row.get('source_name') or row.get('source_type')} [{row.get('source_type')}]"
            for row in rows
            if row.get("source_type")
        }
    )
    if generic_evidence:
        effective = resolve_effective_insights(human_review_bundle or theme_bundle) if (human_review_bundle or theme_bundle) else []
        source_keys = {
            (str(row.get("source_name") or ""), str(row.get("source_type") or ""))
            for row in rows
        }
        source_types = sorted({str(row.get("source_type") or "") for row in rows if row.get("source_type")})
        metrics = [
            ("Research Setup", ""),
            ("Project ID", "；".join(projects)),
            ("Research Profile", f"{profile.profile_id} — {profile.profile_name}"),
            ("Research Goal", research_goal),
            ("Research Questions", "\n".join(f"• {item}" for item in research_questions)),
            ("Research Question", research_context.research_question or ""),
            ("Decision Context", research_context.decision_context or ""),
            ("Analysis Mode", research_context.analysis_mode),
            ("Evidence Sources", "\n".join(f"• {item}" for item in sources)),
            ("Evidence Count", len(rows)),
            ("Source Count", len(source_keys)),
            ("Source Types", "; ".join(source_types)),
            ("Theme Count", len(theme_bundle.themes) if theme_bundle else 0),
            ("Effective Insight Count", len(effective)),
            ("Decision Brief Count", len(decision_briefs)),
            ("Review Recommended Count", sum(item.review_recommended for item in effective)),
            ("Verification Status", "; ".join(
                f"{status}: {count}"
                for status, count in sorted(Counter(item.verification_status for item in effective).items())
            )),
            ("Data Limitations", "• Findings reflect only the supplied qualitative evidence.\n• Counts do not establish prevalence in a wider population.\n• Decision directions require validation before action."),
        ]
    else:
        metrics = [
        ("Research Setup", ""),
        ("项目 ID", "；".join(projects)),
        ("Research Profile", f"{profile.profile_id} — {profile.profile_name}"),
        ("Research Goal", research_goal),
        ("Research Questions", "\n".join(f"• {item}" for item in research_questions)),
        ("Research Question", research_context.research_question or ""),
        ("Decision Context", research_context.decision_context or ""),
        ("Analysis Mode", research_context.analysis_mode),
        ("Sources" if generic_evidence else "Videos", "\n".join(f"• {item}" for item in (sources if generic_evidence else videos))),
        ("Evidence Records" if generic_evidence else "Comments", len(rows)),
        ("采集时间", collected[0] if len(collected) == 1 else "；".join(collected)),
        ("视频数量", len({str(row.get("video_id")) for row in rows if row.get("video_id")})),
        ("总记录数", len(rows)),
        ("顶级评论数量", len(top_rows)),
        ("回复数量", len(reply_rows)),
        ("有回复的主评论数量", len(top_rows_with_replies)),
        ("已分析评论数量", len(classified)),
        ("Codex 辅助分类数量", len(codex_rows)),
        ("人工确认数量", len(human_rows)),
        ("Review Workload", ""),
        ("Total Comments", workload.total_comments),
        ("Low Risk Count", workload.low_risk_count),
        ("Medium Risk Count", workload.medium_risk_count),
        ("High Risk Count", workload.high_risk_count),
        ("Mandatory Review Count", workload.mandatory_review_count),
        ("QA Sample Count", workload.qa_sample_count),
        ("Evidence Review Count", workload.evidence_review_count),
        ("Unique Human Review Workload", workload.unique_human_review_workload),
        ("Human Review Coverage %", workload.human_review_coverage),
        ("颜色说明", "浅黄色 = 主评论；浅蓝色 = 回复"),
        ("数据限制说明", "\n".join(f"• {item}" for item in DATA_LIMITATIONS)),
        ]
    for row_number, (label, value) in enumerate(metrics, start=3):
        sheet.cell(row_number, 1, label)
        sheet.cell(row_number, 2, value)
        sheet.cell(row_number, 1).font = Font(bold=True)
        if label in {"Research Setup", "Review Workload"}:
            sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=2)
            sheet.cell(row_number, 1).fill = HEADER_FILL
            sheet.cell(row_number, 1).font = HEADER_FONT
        else:
            sheet.cell(row_number, 1).fill = SECTION_FILL
        for cell in sheet[row_number]:
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        if label in {"Research Questions", "Research Question", "Decision Context", "Videos", "Sources", "Evidence Sources", "Data Limitations", "数据限制说明"}:
            sheet.row_dimensions[row_number].height = 90
        else:
            sheet.row_dimensions[row_number].height = 24
        if label == "Human Review Coverage %":
            sheet.cell(row_number, 2).number_format = "0.0%"
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 84
    sheet.freeze_panes = "A3"


def _add_conversation(workbook: Workbook, rows: list[dict[str, Any]]) -> None:
    sheet = workbook.create_sheet(SHEET_NAMES[1])
    generic_evidence = _is_generic_evidence(rows)
    columns = (
        (
            ("thread", "证据组 Thread"),
            ("type", "记录类型 Type"),
            ("comment_text", "证据正文 Evidence Text"),
            ("reply_to", "关联证据 Related Evidence"),
            ("comment_like_count", "互动量 Engagement"),
            ("reply_count_display", "关联记录数 Related Count"),
            ("author_display_name", "来源主体 Source Actor"),
            ("comment_published_at", "时间 Timestamp"),
            ("video_title", "来源 Source"),
            ("channel_title", "组织 Organization"),
            ("comment_id", "证据 ID evidence_id"),
            ("parent_comment_id", "父证据 ID parent_evidence_id"),
        )
        if generic_evidence
        else CONVERSATION_COLUMNS
    )
    _set_header(sheet, (header for _, header in columns))
    top_by_id = {
        (str(row.get("video_id") or ""), str(row.get("comment_id") or "")): row
        for row in rows
        if not _as_bool(row.get("is_reply"))
    }
    for thread, row in _ordered_conversations(rows):
        is_reply = _as_bool(row.get("is_reply"))
        parent = top_by_id.get(
            (str(row.get("video_id") or ""), str(row.get("parent_comment_id") or ""))
        )
        values = {
            "thread": thread,
            "type": (
                "↳ 关联记录" if is_reply else "证据记录"
            ) if generic_evidence else ("↳ 回复" if is_reply else "主评论"),
            "comment_text": row.get("comment_text"),
            "reply_to": _summary(parent.get("comment_text")) if is_reply and parent else "",
            "comment_like_count": _as_int(row.get("comment_like_count")),
            "reply_count_display": "" if is_reply else _as_int(row.get("reply_count")),
            "author_display_name": row.get("author_display_name"),
            "comment_published_at": row.get("comment_published_at"),
            "video_title": row.get("video_title"),
            "channel_title": row.get("channel_title"),
            "comment_id": row.get("comment_id"),
            "parent_comment_id": row.get("parent_comment_id"),
        }
        sheet.append([values[key] for key, _ in columns])
        excel_row = sheet.max_row
        fill = REPLY_FILL if is_reply else MAIN_COMMENT_FILL
        for cell in sheet[excel_row]:
            cell.fill = fill
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if is_reply:
            sheet.cell(excel_row, 3).alignment = Alignment(vertical="top", wrap_text=True, indent=2)
        sheet.row_dimensions[excel_row].height = estimate_row_height(
            (
                (values["comment_text"], 61 if is_reply else 65),
                (values["reply_to"], 42),
            )
        )

    sheet.auto_filter.ref = sheet.dimensions
    _set_widths(
        sheet,
        {
            1: 15,
            2: 12,
            3: 65,
            4: 42,
            5: 11,
            6: 12,
            7: 20,
            8: 22,
            9: 34,
            10: 24,
            11: 26,
            12: 26,
        },
    )


def _add_classifications(
    workbook: Workbook,
    classified: list[dict[str, Any]],
    *,
    generic_evidence: bool = False,
) -> None:
    sheet = workbook.create_sheet(SHEET_NAMES[2])
    columns = GENERIC_CLASSIFICATION_COLUMNS if generic_evidence else CLASSIFICATION_COLUMNS
    _set_header(sheet, (header for _, header in columns))
    for row in classified:
        sheet.append([row.get(field) for field, _ in columns])
    _finish_table(sheet)
    width_by_field = {
        "comment_id": 26,
        "video_id": 18,
        "is_reply": 10,
        "comment_text": 64,
        "comment_like_count": 12,
        "rule_classification_label": 24,
        "rule_candidate_categories": 36,
        "rule_candidate_subtopics": 34,
        "detected_entities": 34,
        "detected_comparison_phrases": 32,
        "detected_lifecycle_signals": 30,
        "rule_evidence": 54,
        "codex_subject": 22,
        "codex_primary_category": 25,
        "codex_secondary_categories": 35,
        "codex_subtopics": 35,
        "codex_lifecycle_stage": 22,
        "codex_mentioned_brands": 28,
        "codex_mentioned_products": 30,
        "codex_evidence": 55,
        "codex_confidence": 16,
        "codex_review_risk": 16,
        "codex_manual_review": 18,
        "codex_review_triggers": 48,
        "qa_sample": 14,
        "qa_status": 16,
        "evidence_review_required": 24,
        "evidence_review_status": 24,
        "review_status": 18,
        "reviewer_note": 54,
        "codex_review_reason": 54,
        "human_confirmed_primary_category": 28,
        "human_confirmed_secondary_categories": 36,
        "human_confirmed_subtopics": 34,
        "human_confirmed_lifecycle_stage": 24,
        "reviewed_at": 24,
        "classification_source": 24,
        "classifier_version": 24,
        "rubric_version": 24,
        "classified_at": 25,
        "classification_history": 45,
        "research_profile": 24,
        "qa_sample_reason": 42,
        "qa_sample_seed": 18,
        "qa_sampled_at": 26,
    }
    widths = {
        index: width_by_field.get(field, 24)
        for index, (field, _) in enumerate(columns, start=1)
    }
    _set_widths(sheet, widths)
    confidence_column = next(
        index
        for index, (field, _) in enumerate(columns, start=1)
        if field == "codex_confidence"
    )
    for cell in sheet[get_column_letter(confidence_column)][1:]:
        cell.number_format = "0%"
    # Keep compatibility and verbose trace columns without letting them dominate.
    if not generic_evidence:
        history_column = next(
            index
            for index, (field, _) in enumerate(columns, start=1)
            if field == "classification_history"
        )
        for index in range(history_column, len(columns) + 1):
            sheet.column_dimensions[get_column_letter(index)].hidden = True
    for index, row in enumerate(classified, start=2):
        sheet.row_dimensions[index].height = estimate_row_height(
            (
                (row.get("comment_text"), 64),
                (row.get("codex_evidence"), 55),
                (row.get("codex_review_triggers"), 48),
                (row.get("codex_review_reason"), 54),
                (row.get("reviewer_note"), 54),
            )
        )


def _topic_insight_rows(
    classified: list[dict[str, Any]],
    insights: list[InsightRecord | dict[str, Any]],
) -> list[dict[str, Any]]:
    by_id = {str(row.get("comment_id") or ""): row for row in classified}
    output: list[dict[str, Any]] = []
    if insights:
        for item in insights:
            payload = item.model_dump() if isinstance(item, InsightRecord) else dict(item)
            section = insight_section(payload)
            evidence_ids = [str(value) for value in payload.get("evidence_comment_ids", [])]
            evidence_is_verified = bool(evidence_ids) and all(
                str(by_id.get(comment_id, {}).get("qa_status") or "").casefold()
                in {"confirmed", "corrected"}
                or str(
                    by_id.get(comment_id, {}).get("evidence_review_status") or ""
                ).casefold()
                == "confirmed"
                or str(by_id.get(comment_id, {}).get("review_status") or "").casefold()
                in {"confirmed", "corrected"}
                for comment_id in evidence_ids
            )
            if section == "Confirmed Findings" and not evidence_is_verified:
                section = "Insight Review Queue"
            representative_ids = list(payload.get("representative_comment_ids") or [])
            representative_id = representative_ids[0] if representative_ids else ""
            representative = by_id.get(representative_id, {})
            output.append(
                {
                    "report_section": section,
                    "theme": payload.get("theme", ""),
                    "evidence_count": payload.get("evidence_count", 0),
                    "evidence_videos": "; ".join(payload.get("source_video_ids", [])),
                    "representative_comment": representative.get("comment_text", ""),
                    "representative_comment_id": representative_id,
                    "insight": payload.get("insight", ""),
                    "action_opportunity": payload.get("action_opportunity", ""),
                    "priority": payload.get("priority", ""),
                    "priority_reason": payload.get("priority_reason", ""),
                    "status": (
                        payload.get("status", "")
                        if section == "Confirmed Findings"
                        else "preliminary"
                    ),
                    "evidence_review_status": (
                        payload.get("evidence_review_status", "")
                        if evidence_is_verified
                        else "pending"
                    ),
                    "needs_human_confirmation": (
                        section != "Confirmed Findings"
                        or payload.get("needs_human_confirmation", True)
                    ),
                }
            )
    else:
        section_map = {
            "Confirmed Findings": "Confirmed Findings",
            "Preliminary Codex-assisted Findings": "Preliminary Insights",
            "Review Queue": "Insight Review Queue",
        }
        for summary in build_review_sections(classified):
            representative_id = str(summary.get("representative_comment_id") or "")
            representative = by_id.get(representative_id, {})
            section = section_map.get(
                str(summary.get("report_section") or ""), "Preliminary Insights"
            )
            output.append(
                {
                    "report_section": section,
                    "theme": summary.get("primary_category") or summary.get("finding_status"),
                    "evidence_count": summary.get("count", 0),
                    "evidence_videos": str(representative.get("video_id") or ""),
                    "representative_comment": summary.get("representative_comment", ""),
                    "representative_comment_id": representative_id,
                    "insight": "",
                    "action_opportunity": "",
                    "priority": "",
                    "priority_reason": "",
                    "status": (
                        "confirmed" if section == "Confirmed Findings" else "preliminary"
                    ),
                    "evidence_review_status": representative.get(
                        "evidence_review_status", "not_required"
                    ),
                    "needs_human_confirmation": section != "Confirmed Findings",
                }
            )
    present = {str(row.get("report_section") or "") for row in output}
    for section in ("Confirmed Findings", "Preliminary Insights", "Insight Review Queue"):
        if section not in present:
            output.append(
                {
                    "report_section": section,
                    "theme": "",
                    "evidence_count": 0,
                    "evidence_videos": "",
                    "representative_comment": "",
                    "representative_comment_id": "",
                    "insight": "",
                    "action_opportunity": "",
                    "priority": "",
                    "priority_reason": "",
                    "status": "preliminary" if section != "Confirmed Findings" else "confirmed",
                    "evidence_review_status": "not_required",
                    "needs_human_confirmation": section != "Confirmed Findings",
                }
            )
    order = {"Confirmed Findings": 0, "Preliminary Insights": 1, "Insight Review Queue": 2}
    return sorted(output, key=lambda row: order.get(str(row.get("report_section")), 9))


def _add_topic_summary(
    workbook: Workbook,
    classified: list[dict[str, Any]],
    insights: list[InsightRecord | dict[str, Any]],
    *,
    generic_evidence: bool = False,
) -> None:
    sheet = workbook.create_sheet(
        GENERIC_SHEET_NAMES[3] if generic_evidence else SHEET_NAMES[3]
    )
    topic_fields = tuple(
        (
            field,
            {
                "report_section": "Output Status",
                "evidence_videos": "Evidence Sources",
                "representative_comment": "Representative Evidence",
                "representative_comment_id": "Representative Evidence ID",
                "evidence_review_status": "Verification Status",
                "needs_human_confirmation": "Review Recommended",
            }.get(field, header),
        )
        for field, header in TOPIC_FIELDS
    )
    _set_header(sheet, (header for _, header in topic_fields))
    by_id = {str(row.get("comment_id") or ""): row for row in classified}
    topic_rows = _topic_insight_rows(classified, insights)
    if generic_evidence:
        topic_rows = [row for row in topic_rows if int(row.get("evidence_count") or 0) > 0]
        section_names = {
            "Confirmed Findings": "Verified Insights",
            "Insight Review Queue": "Review Recommended",
        }
        for row in topic_rows:
            row["report_section"] = section_names.get(
                str(row.get("report_section") or ""),
                row.get("report_section"),
            )
            representative = by_id.get(str(row.get("representative_comment_id") or ""), {})
            row["evidence_videos"] = representative.get("source_name") or representative.get("source_type") or ""
            row["evidence_review_status"] = "unreviewed"
            row["needs_human_confirmation"] = (
                str(row.get("report_section") or "") == "Review Recommended"
            )
    for row in topic_rows:
        sheet.append([row.get(field) for field, _ in TOPIC_FIELDS])
    _finish_table(sheet)
    _set_widths(
        sheet,
        {
            1: 28,
            2: 34,
            3: 15,
            4: 28,
            5: 64,
            6: 28,
            7: 66,
            8: 66,
            9: 12,
            10: 58,
            11: 16,
            12: 24,
            13: 24,
        },
    )
    for index in range(2, sheet.max_row + 1):
        section = str(sheet.cell(index, 1).value or "")
        if section in {"Confirmed Findings", "Verified Insights"}:
            fill = PatternFill("solid", fgColor="FFE2F0D9")
        elif section == "Preliminary Insights":
            fill = PatternFill("solid", fgColor="FFDDEBF7")
        elif section in {"Insight Review Queue", "Review Recommended"}:
            fill = PatternFill("solid", fgColor="FFFFF2CC")
        else:
            fill = None
        if fill:
            for cell in sheet[index]:
                cell.fill = fill
        sheet.row_dimensions[index].height = estimate_row_height(
            (
                (sheet.cell(index, 5).value, 64),
                (sheet.cell(index, 7).value, 66),
                (sheet.cell(index, 8).value, 66),
                (sheet.cell(index, 10).value, 58),
            )
        )


def _profile_video_comparison(
    classified: list[dict[str, Any]], profile: ResearchProfile
) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in classified:
        grouped[str(row.get("video_id") or "")].append(row)
    output: list[dict[str, Any]] = []
    for video_id, rows in grouped.items():
        categories = Counter(
            str(row.get("effective_primary_category") or row.get("codex_primary_category") or "Other")
            for row in rows
        )
        topics: Counter[str] = Counter()
        for row in rows:
            topics.update(
                value.strip()
                for value in str(
                    row.get("effective_subtopics") or row.get("codex_subtopics") or ""
                ).split(";")
                if value.strip()
            )
        risk_counts = Counter(str(row.get("codex_review_risk") or "low").casefold() for row in rows)
        top = max(rows, key=lambda row: _as_int(row.get("comment_like_count")))
        result: dict[str, Any] = {
            "video_id": video_id,
            "video_title": rows[0].get("video_title", ""),
            "channel_title": rows[0].get("channel_title", ""),
            "record_count": len(rows),
            "top_level_count": sum(not _as_bool(row.get("is_reply")) for row in rows),
            "reply_count": sum(_as_bool(row.get("is_reply")) for row in rows),
            "top_themes": "; ".join(topic for topic, _ in topics.most_common(3)),
            "top_categories": "; ".join(category for category, _ in categories.most_common(3)),
            "review_risk_distribution": "; ".join(
                f"{risk}: {risk_counts.get(risk, 0)}" for risk in ("low", "medium", "high")
            ),
            "high_engagement_comment": top.get("comment_text", ""),
            "high_engagement_comment_id": top.get("comment_id", ""),
            "unique_discussion_topics": len(topics),
        }
        result.update({category: categories.get(category, 0) for category in profile.primary_categories})
        output.append(result)
    return output


def _add_video_comparison(
    workbook: Workbook,
    classified: list[dict[str, Any]],
    profile: ResearchProfile,
) -> None:
    sheet = workbook.create_sheet(SHEET_NAMES[4])
    base_fields: tuple[tuple[str, str], ...] = (
        ("video_id", "视频 ID"),
        ("video_title", "视频标题"),
        ("channel_title", "频道"),
        ("record_count", "采集评论数"),
        ("top_level_count", "顶级评论数"),
        ("reply_count", "回复数"),
    )
    category_fields = tuple((category, category) for category in profile.primary_categories)
    tail_fields: tuple[tuple[str, str], ...] = (
        ("top_themes", "Top Themes"),
        ("top_categories", "Top Categories"),
        ("review_risk_distribution", "Review Risk Distribution"),
        ("high_engagement_comment", "High Engagement Comments"),
        ("high_engagement_comment_id", "High Engagement Comment ID"),
        ("unique_discussion_topics", "Unique Discussion Topics"),
    )
    fields = base_fields + category_fields + tail_fields
    _set_header(sheet, (header for _, header in fields))
    for row in _profile_video_comparison(classified, profile):
        sheet.append([row.get(field) for field, _ in fields])
    _finish_table(sheet)
    sheet.freeze_panes = "D2"
    widths = {1: 18, 2: 36, 3: 24, 4: 14, 5: 14, 6: 12}
    category_start = 7
    for index in range(category_start, category_start + len(category_fields)):
        widths[index] = 24
    tail_start = category_start + len(category_fields)
    tail_widths = (38, 38, 34, 65, 28, 24)
    for offset, width in enumerate(tail_widths):
        widths[tail_start + offset] = width
    _set_widths(sheet, widths)


def _add_source_comparison(
    workbook: Workbook,
    classified: list[dict[str, Any]],
) -> None:
    sheet = workbook.create_sheet(GENERIC_SHEET_NAMES[4])
    fields = (
        ("source_name", "Source Name"),
        ("source_type", "Source Type"),
        ("evidence_count", "Evidence Count"),
        ("top_categories", "Top Categories"),
        ("top_subtopics", "Top Subtopics"),
        ("review_recommended", "Review Recommended"),
        ("representative_evidence", "Representative Evidence"),
        ("representative_evidence_id", "Representative Evidence ID"),
    )
    grouped: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in classified:
        key = (
            str(row.get("source_name") or "Unspecified source"),
            str(row.get("source_type") or "unspecified"),
        )
        grouped[key].append(row)
    _set_header(sheet, (header for _, header in fields))
    for (source_name, source_type), rows in sorted(grouped.items()):
        categories = Counter(
            str(row.get("effective_primary_category") or row.get("codex_primary_category") or "Other / Unclear")
            for row in rows
        )
        subtopics: Counter[str] = Counter()
        for row in rows:
            subtopics.update(
                item.strip()
                for item in str(row.get("effective_subtopics") or row.get("codex_subtopics") or "").split(";")
                if item.strip()
            )
        representative = max(rows, key=lambda row: _as_int(row.get("comment_like_count")))
        payload = {
            "source_name": source_name,
            "source_type": source_type,
            "evidence_count": len(rows),
            "top_categories": "; ".join(value for value, _ in categories.most_common(3)),
            "top_subtopics": "; ".join(value for value, _ in subtopics.most_common(3)),
            "review_recommended": sum(_as_bool(row.get("codex_manual_review")) for row in rows),
            "representative_evidence": representative.get("comment_text") or representative.get("text") or "",
            "representative_evidence_id": representative.get("evidence_id") or representative.get("comment_id") or "",
        }
        sheet.append([payload[field] for field, _ in fields])
    _finish_table(sheet)
    _set_widths(sheet, {1: 28, 2: 24, 3: 16, 4: 38, 5: 38, 6: 20, 7: 72, 8: 28})
    for index in range(2, sheet.max_row + 1):
        sheet.row_dimensions[index].height = estimate_row_height(
            ((sheet.cell(index, 7).value, 72),)
        )


def _add_raw_data(
    workbook: Workbook,
    rows: list[dict[str, Any]],
    *,
    generic_evidence: bool = False,
) -> None:
    sheet = workbook.create_sheet(SHEET_NAMES[5])
    extra_fields: list[str] = []
    for row in rows:
        for field in row:
            if field not in EXPORT_FIELDS and field not in extra_fields:
                extra_fields.append(field)
    fields = list(GENERIC_RAW_FIELDS) if generic_evidence else [*EXPORT_FIELDS, *extra_fields]
    _set_header(sheet, fields)
    for row in rows:
        sheet.append([row.get(field) for field in fields])
    _finish_table(sheet)
    for index, field in enumerate(fields, start=1):
        width = {
            "evidence_id": 28,
            "text": 80,
            "source_type": 24,
            "source_name": 32,
            "created_at": 26,
            "language": 14,
            "market": 16,
            "region": 18,
            "metadata": 48,
            "project_id": 20,
            "video_id": 18,
            "video_url": 36,
            "video_title": 42,
            "comment_text": 80,
            "author_display_name": 24,
            "comment_id": 28,
            "parent_comment_id": 28,
            "video_published_at": 24,
            "comment_published_at": 24,
            "comment_updated_at": 24,
            "collected_at": 24,
        }.get(field, max(12, min(len(field) + 2, 24)))
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "D2"


def _add_field_guide(workbook: Workbook, *, generic_evidence: bool = False) -> None:
    sheet = workbook.create_sheet(SHEET_NAMES[6])
    _set_header(sheet, ("字段 Field", "所在工作表 Sheet", "说明 Description"))
    generic_descriptions = {
        "evidence_id": "Stable source-neutral Evidence identifier.",
        "text": "Original Evidence text.",
        "source_type": "Type of source that supplied the Evidence.",
        "source_name": "Human-readable source name.",
        "created_at": "Original Evidence timestamp when available.",
        "language": "Recorded language metadata when available.",
        "market": "Recorded market metadata when available.",
        "region": "Recorded region metadata when available.",
        "metadata": "Source-specific metadata retained without reinterpretation.",
    }
    raw_fields = GENERIC_RAW_FIELDS if generic_evidence else EXPORT_FIELDS
    for field in raw_fields:
        sheet.append([
            field,
            "06_原始数据",
            generic_descriptions.get(field, RAW_FIELD_DESCRIPTIONS.get(field, "")),
        ])
    generic_report_descriptions = {
        "codex_subject": "Primary subject discussed in the Evidence.",
        "codex_primary_category": "Primary semantic category under the active Research Profile.",
        "codex_secondary_categories": "Additional categories directly supported by the Evidence.",
        "codex_subtopics": "Source-neutral subtopics directly supported by the Evidence.",
        "codex_lifecycle_stage": "Stage defined by the active Research Profile; independent of category.",
        "codex_confidence": "Confidence in the primary-category judgment; not a statistical probability.",
        "codex_review_risk": "Independent semantic review risk: low, medium, or high.",
        "codex_manual_review": "Whether optional verification is recommended.",
        "codex_review_triggers": "Concise reasons that raised review risk.",
        "codex_review_reason": "Human-readable explanation for the verification recommendation.",
        "codex_evidence": "Short excerpts copied from the original Evidence.",
        "classification_source": "Classification workflow that produced the semantic result.",
        "classifier_version": "Version identifier for the semantic classifier.",
        "rubric_version": "Version identifier for the applied classification rubric.",
        "classified_at": "Timestamp recorded when the classification was produced.",
        "research_profile": "Research Profile used to validate categories, subjects, and stages.",
    }
    descriptions = (
        generic_report_descriptions.items()
        if generic_evidence
        else REPORT_FIELD_DESCRIPTIONS.items()
    )
    for field, description in descriptions:
        source_sheet = (
            "02_对话视图"
            if field in {"Thread", "reply_to"}
            else "03_语义标注"
        )
        sheet.append([field, source_sheet, description])
    concepts = {
        "Research Profile": "定义研究目标、问题、taxonomy、阶段、复核策略和 insight questions 的配置。",
        "Category": "当前 profile 下对单条证据主要交流意图的语义标签。",
        "Stage": "与主分类独立的用户、玩家、受众或采用阶段。",
        "Confidence": "主分类判断把握；confidence ≠ probability，confidence ≠ no human review required。",
        "Review Risk": "独立于 confidence 的 low / medium / high 复核风险。",
        "Manual Review": "分类层的高风险质量控制信号；不构成所有下游 Insight 的统一人工门槛。",
        "QA Sample": "从 low/medium risk 中按固定 seed 抽取的质量检查样本。",
        "Evidence Review": "核心证据、代表评论或推荐证据的人工确认。",
        "Preliminary": "AI-assisted 初步结果；preliminary ≠ confirmed finding。",
        "Verified": "经过可选人工核验或证据核查的结果；unreviewed 仍可作为明确标注的初步输出。",
        "Verification Status": "Insight 的可选人工核查状态：unreviewed、reviewed、edited 或 rejected。",
        "Review Recommended": "基于低置信度、证据较少、emergent taxonomy coverage 或追溯警告给出的可解释核查建议；不等于强制审核。",
        "Effective Insight": "统一的下游 Insight：默认包含 unreviewed、reviewed、edited，排除 rejected；显式 strict mode 除外。",
        "Priority": "综合证据频次、来源复现、研究目标相关性、影响、证据质量与核验状态的 P0/P1/P2。",
    }
    if generic_evidence:
        concepts.pop("Manual Review", None)
        concepts.pop("QA Sample", None)
        concepts.pop("Evidence Review", None)
    for field, description in concepts.items():
        sheet.append([field, "全工作簿 Workbook", description])
    _finish_table(sheet)
    _set_widths(sheet, {1: 30, 2: 32, 3: 90})


def _add_theme_discovery(
    workbook: Workbook,
    theme_bundle: ThemeInsightBundle | None,
) -> None:
    sheet = workbook.create_sheet(SHEET_NAMES[7])
    headers = (
        "Theme",
        "Description",
        "Evidence Count",
        "Related Categories",
        "Discovery Method",
        "Taxonomy Coverage",
        "Coverage Reason",
        "Confidence",
        "Review Status",
        "Representative Comments",
        "Representative Evidence IDs",
        "Evidence IDs",
    )
    _set_header(sheet, headers)
    if theme_bundle and theme_bundle.themes:
        for theme in theme_bundle.themes:
            sheet.append(
                [
                    theme.theme_name,
                    theme.theme_description,
                    theme.evidence_count,
                    "; ".join(theme.related_categories),
                    theme.discovery_method,
                    theme.taxonomy_coverage,
                    theme.coverage_reason,
                    theme.confidence,
                    theme.review_status,
                    "\n\n".join(theme.representative_comments),
                    "; ".join(theme.representative_evidence_ids),
                    "; ".join(theme.evidence_comment_ids),
                ]
            )
    else:
        sheet.append(
            [
                "",
                "No Theme Discovery bundle supplied. Generate preliminary themes with generate-insights.",
                0,
                "",
                "",
                "",
                "",
                "",
                "preliminary",
                "",
                "",
                "",
            ]
        )
    _finish_table(sheet)
    _set_widths(
        sheet,
        {
            1: 30,
            2: 58,
            3: 15,
            4: 34,
            5: 20,
            6: 20,
            7: 58,
            8: 14,
            9: 16,
            10: 72,
            11: 38,
            12: 55,
        },
    )
    for index in range(2, sheet.max_row + 1):
        sheet.row_dimensions[index].height = estimate_row_height(
            (
                (sheet.cell(index, 2).value, 58),
                (sheet.cell(index, 7).value, 58),
                (sheet.cell(index, 10).value, 72),
                (sheet.cell(index, 12).value, 55),
            )
        )


def _add_insight_drafts(
    workbook: Workbook,
    theme_bundle: ThemeInsightBundle | None,
    human_review_bundle: HumanReviewBundle | None,
    *,
    generic_evidence: bool = False,
) -> None:
    sheet = workbook.create_sheet(SHEET_NAMES[8])
    headers = (
        "Theme",
        "Observation",
        "Interpretation",
        "Validation Cue" if generic_evidence else "Review Note (legacy Potential Opportunity)",
        "Confidence",
        "Verification Status",
        "Review Recommended",
        "Review Reasons",
        "Review Status",
    )
    _set_header(sheet, headers)
    if theme_bundle and theme_bundle.insight_drafts:
        themes_by_id = {theme.theme_id: theme for theme in theme_bundle.themes}
        effective_source = human_review_bundle or theme_bundle
        effective_by_id = {
            item.source_insight_id: item
            for item in resolve_effective_insights(
                effective_source,
                include_rejected=True,
            )
        }
        for insight in theme_bundle.insight_drafts:
            theme = themes_by_id[insight.theme_id]
            effective = effective_by_id[insight.insight_id]
            sheet.append(
                [
                    theme.theme_name,
                    insight.observation,
                    insight.interpretation,
                    insight.potential_opportunity,
                    insight.confidence,
                    effective.verification_status,
                    effective.review_recommended,
                    "\n".join(effective.review_reasons),
                    insight.review_status,
                ]
            )
    else:
        sheet.append(
            [
                "",
                "No Insight Draft bundle supplied.",
                "AI-assisted interpretations are unreviewed and remain eligible for downstream use.",
                "",
                "",
                "unreviewed",
                False,
                "",
                "preliminary",
            ]
        )
    _finish_table(sheet)
    _set_widths(
        sheet,
        {1: 30, 2: 62, 3: 70, 4: 75, 5: 16, 6: 20, 7: 20, 8: 54, 9: 18},
    )
    for index in range(2, sheet.max_row + 1):
        sheet.row_dimensions[index].height = estimate_row_height(
            (
                (sheet.cell(index, 2).value, 62),
                (sheet.cell(index, 3).value, 70),
                (sheet.cell(index, 4).value, 75),
                (sheet.cell(index, 8).value, 54),
            )
        )


def _add_human_review(
    workbook: Workbook,
    review_bundle: HumanReviewBundle | None,
    *,
    generic_evidence: bool = False,
) -> None:
    sheet = workbook.create_sheet(GENERIC_SHEET_NAMES[9] if generic_evidence else SHEET_NAMES[9])
    headers = (
        "Verification ID" if generic_evidence else "Review ID",
        "Theme",
        "Verification Action" if generic_evidence else "Review Action",
        "Original Observation",
        "Final Observation",
        "Original Interpretation",
        "Final Interpretation",
        "Evidence Count",
        "Verification Note / Rejection Reason" if generic_evidence else "Reviewer Note / Rejection Reason",
        "Verification Status",
        "Reviewed At",
    )
    _set_header(sheet, headers)
    if review_bundle and review_bundle.review_items:
        decisions = {
            decision.review_id: decision
            for decision in review_bundle.review_decisions
        }
        confirmed = {
            insight.review_id: insight
            for insight in review_bundle.confirmed_insights
        }
        for item in review_bundle.review_items:
            decision = decisions.get(item.review_id)
            final = confirmed.get(item.review_id)
            if decision is None:
                action = ""
                theme_name = item.theme.theme_name
                final_observation = ""
                final_interpretation = ""
                evidence_count = item.theme.evidence_count
                reviewer_note = ""
                review_status = "Unreviewed"
                reviewed_at = ""
            elif decision.action == "reject":
                action = "Reject"
                theme_name = item.theme.theme_name
                final_observation = ""
                final_interpretation = ""
                evidence_count = item.theme.evidence_count
                reviewer_note = decision.rejection_reason or decision.reviewer_note
                review_status = "Rejected"
                reviewed_at = decision.reviewed_at
            else:
                action = "Confirm" if decision.action == "confirm" else "Edit"
                theme_name = final.theme_name if final else item.theme.theme_name
                final_observation = final.observation if final else ""
                final_interpretation = final.interpretation if final else ""
                evidence_count = len(final.evidence_ids) if final else 0
                reviewer_note = decision.reviewer_note
                review_status = (
                    "Reviewed"
                    if decision.action == "confirm"
                    else "Edited"
                )
                reviewed_at = decision.reviewed_at
            sheet.append(
                [
                    item.review_id,
                    theme_name,
                    action,
                    item.insight_draft.observation,
                    final_observation,
                    item.insight_draft.interpretation,
                    final_interpretation,
                    evidence_count,
                    reviewer_note,
                    review_status,
                    reviewed_at,
                ]
            )
    else:
        sheet.append(
            [
                "",
                "No Optional Verification bundle supplied." if generic_evidence else "No optional Human Review bundle supplied.",
                "",
                "AI Theme and Insight Draft artifacts remain preliminary.",
                "",
                "",
                "",
                0,
                "",
                "Unreviewed",
                "",
            ]
        )
    _finish_table(sheet)
    _set_widths(
        sheet,
        {
            1: 28,
            2: 34,
            3: 16,
            4: 62,
            5: 62,
            6: 72,
            7: 72,
            8: 16,
            9: 52,
            10: 20,
            11: 28,
        },
    )
    for index in range(2, sheet.max_row + 1):
        sheet.row_dimensions[index].height = estimate_row_height(
            (
                (sheet.cell(index, 4).value, 62),
                (sheet.cell(index, 5).value, 62),
                (sheet.cell(index, 6).value, 72),
                (sheet.cell(index, 7).value, 72),
                (sheet.cell(index, 9).value, 52),
            )
        )


def _format_brief_items(items: Iterable[Any], fields: tuple[str, ...]) -> str:
    lines: list[str] = []
    for index, item in enumerate(items, start=1):
        values = [str(getattr(item, field, "") or "") for field in fields]
        lines.append(f"{index}. " + "\n".join(values))
    return "\n\n".join(lines)


def _add_decision_support(
    workbook: Workbook,
    decision_briefs: list[DecisionBriefRecord],
) -> None:
    sheet = workbook.create_sheet(SHEET_NAMES[10])
    headers = (
        "Decision Brief ID",
        "Research Question",
        "Source Effective Insights",
        "Verification Status",
        "Review Recommended",
        "Review Reasons",
        "Why It Matters",
        "Decision Question",
        "Possible Directions",
        "Evidence Gaps",
        "Validation Plan",
        "Supporting Evidence IDs",
        "Uncertainty Note",
        "Status",
    )
    _set_header(sheet, headers)
    if decision_briefs:
        for brief in decision_briefs:
            sheet.append(
                [
                    brief.decision_brief_id,
                    brief.research_question or "",
                    "\n".join(brief.source_effective_insight_ids),
                    ", ".join(brief.verification_statuses),
                    brief.review_recommended,
                    "\n".join(brief.review_reasons),
                    brief.why_it_matters,
                    brief.decision_question,
                    _format_brief_items(
                        brief.possible_directions,
                        ("direction", "rationale", "evidence_basis", "tradeoff_or_risk"),
                    ),
                    _format_brief_items(
                        brief.evidence_gaps,
                        ("gap", "why_it_matters"),
                    ),
                    _format_brief_items(
                        brief.validation_plan,
                        (
                            "validation_question",
                            "method",
                            "required_evidence",
                            "evaluation_criterion",
                            "limitation",
                        ),
                    ),
                    "\n".join(brief.supporting_evidence_ids),
                    brief.uncertainty_note,
                    brief.status,
                ]
            )
    else:
        sheet.append(
            [
                "",
                "",
                "",
                "",
                False,
                "",
                "No Decision Brief supplied.",
                "",
                "",
                "",
                "",
                "",
                "",
                "AI-generated decision support",
            ]
        )
    _finish_table(sheet)
    _set_widths(
        sheet,
        {
            1: 26, 2: 36, 3: 34, 4: 22, 5: 18, 6: 38, 7: 58,
            8: 48, 9: 64, 10: 58, 11: 64, 12: 34, 13: 58, 14: 28,
        },
    )
    for index in range(2, sheet.max_row + 1):
        sheet.row_dimensions[index].height = estimate_row_height(
            tuple(
                (sheet.cell(index, column).value, width)
                for column, width in ((7, 58), (8, 48), (9, 64), (10, 58), (11, 64), (13, 58))
            )
        )


def create_readable_workbook(
    records: Iterable[CommentRecord | dict[str, Any]],
    path: Path,
    *,
    profile_id: str = DEFAULT_PROFILE_ID,
    research_goal: str | None = None,
    research_questions: list[str] | None = None,
    insights: list[InsightRecord | dict[str, Any]] | None = None,
    theme_bundle: ThemeInsightBundle | None = None,
    human_review_bundle: HumanReviewBundle | None = None,
    decision_briefs: list[DecisionBriefRecord] | None = None,
    review_sample_seed: int = 20260816,
    research_context: ResearchContext | None = None,
) -> Path:
    rows = [_record_row(record) for record in records]
    profile = load_profile(profile_id)
    insight_rows = list(insights or [])
    classified = prepare_report_rows(
        rows,
        profile_id=profile.profile_id,
        review_sample_seed=review_sample_seed,
    )
    classified = mark_insight_evidence(classified, insight_rows)
    classified = apply_review_sampling(
        classified,
        profile,
        seed=review_sample_seed,
    )
    if human_review_bundle and human_review_bundle.profile_id != profile.profile_id:
        raise ValueError(
            "Human Review bundle profile does not match workbook profile."
        )
    if (
        theme_bundle
        and human_review_bundle
        and theme_bundle != human_review_bundle.source_bundle
    ):
        raise ValueError(
            "Theme bundle does not match the source bundle in Human Review results."
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    effective_theme_bundle = theme_bundle or (
        human_review_bundle.source_bundle if human_review_bundle else None
    )
    effective_research_context = research_context or (
        effective_theme_bundle.research_context
        if effective_theme_bundle
        else ResearchContext()
    )
    generic_evidence = _is_generic_evidence(rows)
    decision_brief_rows = list(decision_briefs or [])

    workbook = Workbook()
    workbook.remove(workbook.active)
    _add_overview(
        workbook,
        rows,
        classified,
        profile=profile,
        research_goal=research_goal or profile.research_goal,
        research_questions=research_questions or profile.research_questions,
        research_context=effective_research_context,
        theme_bundle=effective_theme_bundle,
        human_review_bundle=human_review_bundle,
        decision_briefs=decision_brief_rows,
    )
    _add_conversation(workbook, rows)
    _add_classifications(workbook, classified, generic_evidence=generic_evidence)
    _add_topic_summary(
        workbook,
        classified,
        insight_rows,
        generic_evidence=generic_evidence,
    )
    if generic_evidence:
        _add_source_comparison(workbook, classified)
    else:
        _add_video_comparison(workbook, classified, profile)
    _add_raw_data(workbook, rows, generic_evidence=generic_evidence)
    _add_field_guide(workbook, generic_evidence=generic_evidence)
    _add_theme_discovery(workbook, effective_theme_bundle)
    _add_insight_drafts(
        workbook,
        effective_theme_bundle,
        human_review_bundle,
        generic_evidence=generic_evidence,
    )
    _add_human_review(
        workbook,
        human_review_bundle,
        generic_evidence=generic_evidence,
    )
    _add_decision_support(workbook, decision_brief_rows)
    workbook.save(path)
    return path
