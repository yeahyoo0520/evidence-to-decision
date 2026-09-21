from __future__ import annotations

import csv
import json
import re
from collections import Counter, defaultdict
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from .classification import (
    ClassificationResult as ScoredClassificationResult,
    classify_text as classify_scored_text,
    detect_subtopics as detect_scored_subtopics,
)
from .evidence import evidence_to_research_rows, load_evidence_json
from .models import CommentRecord

CATEGORY_LABELS: tuple[tuple[str, str], ...] = (
    ("Product Praise", "产品认可"),
    ("Product Concern", "产品顾虑"),
    ("Product Question", "产品问题"),
    ("Purchase Question", "购买疑问"),
    ("Usage Scenario", "使用场景"),
    ("Competitor Comparison", "竞品比较"),
    ("Purchase Intent", "购买意向"),
    ("Feature Request", "功能需求"),
    ("After-sales Issue", "售后问题"),
    ("Other", "其他"),
)
CATEGORY_CN = dict(CATEGORY_LABELS)

DATA_LIMITATIONS: tuple[str, ...] = (
    "仅包含 YouTube Data API 返回的公开可访问评论，不包含已删除、私密、受限或未返回的评论。",
    "评论是研究样本，不应被描述为具有代表性的消费者调查。",
    "分类基于确定性规则；低置信度结果需要人工复核。",
    "主题汇总仅复述评论中可追溯的信号，不构成自动营销结论。",
)

COMPARISON_PATTERNS = (
    " vs ",
    " versus ",
    "better than",
    "compare",
    "compared with",
    "compared to",
    "which is better",
    "对比",
    "相比",
    "哪个好",
    "孰优",
)
AFTER_SALES_PATTERNS = (
    "warranty",
    "return",
    "refund",
    "replacement",
    "customer support",
    "after-sales",
    "after sales",
    "售后",
    "退货",
    "退款",
    "保修",
    "换货",
)
FEATURE_REQUEST_PATTERNS = (
    "wish it",
    "please add",
    "should add",
    "should have",
    "would like",
    "feature request",
    "希望增加",
    "希望有",
    "能不能增加",
    "建议增加",
)
PURCHASE_INTENT_PATTERNS = (
    "i ordered",
    "i bought",
    "i'm buying",
    "i am buying",
    "will buy",
    "will get one",
    "take my money",
    "下单了",
    "已经买了",
    "准备买",
    "打算买",
    "想买",
)
PURCHASE_QUESTION_PATTERNS = (
    "worth the price",
    "worth buying",
    "worth it",
    "how much",
    "what price",
    "where can i buy",
    "where to buy",
    "available in",
    "discount",
    "shipping",
    "多少钱",
    "价格多少",
    "值得买吗",
    "哪里买",
    "有优惠",
    "什么时候上市",
)
PRODUCT_FEATURE_PATTERNS = (
    "battery",
    "waterproof",
    "water resistant",
    "bluetooth",
    "multipoint",
    "connect",
    "connection",
    "compatible",
    "compatibility",
    "microphone",
    " mic ",
    "sound",
    "volume",
    "fit",
    "glasses",
    "latency",
    "charge",
    "功能",
    "续航",
    "防水",
    "蓝牙",
    "连接",
    "兼容",
    "麦克风",
    "音质",
    "音量",
    "眼镜",
    "延迟",
    "充电",
)
QUESTION_PREFIXES = (
    "how ",
    "does ",
    "do ",
    "can ",
    "could ",
    "is ",
    "are ",
    "what ",
    "which ",
    "will ",
)
QUESTION_CN_SIGNALS = ("是否", "能否", "怎么", "如何", "吗", "么")
USAGE_PATTERNS = (
    "running",
    "run ",
    "cycling",
    "bike",
    "commute",
    "office",
    "gym",
    "workout",
    "travel",
    "swimming",
    "hiking",
    "跑步",
    "骑行",
    "通勤",
    "办公室",
    "健身",
    "旅行",
    "游泳",
    "徒步",
)
USAGE_RELATION_PATTERNS = ("use ", "using ", "wear ", "for ", "用于", "用来", "戴着")
CONCERN_PATTERNS = (
    "problem",
    "issue",
    "uncomfortable",
    "hurt",
    "pain",
    "leak",
    "overheat",
    "too expensive",
    "too loud",
    "not work",
    "stopped working",
    "poor",
    "bad",
    "担心",
    "问题",
    "不舒服",
    "疼",
    "漏音",
    "过热",
    "太贵",
    "失灵",
    "不好",
)
PRAISE_PATTERNS = (
    "love",
    "great",
    "excellent",
    "amazing",
    "comfortable",
    "awesome",
    "fantastic",
    "works well",
    "好用",
    "喜欢",
    "舒服",
    "很棒",
    "优秀",
    "不错",
)

SUBTOPIC_RULES: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("价格与价值 Price/value", ("price", "cost", "worth", "expensive", "价格", "多少钱", "太贵")),
    ("购买渠道与供货 Availability", ("where to buy", "available", "shipping", "discount", "哪里买", "上市", "优惠")),
    ("电池与充电 Battery/charging", ("battery", "charge", "charging", "续航", "电池", "充电")),
    ("音质与音量 Audio/volume", ("sound", "audio", "volume", "音质", "声音", "音量", "漏音")),
    ("麦克风与通话 Microphone/calls", ("microphone", " mic ", "call", "麦克风", "通话")),
    ("佩戴与舒适 Comfort/fit", ("comfortable", "uncomfortable", "fit", "glasses", "hurt", "舒适", "舒服", "佩戴", "眼镜", "疼")),
    ("连接与兼容 Connectivity", ("bluetooth", "multipoint", "connect", "compatible", "latency", "蓝牙", "连接", "兼容", "延迟")),
    ("防水与耐用 Waterproof/durability", ("waterproof", "water resistant", "broken", "durable", "防水", "损坏", "耐用")),
    ("跑步与健身 Running/fitness", ("running", "run ", "gym", "workout", "跑步", "健身")),
    ("骑行 Cycling", ("cycling", "bike", "骑行", "自行车")),
    ("通勤与办公 Commute/office", ("commute", "office", "通勤", "办公室")),
    ("售后与退换 Warranty/returns", ("warranty", "return", "refund", "replacement", "售后", "退货", "退款", "保修", "换货")),
    ("产品比较 Product comparison", COMPARISON_PATTERNS),
)

STOPWORDS = {
    "the",
    "and",
    "for",
    "that",
    "this",
    "with",
    "you",
    "your",
    "are",
    "was",
    "have",
    "has",
    "but",
    "not",
    "they",
    "from",
    "its",
    "just",
    "can",
    "will",
    "would",
    "what",
    "when",
    "how",
    "why",
    "video",
    "really",
    "one",
    "get",
    "got",
    "about",
    "too",
    "all",
    "more",
    "than",
    "use",
    "using",
    "these",
    "those",
    "into",
}


@dataclass(frozen=True)
class LegacyClassificationResult:
    primary_category: str
    category_cn: str
    subtopic: str
    category_confidence: float
    category_evidence: str
    manual_review: bool

    def as_dict(self) -> dict[str, Any]:
        return asdict(self)


def _normalized(text: str) -> str:
    collapsed = re.sub(r"\s+", " ", text.casefold()).strip()
    return f" {collapsed} "


def _first_match(normalized: str, patterns: Iterable[str]) -> str | None:
    return next((pattern for pattern in patterns if pattern in normalized), None)


def _question_signal(text: str) -> str | None:
    stripped = re.sub(r"\s+", " ", text.casefold()).strip()
    if "?" in stripped or "？" in stripped:
        return "question mark"
    prefix = next((item for item in QUESTION_PREFIXES if stripped.startswith(item)), None)
    if prefix:
        return prefix.strip()
    return next((item for item in QUESTION_CN_SIGNALS if item in stripped), None)


def detect_subtopic(text: str) -> str:
    """Return the first independently detected Phase 2A.1 subtopic."""
    topics = detect_scored_subtopics(text)
    return topics[0] if topics else "general"


def _result(
    category: str,
    confidence: float,
    evidence: str,
    *,
    text: str,
) -> LegacyClassificationResult:
    return LegacyClassificationResult(
        primary_category=category,
        category_cn=CATEGORY_CN[category],
        subtopic=detect_subtopic(text),
        category_confidence=confidence,
        category_evidence=evidence,
        manual_review=confidence < 0.65,
    )


def _legacy_classify_text(text: str) -> LegacyClassificationResult:
    """Classify only explicit signals found in the supplied comment text."""
    normalized = _normalized(text)

    match = _first_match(normalized, COMPARISON_PATTERNS)
    if match:
        return _result(
            "Competitor Comparison",
            0.92,
            f'Explicit comparison relation matched: "{match.strip()}".',
            text=text,
        )

    match = _first_match(normalized, AFTER_SALES_PATTERNS)
    if match:
        return _result(
            "After-sales Issue",
            0.90,
            f'After-sales signal matched: "{match.strip()}".',
            text=text,
        )

    match = _first_match(normalized, FEATURE_REQUEST_PATTERNS)
    if match:
        return _result(
            "Feature Request",
            0.88,
            f'Explicit feature request matched: "{match.strip()}".',
            text=text,
        )

    match = _first_match(normalized, PURCHASE_INTENT_PATTERNS)
    if match:
        return _result(
            "Purchase Intent",
            0.88,
            f'Purchase-intent phrase matched: "{match.strip()}".',
            text=text,
        )

    match = _first_match(normalized, PURCHASE_QUESTION_PATTERNS)
    if match:
        return _result(
            "Purchase Question",
            0.86,
            f'Purchase decision signal matched: "{match.strip()}".',
            text=text,
        )

    feature_match = _first_match(normalized, PRODUCT_FEATURE_PATTERNS)
    question_match = _question_signal(text)
    if feature_match and question_match:
        return _result(
            "Product Question",
            0.82,
            f'Product feature "{feature_match.strip()}" is asked about explicitly.',
            text=text,
        )

    usage_match = _first_match(normalized, USAGE_PATTERNS)
    usage_relation = _first_match(normalized, USAGE_RELATION_PATTERNS)
    if usage_match and usage_relation:
        return _result(
            "Usage Scenario",
            0.82,
            f'Usage scenario matched: "{usage_match.strip()}".',
            text=text,
        )

    match = _first_match(normalized, CONCERN_PATTERNS)
    if match:
        return _result(
            "Product Concern",
            0.82,
            f'Product concern signal matched: "{match.strip()}".',
            text=text,
        )

    match = _first_match(normalized, PRAISE_PATTERNS)
    if match:
        return _result(
            "Product Praise",
            0.80,
            f'Positive product signal matched: "{match.strip()}".',
            text=text,
        )

    if question_match:
        return _result(
            "Other",
            0.40,
            "A question was detected, but no product-feature or purchase-decision signal was found.",
            text=text,
        )

    return _result(
        "Other",
        0.30,
        "No deterministic category rule matched.",
        text=text,
    )


def classify_text(text: str) -> ScoredClassificationResult:
    """Classify text with the Phase 2A.1 scored, multi-label engine."""
    return classify_scored_text(text)


def classify_rows(rows: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    classified: list[dict[str, Any]] = []
    for row in rows:
        result = classify_text(str(row.get("comment_text") or ""))
        classified.append({**row, **result.as_dict()})
    return classified


def _as_int(value: Any) -> int:
    try:
        return int(value or 0)
    except (TypeError, ValueError):
        return 0


def build_topic_summary(classified_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    result: list[dict[str, Any]] = []
    limitation = "规则分类仅作初步结果；代表评论按点赞数选取，不代表整体消费者意见。"
    scopes = (
        ("确认结果（规则初筛）", False, True),
        ("待人工复核（不计入主汇总）", True, False),
    )
    for scope_name, review_value, is_primary_summary in scopes:
        scoped = [
            row
            for row in classified_rows
            if _as_bool(row.get("manual_review")) is review_value
        ]
        for category, category_cn in CATEGORY_LABELS:
            matching = [row for row in scoped if row.get("primary_category") == category]
            subtopics: Counter[str] = Counter()
            for row in matching:
                for topic in str(row.get("subtopics") or row.get("subtopic") or "").split(";"):
                    if topic.strip() and topic.strip() != "general":
                        subtopics[topic.strip()] += 1
            representative = max(
                matching,
                key=lambda row: _as_int(row.get("comment_like_count")),
                default=None,
            )
            result.append(
                {
                    "summary_scope": scope_name,
                    "is_primary_summary": is_primary_summary,
                    "primary_category": category,
                    "category_cn": category_cn,
                    "count": len(matching),
                    "share": len(matching) / len(scoped) if scoped else 0,
                    "high_frequency_subtopics": "；".join(
                        f"{name} ({count})" for name, count in subtopics.most_common(3)
                    ),
                    "representative_comment": (representative or {}).get("comment_text"),
                    "representative_comment_id": (representative or {}).get("comment_id"),
                    "data_limitation": limitation,
                }
            )
    return result


def build_video_comparison(classified_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in classified_rows:
        grouped[str(row.get("video_id") or "")].append(row)

    result: list[dict[str, Any]] = []
    for video_id, rows in grouped.items():
        top = [row for row in rows if not _as_bool(row.get("is_reply"))]
        replies = [row for row in rows if _as_bool(row.get("is_reply"))]
        category_counts = Counter(
            str(
                row.get("effective_primary_category")
                or row.get("primary_category")
                or "Other"
            )
            for row in rows
        )
        subtopics: Counter[str] = Counter()
        for row in rows:
            for topic in str(
                row.get("effective_subtopics")
                or row.get("subtopics")
                or row.get("subtopic")
                or ""
            ).split(";"):
                if topic.strip() and topic.strip() != "general":
                    subtopics[topic.strip()] += 1
        top_liked = max(rows, key=lambda row: _as_int(row.get("comment_like_count")), default={})
        item: dict[str, Any] = {
            "video_id": video_id,
            "video_title": rows[0].get("video_title") if rows else "",
            "channel_title": rows[0].get("channel_title") if rows else "",
            "record_count": len(rows),
            "top_level_count": len(top),
            "reply_count": len(replies),
        }
        item.update({category: category_counts.get(category, 0) for category, _ in CATEGORY_LABELS})
        item["top_liked_comment"] = top_liked.get("comment_text")
        item["top_liked_comment_id"] = top_liked.get("comment_id")
        item["most_common_subtopic"] = subtopics.most_common(1)[0][0] if subtopics else ""
        result.append(item)
    return result


def _as_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value).strip().casefold() in {"1", "true", "yes"}


def load_rows(path: Path) -> list[dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".json":
        return evidence_to_research_rows(load_evidence_json(path))
    if suffix == ".jsonl":
        return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
    if suffix == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return list(csv.DictReader(handle))
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        if "06_原始数据" in workbook.sheetnames:
            sheet = workbook["06_原始数据"]
        elif "comments" in workbook.sheetnames:
            sheet = workbook["comments"]
        else:
            sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        header_row = next(iterator, None)
        if not header_row:
            return []
        header = [str(value) if value is not None else "" for value in header_row]
        return [dict(zip(header, values, strict=False)) for values in iterator]
    raise ValueError(f"Unsupported input format: {path.suffix}")


def rows_to_records(rows: Iterable[dict[str, Any]]) -> list[CommentRecord]:
    records: list[CommentRecord] = []
    model_fields = set(CommentRecord.model_fields)
    for row in rows:
        normalized = {
            key: (None if value == "" else value)
            for key, value in row.items()
            if key in model_fields
        }
        records.append(CommentRecord.model_validate(normalized))
    return records


def top_terms(rows: Iterable[dict[str, Any]], limit: int = 30) -> list[tuple[str, int]]:
    terms: Counter[str] = Counter()
    for row in rows:
        text = str(row.get("comment_text") or "")
        for token in re.findall(r"[A-Za-z][A-Za-z'-]{2,}", text.casefold()):
            if token not in STOPWORDS:
                terms[token] += 1
    return terms.most_common(limit)
