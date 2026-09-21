from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from .models import CollectionSummary, CommentRecord

EXPORT_FIELDS = [
    "project_id",
    "video_id",
    "video_url",
    "video_title",
    "channel_title",
    "video_published_at",
    "video_view_count",
    "video_like_count",
    "video_comment_count",
    "comment_id",
    "parent_comment_id",
    "is_reply",
    "comment_text",
    "comment_like_count",
    "reply_count",
    "comment_published_at",
    "comment_updated_at",
    "author_display_name",
    "author_channel_id",
    "collected_at",
    "collection_order",
]


def _row(record: CommentRecord) -> dict[str, Any]:
    data = record.export_dict()
    return {field: data.get(field) for field in EXPORT_FIELDS}


def export_csv(records: list[CommentRecord], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=EXPORT_FIELDS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(_row(record) for record in records)
    return path


def export_jsonl(records: list[CommentRecord], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for record in records:
            handle.write(json.dumps(_row(record), ensure_ascii=False) + "\n")
    return path


def export_xlsx(
    records: list[CommentRecord],
    path: Path,
    *,
    generate_report: bool = True,
    profile_id: str = "consumer_product",
    research_goal: str | None = None,
    research_questions: list[str] | None = None,
) -> Path:
    if generate_report:
        from .readable_workbook import create_readable_workbook

        return create_readable_workbook(
            records,
            path,
            profile_id=profile_id,
            research_goal=research_goal,
            research_questions=research_questions,
        )

    return export_raw_xlsx(records, path)


def export_raw_xlsx(records: list[CommentRecord], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "comments"
    sheet.append(EXPORT_FIELDS)
    for record in records:
        row = _row(record)
        sheet.append([row.get(field) for field in EXPORT_FIELDS])

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = {
        "video_url": 36,
        "video_title": 42,
        "comment_text": 80,
        "author_display_name": 24,
        "comment_id": 28,
        "parent_comment_id": 28,
    }
    for index, field in enumerate(EXPORT_FIELDS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(field, max(12, min(len(field) + 2, 24)))
    workbook.save(path)
    return path


def export_summary(summary: CollectionSummary, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(summary.model_dump_json(indent=2), encoding="utf-8")
    return path
