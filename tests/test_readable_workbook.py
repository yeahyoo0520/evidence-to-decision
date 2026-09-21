from pathlib import Path

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from youtube_comment_research.cli import app
from youtube_comment_research.exporters import EXPORT_FIELDS, export_xlsx
from youtube_comment_research.readable_workbook import (
    MAIN_COMMENT_FILL,
    REPLY_FILL,
    SHEET_NAMES,
)
from youtube_comment_research.reporting import load_rows, rows_to_records

FIXTURE_PATH = Path(__file__).parent / "fixtures" / "phase2a_comments.csv"


@pytest.fixture
def fixture_records():
    return rows_to_records(load_rows(FIXTURE_PATH))


@pytest.fixture
def readable_path(tmp_path, fixture_records):
    return export_xlsx(fixture_records, tmp_path / "readable.xlsx")


def _header_map(sheet) -> dict[str, int]:
    return {str(cell.value): cell.column for cell in sheet[1]}


def test_seven_sheets_exist(readable_path) -> None:
    workbook = load_workbook(readable_path)
    assert set(workbook.sheetnames) == set(SHEET_NAMES)


def test_sheets_are_in_required_order(readable_path) -> None:
    workbook = load_workbook(readable_path)
    assert workbook.sheetnames == SHEET_NAMES


def test_main_comment_is_followed_by_its_replies(readable_path) -> None:
    workbook = load_workbook(readable_path)
    sheet = workbook["02_对话视图"]
    headers = _header_map(sheet)
    comment_id = headers["评论 ID comment_id"]
    identifiers = [sheet.cell(row, comment_id).value for row in range(2, sheet.max_row + 1)]
    assert identifiers[:3] == ["top-1", "reply-1", "reply-2"]


def test_reply_keeps_correct_parent_comment_id(readable_path) -> None:
    workbook = load_workbook(readable_path)
    sheet = workbook["02_对话视图"]
    headers = _header_map(sheet)
    parent_id = headers["父评论 ID parent_comment_id"]
    assert sheet.cell(3, parent_id).value == "top-1"
    assert sheet.cell(4, parent_id).value == "top-1"


def test_main_and_reply_rows_have_different_fills(readable_path) -> None:
    workbook = load_workbook(readable_path)
    sheet = workbook["02_对话视图"]
    assert sheet.cell(2, 1).fill.fgColor.rgb == MAIN_COMMENT_FILL.fgColor.rgb
    assert sheet.cell(3, 1).fill.fgColor.rgb == REPLY_FILL.fgColor.rgb
    assert sheet.cell(2, 1).fill.fgColor.rgb != sheet.cell(3, 1).fill.fgColor.rgb


def test_reply_uses_reply_label(readable_path) -> None:
    workbook = load_workbook(readable_path)
    sheet = workbook["02_对话视图"]
    headers = _header_map(sheet)
    type_column = headers["类型 Type"]
    assert sheet.cell(3, type_column).value == "↳ 回复"


def test_reply_target_and_reply_count_display(readable_path) -> None:
    workbook = load_workbook(readable_path)
    sheet = workbook["02_对话视图"]
    headers = _header_map(sheet)
    reply_to = headers["回复对象（主评论） Reply to"]
    reply_count = headers["回复数 Replies"]
    assert sheet.cell(2, reply_count).value == 2
    assert sheet.cell(3, reply_count).value in ("", None)
    assert sheet.cell(3, reply_to).value == "I love my Shokz for running."


def test_raw_data_keeps_every_record_and_field(readable_path, fixture_records) -> None:
    workbook = load_workbook(readable_path)
    sheet = workbook["06_原始数据"]
    headers = [cell.value for cell in sheet[1]]
    assert headers == EXPORT_FIELDS
    assert sheet.max_row - 1 == len(fixture_records)


def test_classification_sheet_has_complete_traceable_fields(readable_path, fixture_records) -> None:
    workbook = load_workbook(readable_path)
    sheet = workbook["03_语义标注"]
    headers = [cell.value for cell in sheet[1]]
    required = {
        "comment_id",
        "primary_category",
        "category_cn",
        "subtopic",
        "category_confidence",
        "category_evidence",
        "manual_review",
    }
    assert required.issubset(headers)
    assert sheet.max_row - 1 == len(fixture_records)
    comment_ids = {sheet.cell(row, 1).value for row in range(2, sheet.max_row + 1)}
    assert comment_ids == {record.comment_id for record in fixture_records}


def test_generated_xlsx_can_be_reopened(readable_path) -> None:
    reopened = load_workbook(readable_path, read_only=True, data_only=True)
    assert reopened["01_项目总览"]["A1"].value == "YouTube 评论研究项目总览"


def test_report_cli_regenerates_workbook_without_api_key(tmp_path, monkeypatch) -> None:
    monkeypatch.delenv("YOUTUBE_API_KEY", raising=False)
    output = tmp_path / "regenerated.xlsx"
    result = CliRunner().invoke(
        app,
        ["report", str(FIXTURE_PATH), "--output", str(output)],
    )
    assert result.exit_code == 0, result.output
    assert load_workbook(output).sheetnames == SHEET_NAMES
