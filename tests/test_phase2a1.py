from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from youtube_comment_research.classification import classify_text
from youtube_comment_research.evaluation import evaluate_gold_labels
from youtube_comment_research.exporters import export_xlsx
from youtube_comment_research.models import CommentRecord
from youtube_comment_research.readable_workbook import SHEET_NAMES
from youtube_comment_research.reporting import build_topic_summary


FIXTURES = Path(__file__).parent / "fixtures"
GOLD_PATH = FIXTURES / "shokz_gold_labels.csv"


def test_openrun_does_not_match_run() -> None:
    result = classify_text("OpenRun")
    assert result.primary_category != "Usage Scenario"
    assert "running" not in result.subtopics


def test_openfit_does_not_match_fit() -> None:
    result = classify_text("OpenFit")
    assert "fit/stability" not in result.subtopics
    assert "comfort" not in result.subtopics


def test_call_does_not_match_inside_another_word() -> None:
    assert "microphone/calls" not in classify_text("Basically excellent.").subtopics


def test_brand_alone_is_not_comparison() -> None:
    assert classify_text("Shokz Bose JBL Sony").primary_category != "Competitor Comparison"


def test_explicit_comparison_phrase_is_comparison() -> None:
    result = classify_text("I prefer Shokz to Bose.")
    assert result.primary_category == "Competitor Comparison"
    assert "compare.prefer_to" in result.category_evidence


def test_enjoying_is_product_praise() -> None:
    assert classify_text("Enjoying my OpenRun Pro").primary_category == "Product Praise"


def test_talk_question_is_product_question() -> None:
    result = classify_text("Can you talk on them")
    assert result.primary_category == "Product Question"
    assert "microphone/calls" in result.subtopics


def test_earplug_compatibility_question() -> None:
    result = classify_text("Does these work with earplugs?")
    assert result.primary_category == "Product Question"
    assert "earplug compatibility" in result.subtopics


def test_hearing_aid_compatibility_question() -> None:
    result = classify_text("Compatible with ITC hearing aids?")
    assert result.primary_category == "Product Question"
    assert "hearing aid compatibility" in result.subtopics


def test_olive_green_is_feature_request() -> None:
    result = classify_text("Can we have an olive green colour?")
    assert result.primary_category == "Feature Request"
    assert "colour/design" in result.subtopics


def test_waiting_for_us_launch_lifecycle() -> None:
    result = classify_text("Hoping OpenFit 2+ comes to USA soon.")
    assert result.lifecycle_stage == "Waiting for Launch"
    assert result.primary_category == "Purchase Intent"


def test_purchased_does_not_automatically_mean_purchase_intent() -> None:
    result = classify_text("I bought OpenRun Pro 2 and I am satisfied.")
    assert result.lifecycle_stage == "Purchased"
    assert result.primary_category == "Product Praise"
    assert result.primary_category != "Purchase Intent"


def test_returned_competitor_is_not_shokz_after_sales() -> None:
    result = classify_text("I returned my AirPods and bought Shokz.")
    assert result.primary_category != "After-sales Issue"
    assert result.lifecycle_stage == "Purchased"
    assert result.manual_review is True


def test_multitopic_comment_keeps_secondary_category() -> None:
    result = classify_text("I love these, but the plastic broke and I am disappointed.")
    assert result.secondary_categories
    assert {"Product Praise", "Product Concern"}.issubset(
        {result.primary_category, *result.secondary_categories}
    )


def test_multitopic_comment_requires_manual_review() -> None:
    result = classify_text("I love these, but the plastic broke and I am disappointed.")
    assert result.manual_review is True
    assert "Multiple substantive categories" in result.review_reason


def test_single_weak_keyword_confidence_is_at_most_point_six() -> None:
    result = classify_text("rain")
    assert result.category_confidence <= 0.60
    assert result.manual_review is True
    assert result.review_reason


def test_high_confidence_result_contains_strong_evidence() -> None:
    result = classify_text("Can you talk on them")
    assert result.category_confidence >= 0.80
    matches = json.loads(result.as_dict()["matched_rules"])
    assert any(
        match["category"] == result.primary_category and match["strength"] == "strong"
        for match in matches
    )


def test_matched_rules_include_required_trace_fields() -> None:
    match = json.loads(classify_text("Shokz is better than Bose").as_dict()["matched_rules"])[0]
    assert {
        "matched_phrase",
        "matched_span",
        "rule_id",
        "category",
        "score",
        "target_entity",
    }.issubset(match)


def test_gold_label_evaluation_runs() -> None:
    result = evaluate_gold_labels(GOLD_PATH)
    assert result.evaluated_count == 30
    assert 0 <= result.primary_category_accuracy <= 1
    assert 0 <= result.manual_review_recall <= 1
    assert isinstance(result.per_category_confusion_summary, dict)


def test_topic_summary_separates_manual_review_rows() -> None:
    rows = [
        {
            "primary_category": "Product Praise",
            "category_cn": "产品认可",
            "subtopics": "comfort",
            "manual_review": False,
            "comment_text": "works great",
            "comment_id": "confirmed",
            "comment_like_count": 2,
        },
        {
            "primary_category": "Product Praise",
            "category_cn": "产品认可",
            "subtopics": "comfort",
            "manual_review": True,
            "comment_text": "interesting",
            "comment_id": "review",
            "comment_like_count": 3,
        },
    ]
    summary = build_topic_summary(rows)
    praise_rows = [row for row in summary if row["primary_category"] == "Product Praise"]
    assert len(praise_rows) == 2
    assert {row["is_primary_summary"] for row in praise_rows} == {True, False}
    assert sum(row["count"] for row in praise_rows) == 2


def test_long_comment_rows_are_taller_and_text_is_preserved(tmp_path: Path) -> None:
    base = {
        "project_id": "phase2a1-fixture",
        "video_id": "video-one",
        "video_url": "https://www.youtube.com/watch?v=video-one",
        "video_title": "Fixture video",
        "channel_title": "Fixture channel",
        "comment_id": "short",
        "is_reply": False,
        "comment_text": "Works great.",
        "comment_like_count": 0,
        "reply_count": 0,
        "author_display_name": "Commenter 001",
        "collected_at": "2026-07-26T00:00:00Z",
        "collection_order": 1,
    }
    long_text = (
        "这是一条包含中文、English words 和 emoji 😊 的长评论。\n"
        + "It discusses comfort, battery, charging, sound quality and daily use. " * 12
    )
    long_row = {**base, "comment_id": "long", "comment_text": long_text, "collection_order": 2}
    records = [CommentRecord.model_validate(base), CommentRecord.model_validate(long_row)]
    path = export_xlsx(records, tmp_path / "long-comments.xlsx")
    workbook = load_workbook(path)

    conversation = workbook[SHEET_NAMES[1]]
    assert conversation.row_dimensions[3].height > conversation.row_dimensions[2].height
    assert conversation.cell(3, 3).value == long_text
    assert conversation.row_dimensions[3].height <= 150

    classification = workbook[SHEET_NAMES[2]]
    headers = {cell.value: cell.column for cell in classification[1]}
    assert classification.row_dimensions[3].height > classification.row_dimensions[2].height
    assert classification.cell(3, headers["comment_text"]).value == long_text
    assert classification.row_dimensions[3].height <= 150


def test_phase2a1_workbook_keeps_seven_sheet_contract(tmp_path: Path) -> None:
    record = CommentRecord.model_validate(
        {
            "project_id": "phase2a1-fixture",
            "video_id": "video-one",
            "video_url": "https://www.youtube.com/watch?v=video-one",
            "video_title": "Fixture video",
            "channel_title": "Fixture channel",
            "comment_id": "comment-one",
            "is_reply": False,
            "comment_text": "Enjoying my OpenRun Pro",
            "comment_like_count": 1,
            "reply_count": 0,
            "author_display_name": "Commenter 001",
            "collected_at": "2026-07-26T00:00:00Z",
            "collection_order": 1,
        }
    )
    path = export_xlsx([record], tmp_path / "seven-sheets.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == SHEET_NAMES
    headers = {cell.value for cell in workbook[SHEET_NAMES[2]][1]}
    assert {
        "primary_category",
        "primary_category_cn",
        "secondary_categories",
        "secondary_categories_cn",
        "subtopics",
        "lifecycle_stage",
        "mentioned_brands",
        "mentioned_products",
        "category_confidence",
        "category_evidence",
        "matched_rules",
        "manual_review",
        "review_reason",
    }.issubset(headers)
