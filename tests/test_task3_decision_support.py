from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from youtube_comment_research.cli import app
from youtube_comment_research.decision_provider import CodexFileDecisionSupportProvider
from youtube_comment_research.decision_support import (
    DecisionBriefRecord,
    build_decision_briefs,
    load_decision_briefs,
    write_decision_briefs,
)
from youtube_comment_research.evidence import ResearchContext
from youtube_comment_research.errors import ClassificationWorkflowError
from youtube_comment_research.readable_workbook import create_readable_workbook
from youtube_comment_research.reporting import load_rows
from youtube_comment_research.verification import EffectiveInsightRecord


FIXTURE = Path(__file__).parent / "fixtures" / "general_feedback_evidence.json"
FIXED_TIME = "2026-09-11T00:00:00+00:00"


def _effective(
    suffix: str = "1",
    *,
    status: str = "unreviewed",
    eligible: bool = True,
    review_recommended: bool = False,
    context: ResearchContext | None = None,
) -> EffectiveInsightRecord:
    return EffectiveInsightRecord(
        effective_insight_id=f"effective-{suffix}",
        source_insight_id=f"insight-{suffix}",
        source_theme_id=f"theme-{suffix}",
        theme_name=f"Adoption pattern {suffix}",
        theme_description="Evidence describes friction during initial setup.",
        final_observation=(
            "Human-edited evidence shows setup instructions are inconsistent."
            if status == "edited"
            else "Several records describe friction during initial setup."
        ),
        final_interpretation="The pattern may affect early adoption.",
        evidence_ids=[f"gf-00{suffix}"],
        representative_evidence_ids=[f"gf-00{suffix}"],
        verification_status=status,
        reviewer_note="Human wording retained." if status == "edited" else "",
        rejection_reason="Rejected." if status == "rejected" else "",
        review_recommended=review_recommended,
        review_reasons=["Only 1 supporting evidence item"] if review_recommended else [],
        downstream_eligible=eligible,
        confidence="medium",
        taxonomy_coverage="covered",
        coverage_reason="The taxonomy expresses setup friction.",
        research_context=context or ResearchContext(),
        provenance="human_edited" if status == "edited" else "ai_generated",
    )


def _proposal(ids: list[str]) -> dict[str, object]:
    return {
        "source_effective_insight_ids": ids,
        "why_it_matters": "If this pattern holds, setup friction may slow early adoption.",
        "decision_question": "Should setup guidance be clarified before the next iteration?",
        "possible_directions": [
            {
                "direction": "Consider clarifying the setup sequence.",
                "rationale": "The source evidence describes setup friction.",
                "evidence_basis": "The cited records mention inconsistent setup guidance.",
                "tradeoff_or_risk": "More detail may increase information density and needs validation.",
            }
        ],
        "evidence_gaps": [
            {
                "gap": "No behavioral outcome data is available.",
                "why_it_matters": "Feedback alone cannot show whether setup friction changes adoption.",
                "related_insight_ids": ids,
                "related_evidence_ids": [],
            }
        ],
        "validation_plan": [
            {
                "validation_question": "Does clearer guidance reduce repeated setup friction?",
                "method": "Follow-up observation",
                "required_evidence": "Feedback and completion observations after a message change.",
                "evaluation_criterion": "Whether repeated setup concerns decline across additional sources.",
                "limitation": "A short observation period may miss longer-term effects.",
            }
        ],
        "uncertainty_note": "This brief uses qualitative feedback and does not establish adoption impact.",
    }


class FakeProvider:
    provider_id = "fake"

    def __init__(self, proposals: list[dict[str, object]]) -> None:
        self.proposals = proposals
        self.seen: list[dict[str, object]] = []

    def generate_decision_briefs(self, effective_insights):
        self.seen = effective_insights
        return {"briefs": self.proposals}


@pytest.mark.parametrize("status", ["unreviewed", "reviewed", "edited"])
def test_eligible_verification_states_generate_brief_and_preserve_content(status: str) -> None:
    source = _effective(status=status)
    provider = FakeProvider([_proposal([source.effective_insight_id])])
    brief = build_decision_briefs([source], provider, generated_at=FIXED_TIME)[0]
    assert brief.verification_statuses == [status]
    assert provider.seen[0]["final_observation"] == source.final_observation
    if status == "edited":
        assert "Human-edited" in str(provider.seen[0]["final_observation"])


@pytest.mark.parametrize(
    "source",
    [_effective(status="rejected", eligible=False), _effective(eligible=False)],
)
def test_rejected_or_ineligible_insight_does_not_enter_decision_support(source) -> None:
    provider = FakeProvider([])
    assert build_decision_briefs([source], provider, generated_at=FIXED_TIME) == []
    assert provider.seen == []


def test_review_recommendation_warns_but_does_not_block() -> None:
    source = _effective(review_recommended=True)
    brief = build_decision_briefs(
        [source], FakeProvider([_proposal([source.effective_insight_id])]), generated_at=FIXED_TIME
    )[0]
    assert brief.review_recommended
    assert brief.review_reasons == ["Only 1 supporting evidence item"]


def test_supporting_evidence_is_computed_from_source_insights() -> None:
    sources = [_effective("1"), _effective("2")]
    brief = build_decision_briefs(
        sources,
        FakeProvider([_proposal([item.effective_insight_id for item in sources])]),
        generated_at=FIXED_TIME,
    )[0]
    assert brief.supporting_evidence_ids == ["gf-001", "gf-002"]
    assert brief.source_theme_ids == ["theme-1", "theme-2"]


def test_supporting_evidence_scope_is_not_automatically_expanded_to_dataset() -> None:
    source = _effective("1")
    all_dataset_ids = {"gf-001", "gf-002", "gf-003", "gf-004"}
    brief = build_decision_briefs(
        [source],
        FakeProvider([_proposal([source.effective_insight_id])]),
        generated_at=FIXED_TIME,
    )[0]
    assert set(brief.supporting_evidence_ids) == set(source.evidence_ids)
    assert set(brief.supporting_evidence_ids) != all_dataset_ids


def test_unknown_insight_and_out_of_scope_gap_evidence_are_rejected() -> None:
    source = _effective()
    with pytest.raises(ClassificationWorkflowError, match="unavailable Effective"):
        build_decision_briefs([source], FakeProvider([_proposal(["missing"])]))
    proposal = _proposal([source.effective_insight_id])
    proposal["evidence_gaps"][0]["related_evidence_ids"] = ["invented"]
    with pytest.raises(ClassificationWorkflowError, match="outside its source"):
        build_decision_briefs([source], FakeProvider([proposal]))


def test_required_decision_sections_are_nonempty_and_structured() -> None:
    source = _effective()
    brief = build_decision_briefs(
        [source], FakeProvider([_proposal([source.effective_insight_id])]), generated_at=FIXED_TIME
    )[0]
    assert brief.why_it_matters
    assert brief.decision_question.endswith("?")
    assert brief.decision_question != source.final_interpretation
    assert brief.possible_directions[0].tradeoff_or_risk
    assert brief.evidence_gaps[0].gap
    assert brief.validation_plan[0].evaluation_criterion
    assert brief.uncertainty_note
    assert brief.status == "AI-generated decision support"


def test_evidence_gaps_may_be_empty_but_validation_plan_remains_required() -> None:
    source = _effective()
    proposal = _proposal([source.effective_insight_id])
    proposal["evidence_gaps"] = []
    brief = build_decision_briefs([source], FakeProvider([proposal]))[0]
    assert brief.evidence_gaps == []
    assert brief.validation_plan


def test_simple_insight_restatement_is_rejected_as_decision_question() -> None:
    source = _effective()
    proposal = _proposal([source.effective_insight_id])
    proposal["decision_question"] = source.final_interpretation + "?"
    with pytest.raises(ClassificationWorkflowError, match="not repeat"):
        build_decision_briefs([source], FakeProvider([proposal]))


@pytest.mark.parametrize(
    ("context", "expected_mode"),
    [
        (ResearchContext(), "exploratory"),
        (
            ResearchContext(
                research_question="What affects adoption?",
                decision_context="Prioritize the next onboarding iteration.",
                analysis_mode="decision_focused",
            ),
            "decision_focused",
        ),
    ],
)
def test_research_context_modes_enter_brief(context: ResearchContext, expected_mode: str) -> None:
    source = _effective(context=context)
    brief = build_decision_briefs(
        [source], FakeProvider([_proposal([source.effective_insight_id])]), generated_at=FIXED_TIME
    )[0]
    assert brief.analysis_mode == expected_mode
    assert brief.research_question == context.research_question
    assert brief.decision_context == context.decision_context


def test_different_contexts_cannot_be_forced_into_one_brief() -> None:
    sources = [
        _effective("1", context=ResearchContext(research_question="Question A")),
        _effective("2", context=ResearchContext(research_question="Question B")),
    ]
    with pytest.raises(ClassificationWorkflowError, match="different Research Contexts"):
        build_decision_briefs(
            sources,
            FakeProvider([_proposal([item.effective_insight_id for item in sources])]),
        )


def test_provider_can_keep_unrelated_insights_as_separate_briefs() -> None:
    sources = [_effective("1"), _effective("2")]
    proposals = [_proposal([item.effective_insight_id]) for item in sources]
    briefs = build_decision_briefs(sources, FakeProvider(proposals), generated_at=FIXED_TIME)
    assert len(briefs) == 2
    assert all(len(item.source_effective_insight_ids) == 1 for item in briefs)


def test_json_round_trip_and_file_loader(tmp_path: Path) -> None:
    source = _effective()
    records = build_decision_briefs(
        [source], FakeProvider([_proposal([source.effective_insight_id])]), generated_at=FIXED_TIME
    )
    output = write_decision_briefs(tmp_path / "briefs.json", records)
    loaded = load_decision_briefs(output)
    assert loaded == records
    assert DecisionBriefRecord.model_validate_json(records[0].model_dump_json()) == records[0]


def test_file_provider_prepares_reviewable_request_and_consumes_response(tmp_path: Path) -> None:
    source = _effective()
    provider = CodexFileDecisionSupportProvider(tmp_path)
    with pytest.raises(ClassificationWorkflowError, match="response is pending"):
        build_decision_briefs([source], provider)
    request = json.loads((tmp_path / "decision_request.json").read_text(encoding="utf-8"))
    assert request["effective_insights"][0]["effective_insight_id"] == source.effective_insight_id
    (tmp_path / "decision_response.json").write_text(
        json.dumps({"briefs": [_proposal([source.effective_insight_id])]}), encoding="utf-8"
    )
    assert len(build_decision_briefs([source], provider, generated_at=FIXED_TIME)) == 1


def test_cli_help_exposes_decision_support_contract() -> None:
    result = CliRunner().invoke(app, ["generate-decision-brief", "--help"])
    assert result.exit_code == 0
    assert "--human-review" in result.output
    assert "--decision-provider" in result.output
    assert "--work-dir" in result.output


def test_cli_consumes_file_response_and_writes_valid_output(tmp_path: Path) -> None:
    source = _effective()
    input_path = tmp_path / "effective.jsonl"
    input_path.write_text(source.model_dump_json() + "\n", encoding="utf-8")
    work_dir = tmp_path / "decision-work"
    work_dir.mkdir()
    (work_dir / "decision_response.json").write_text(
        json.dumps({"briefs": [_proposal([source.effective_insight_id])]}),
        encoding="utf-8",
    )
    output = tmp_path / "briefs.json"
    result = CliRunner().invoke(
        app,
        [
            "generate-decision-brief",
            str(input_path),
            "--output",
            str(output),
            "--work-dir",
            str(work_dir),
        ],
    )
    assert result.exit_code == 0, result.output
    assert len(load_decision_briefs(output)) == 1
    assert "AI-generated decision support" in result.output


def test_static_schema_covers_every_record_field() -> None:
    schema_path = Path(__file__).parents[1] / "schemas" / "decision_brief.schema.json"
    schema = json.loads(schema_path.read_text(encoding="utf-8"))
    assert set(schema["required"]) == set(DecisionBriefRecord.model_fields)
    assert schema["properties"]["status"]["const"] == "AI-generated decision support"


def test_workbook_has_readable_eleventh_sheet_and_reopens(tmp_path: Path) -> None:
    source = _effective()
    briefs = build_decision_briefs(
        [source], FakeProvider([_proposal([source.effective_insight_id])]), generated_at=FIXED_TIME
    )
    output = create_readable_workbook(
        load_rows(FIXTURE),
        tmp_path / "decision-support.xlsx",
        profile_id="general_feedback",
        decision_briefs=briefs,
    )
    workbook = load_workbook(output)
    assert workbook.sheetnames[-1] == "11_决策支持"
    sheet = workbook["11_决策支持"]
    headers = [cell.value for cell in sheet[1]]
    assert {"Why It Matters", "Decision Question", "Validation Plan", "Uncertainty Note"} <= set(headers)
    assert sheet["G2"].value == briefs[0].why_it_matters
    assert sheet.row_dimensions[2].height >= 30
