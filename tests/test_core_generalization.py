from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from typer._click.utils import strip_ansi
from typer.testing import CliRunner

from youtube_comment_research.cli import app
from youtube_comment_research.evidence import (
    EvidenceRecord,
    ResearchContext,
    evidence_to_research_rows,
    load_evidence_json,
)
from youtube_comment_research.decision_support import build_decision_briefs
from youtube_comment_research.human_review import prepare_review_packet
from youtube_comment_research.models import CommentRecord
from youtube_comment_research.profiles import load_profile
from youtube_comment_research.readable_workbook import (
    GENERIC_RAW_FIELDS,
    GENERIC_SHEET_NAMES,
    SHEET_NAMES,
    create_readable_workbook,
)
from youtube_comment_research.reporting import load_rows
from youtube_comment_research.semantic_workflow import (
    merge_classifications,
    prepare_classification_batches,
)
from youtube_comment_research.theme_discovery import generate_theme_insights
from youtube_comment_research.verification import resolve_effective_insights


FIXTURE = Path(__file__).parent / "fixtures" / "general_feedback_evidence.json"


class GenericThemeProvider:
    provider_id = "generic-test-provider"

    def __init__(self) -> None:
        self.contexts: list[ResearchContext] = []
        self.evidence_ids: list[str] = []

    def discover_candidate_themes(
        self,
        comments: list[dict[str, Any]],
        *,
        profile: object,
        research_context: ResearchContext,
    ) -> object:
        self.contexts.append(research_context)
        self.evidence_ids = [item["evidence_id"] for item in comments]
        return {
            "themes": [
                {
                    "theme_name": "Adoption friction and information gaps",
                    "description": "Evidence combines adoption barriers, questions, and service friction.",
                    "evidence_ids": self.evidence_ids,
                    "representative_evidence_ids": ["gf-003", "gf-008"],
                }
            ]
        }

    def consolidate_themes(
        self,
        candidates: list[dict[str, Any]],
        comments: list[dict[str, Any]],
        *,
        profile: object,
        research_context: ResearchContext,
    ) -> object:
        self.contexts.append(research_context)
        return {
            "themes": [
                {
                    "theme_name": "Adoption friction and information gaps",
                    "description": "Evidence combines adoption barriers, questions, and service friction.",
                    "source_candidate_ids": [candidates[0]["candidate_id"]],
                    "representative_evidence_ids": ["gf-003", "gf-008"],
                    "taxonomy_coverage": "covered",
                    "coverage_reason": "The general profile explicitly covers barriers, questions, and service issues.",
                }
            ]
        }


class GenericDecisionProvider:
    provider_id = "generic-decision-test-provider"

    def generate_decision_briefs(self, effective_insights):
        return {
            "briefs": [
                {
                    "source_effective_insight_ids": [item["effective_insight_id"]],
                    "why_it_matters": "If the pattern holds, adoption friction may delay successful use.",
                    "decision_question": "Which adoption friction should be validated before the next onboarding iteration?",
                    "possible_directions": [
                        {
                            "direction": "Consider clarifying the highest-friction onboarding information.",
                            "rationale": "The evidence contains repeated questions and barriers.",
                            "evidence_basis": "The source Insight cites setup, access, and approval friction.",
                            "tradeoff_or_risk": "Additional guidance may add complexity and should be tested.",
                        }
                    ],
                    "evidence_gaps": [
                        {
                            "gap": "The qualitative evidence does not measure completion outcomes.",
                            "why_it_matters": "The decision question concerns successful adoption, not mentions alone.",
                            "related_insight_ids": [item["effective_insight_id"]],
                            "related_evidence_ids": [],
                        }
                    ],
                    "validation_plan": [
                        {
                            "validation_question": "Does the same friction recur across additional sources?",
                            "method": "Collect additional feedback and inspect onboarding behavior.",
                            "required_evidence": "Additional feedback plus completion observations.",
                            "evaluation_criterion": "Whether the concern recurs and aligns with incomplete onboarding.",
                            "limitation": "Observed association would not by itself prove causation.",
                        }
                    ],
                    "uncertainty_note": "The current evidence is qualitative and does not establish outcome impact.",
                }
                for item in effective_insights
            ]
        }

def _classification(item: dict[str, Any], category: str) -> dict[str, Any]:
    profile = load_profile("general_feedback")
    return {
        "comment_id": item["comment_id"],
        "subject": "offering",
        "primary_category": category,
        "secondary_categories": [],
        "subtopics": ["workflow"],
        "lifecycle_stage": "Using",
        "mentioned_brands": [],
        "mentioned_products": [],
        "confidence": 0.82,
        "evidence": [item["comment_text"]],
        "review_risk": "low",
        "review_triggers": [],
        "manual_review": False,
        "review_reason": "",
        "classification_source": "codex-assisted",
        "classifier_version": "generic-fixture-provider-v1",
        "rubric_version": profile.rubric_version,
    }


def test_evidence_record_supports_requested_sources_and_timestamp_alias() -> None:
    record = EvidenceRecord.model_validate(
        {
            "evidence_id": "source-1",
            "text": "A source-neutral statement.",
            "source_type": "customer_inquiry",
            "timestamp": "2026-09-01T00:00:00Z",
            "language": "en",
            "market": "US",
            "region": "North America",
            "metadata": {"case_id": "case-1"},
        }
    )
    assert record.created_at is not None
    assert record.source_type == "customer_inquiry"


def test_generic_fixture_has_source_neutral_records_and_unique_ids() -> None:
    records = load_evidence_json(FIXTURE)
    assert 12 <= len(records) <= 20
    assert len({record.evidence_id for record in records}) == len(records)
    assert {record.source_type for record in records} >= {
        "product_review",
        "customer_inquiry",
        "crm_note",
        "survey_response",
        "interview_note",
        "social_feedback",
    }


def test_evidence_adapter_uses_one_identity_chain() -> None:
    rows = evidence_to_research_rows(load_evidence_json(FIXTURE))
    assert all(row["evidence_id"] == row["comment_id"] for row in rows)
    assert all(row["text"] == row["comment_text"] for row in rows)


def test_general_feedback_profile_is_generic_and_game_profile_remains() -> None:
    general = load_profile("general_feedback")
    game = load_profile("game_publishing")
    assert "Need / Request" in general.primary_categories
    assert "Other / Unclear" in general.primary_categories
    assert "Game / Content Praise" in game.primary_categories
    assert "game" not in {value.casefold() for value in general.subject_options}


@pytest.mark.parametrize("analysis_mode", ["exploratory", "decision_focused"])
def test_research_context_supports_both_modes(analysis_mode: str) -> None:
    context = ResearchContext(
        research_question="What prevents successful adoption?",
        decision_context="Choose the next onboarding improvement.",
        analysis_mode=analysis_mode,
    )
    assert context.analysis_mode == analysis_mode


def test_generic_evidence_runs_existing_classification_theme_review_and_workbook(
    tmp_path: Path,
) -> None:
    batch_dir = tmp_path / "classification"
    prepared = prepare_classification_batches(
        FIXTURE,
        batch_dir,
        profile_id="general_feedback",
    )
    categories = list(load_profile("general_feedback").primary_categories)
    for batch in prepared.batch_files:
        items = [json.loads(line) for line in batch.read_text(encoding="utf-8").splitlines()]
        output = batch_dir / f"classified_{batch.name}"
        output.write_text(
            "".join(
                json.dumps(_classification(item, categories[index % len(categories)])) + "\n"
                for index, item in enumerate(items)
            ),
            encoding="utf-8",
        )

    merged = tmp_path / "merged.csv"
    result = merge_classifications(
        FIXTURE,
        batch_dir,
        merged,
        profile_id="general_feedback",
    )
    assert result.classified_records == 15
    merged_rows = load_rows(merged)
    assert all(row["evidence_id"] == row["comment_id"] for row in merged_rows)

    context = ResearchContext(
        research_question="What prevents successful adoption?",
        decision_context="Choose the next onboarding improvement.",
        analysis_mode="decision_focused",
    )
    provider = GenericThemeProvider()
    bundle_path = tmp_path / "theme_insights.json"
    bundle = generate_theme_insights(
        merged,
        bundle_path,
        profile_id="general_feedback",
        theme_method="llm",
        provider=provider,
        research_context=context,
    )
    assert bundle.research_context == context
    assert set(bundle.themes[0].evidence_ids) == {
        row["evidence_id"] for row in merged_rows
    }
    assert len(provider.evidence_ids) == 15
    assert provider.contexts == [context, context]

    packet_path = tmp_path / "human_review_packet.json"
    packet = prepare_review_packet(bundle_path, packet_path)
    assert len(packet.items) == 1
    assert packet.items[0].insight_draft.evidence_ids == bundle.themes[0].evidence_ids

    unreviewed = resolve_effective_insights(bundle)[0]
    edited = unreviewed.model_copy(
        update={
            "effective_insight_id": "effective-edited-demo",
            "source_insight_id": "insight-edited-demo",
            "source_theme_id": "theme-edited-demo",
            "final_observation": "Human-edited evidence emphasizes approval and setup friction.",
            "verification_status": "edited",
            "provenance": "human_edited",
        }
    )
    rejected = unreviewed.model_copy(
        update={
            "effective_insight_id": "effective-rejected-demo",
            "source_insight_id": "insight-rejected-demo",
            "source_theme_id": "theme-rejected-demo",
            "verification_status": "rejected",
            "downstream_eligible": False,
            "rejection_reason": "Not supported for downstream use.",
        }
    )
    decision_briefs = build_decision_briefs(
        [unreviewed, edited, rejected], GenericDecisionProvider()
    )
    assert len(decision_briefs) == 2
    assert {brief.verification_statuses[0] for brief in decision_briefs} == {
        "unreviewed",
        "edited",
    }
    assert all(
        "effective-rejected-demo" not in brief.source_effective_insight_ids
        for brief in decision_briefs
    )

    workbook_path = create_readable_workbook(
        merged_rows,
        tmp_path / "generic_research.xlsx",
        profile_id="general_feedback",
        theme_bundle=bundle,
        decision_briefs=decision_briefs,
    )
    workbook = load_workbook(workbook_path)
    assert workbook.sheetnames == GENERIC_SHEET_NAMES
    assert "04_分类概览" in workbook.sheetnames
    assert "04_主题洞察" not in workbook.sheetnames
    assert workbook["01_项目总览"]["A1"].value == "研究证据项目总览 Research Evidence Overview"
    raw_headers = [cell.value for cell in workbook["06_原始数据"][1]]
    assert {"evidence_id", "text", "source_type", "source_name"} <= set(raw_headers)
    assert workbook["02_对话视图"]["C1"].value == "证据正文 Evidence Text"
    assert workbook["11_决策支持"]["N2"].value == "AI-generated decision support"


def test_generic_workbook_is_evidence_first_without_youtube_or_mandatory_gate(
    tmp_path: Path,
) -> None:
    rows = load_rows(FIXTURE)
    output = create_readable_workbook(
        rows,
        tmp_path / "generic-presentation.xlsx",
        profile_id="general_feedback",
        research_context=ResearchContext(
            research_question="What affects adoption?",
            decision_context="Choose what to validate next.",
            analysis_mode="decision_focused",
        ),
    )
    workbook = load_workbook(output)
    assert workbook.sheetnames == GENERIC_SHEET_NAMES
    overview_labels = {workbook["01_项目总览"].cell(row, 1).value for row in range(1, workbook["01_项目总览"].max_row + 1)}
    assert {"Evidence Count", "Source Count", "Source Types", "Research Question", "Decision Context"} <= overview_labels
    assert not {"视频数量", "顶级评论数量", "回复数量"} & overview_labels
    source_headers = [cell.value for cell in workbook["05_来源对比"][1]]
    assert source_headers[:3] == ["Source Name", "Source Type", "Evidence Count"]
    raw_headers = [cell.value for cell in workbook["06_原始数据"][1]]
    assert raw_headers == list(GENERIC_RAW_FIELDS)
    assert not {"video_id", "video_title", "channel_title", "comment_id"} & set(raw_headers)
    visible_text = "\n".join(
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert "Mandatory Manual Review" not in visible_text
    assert "Needs Human Confirmation" not in visible_text
    assert "No optional Human Review bundle supplied" not in visible_text
    assert "legacy" not in visible_text.casefold()
    assert workbook["09_洞察草稿"]["D1"].value == "Validation Cue"


def test_youtube_workbook_keeps_video_presentation(tmp_path: Path) -> None:
    youtube_fixture = Path(__file__).parent / "fixtures" / "phase2a_comments.csv"
    output = create_readable_workbook(
        load_rows(youtube_fixture),
        tmp_path / "youtube-presentation.xlsx",
    )
    workbook = load_workbook(output)
    assert workbook.sheetnames == SHEET_NAMES
    assert "04_主题洞察" in workbook.sheetnames
    assert "05_视频对比" in workbook.sheetnames
    overview_labels = {workbook["01_项目总览"].cell(row, 1).value for row in range(1, workbook["01_项目总览"].max_row + 1)}
    assert {"视频数量", "顶级评论数量", "回复数量"} <= overview_labels
    raw_headers = [cell.value for cell in workbook["06_原始数据"][1]]
    assert {"video_id", "video_title", "comment_id"} <= set(raw_headers)


def test_cli_exposes_generic_input_and_research_context_options() -> None:
    prepare = CliRunner().invoke(app, ["prepare-classification", "--help"])
    insights = CliRunner().invoke(app, ["generate-insights", "--help"])
    prepare_help = strip_ansi(prepare.output)
    insights_help = strip_ansi(insights.output)
    assert prepare.exit_code == 0
    assert "Generic Evidence JSON" in prepare_help
    assert insights.exit_code == 0
    assert "--research-question" in insights_help
    assert "--decision-context" in insights_help
    assert "--analysis-mode" in insights_help


def test_youtube_comment_model_remains_backward_compatible() -> None:
    record = CommentRecord(
        project_id="legacy",
        video_id="video",
        video_url="https://www.youtube.com/watch?v=abcdefghijk",
        video_title="Legacy video",
        comment_id="comment-1",
        comment_text="Legacy comment",
        collection_order=1,
    )
    assert record.comment_id == "comment-1"
    assert record.comment_text == "Legacy comment"
