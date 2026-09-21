from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from youtube_comment_research.cli import app
from youtube_comment_research.evidence import ResearchContext
from youtube_comment_research.errors import ClassificationWorkflowError
from youtube_comment_research.human_review import (
    HumanReviewDecision,
    HumanReviewResponse,
    apply_review_response,
    build_review_packet,
)
from youtube_comment_research.readable_workbook import create_readable_workbook
from youtube_comment_research.reporting import load_rows
from youtube_comment_research.theme_discovery import (
    InsightDraft,
    ThemeInsightBundle,
    ThemeRecord,
)
from youtube_comment_research.verification import (
    inspect_effective_insight_evidence,
    resolve_effective_insight_file,
    resolve_effective_insights,
)


FIXTURE = Path(__file__).parent / "fixtures" / "general_feedback_evidence.json"
CHECKSUM = "a" * 64
FIXED_TIME = "2026-09-11T00:00:00+00:00"


def _bundle(
    *,
    evidence_count: int = 3,
    confidence: str = "high",
    coverage: str = "covered",
    representatives: bool = True,
) -> ThemeInsightBundle:
    ids = [f"gf-{index:03d}" for index in range(1, evidence_count + 1)]
    representative_ids = ids[:1] if representatives else []
    representative_comments = (
        ["The setup guide was clear and I was running in ten minutes."]
        if representatives
        else []
    )
    theme = ThemeRecord(
        theme_id="theme-general-feedback",
        theme_name="Adoption experience",
        theme_description="Evidence describes adoption experience.",
        related_categories=["Positive Signal / Praise"],
        evidence_count=evidence_count,
        evidence_comment_ids=ids,
        representative_comments=representative_comments,
        representative_evidence_ids=representative_ids,
        discovery_method="llm_inductive" if coverage == "emergent" else "taxonomy_based",
        taxonomy_coverage=coverage,
        coverage_reason="Coverage was assessed against the general feedback taxonomy.",
        confidence=confidence,
        review_status="preliminary",
    )
    insight = InsightDraft(
        insight_id="insight-general-feedback",
        theme_id=theme.theme_id,
        observation="Several evidence records describe the adoption experience.",
        interpretation="This pattern may indicate that onboarding affects adoption.",
        potential_opportunity="Review the evidence before using the interpretation.",
        evidence_ids=ids,
        confidence=confidence,
        review_status="preliminary",
        draft_label="AI generated draft",
    )
    return ThemeInsightBundle(
        profile_id="general_feedback",
        generated_at=FIXED_TIME,
        source_comment_count=evidence_count,
        research_context=ResearchContext(
            research_question="What affects adoption?",
            decision_context="Prepare the next onboarding iteration.",
            analysis_mode="decision_focused",
        ),
        themes=[theme],
        insight_drafts=[insight],
    )


def _reviewed_bundle(action: str):
    bundle = _bundle()
    packet = build_review_packet(
        bundle,
        source_bundle_sha256=CHECKSUM,
        prepared_at=FIXED_TIME,
    )
    kwargs = {}
    if action == "edit":
        kwargs["observation_override"] = "Human-reviewed adoption evidence is mixed."
    if action == "reject":
        kwargs["rejection_reason"] = "The evidence does not support this interpretation."
    response = HumanReviewResponse(
        source_bundle_sha256=CHECKSUM,
        decisions=[
            HumanReviewDecision(
                review_id=packet.items[0].review_id,
                action=action,
                reviewer_note="Checked against source evidence.",
                reviewed_at=FIXED_TIME,
                **kwargs,
            )
        ],
    )
    return apply_review_response(
        bundle,
        response,
        source_bundle_sha256=CHECKSUM,
        generated_at=FIXED_TIME,
    )


def test_unreviewed_is_downstream_eligible_by_default() -> None:
    result = resolve_effective_insights(_bundle())
    assert len(result) == 1
    assert result[0].verification_status == "unreviewed"
    assert result[0].downstream_eligible
    assert result[0].provenance == "ai_generated"


def test_confirm_maps_to_reviewed_and_preserves_ai_content_and_note() -> None:
    review_bundle = _reviewed_bundle("confirm")
    result = resolve_effective_insights(review_bundle)[0]
    original = review_bundle.source_bundle.insight_drafts[0]
    assert result.verification_status == "reviewed"
    assert result.final_observation == original.observation
    assert result.final_interpretation == original.interpretation
    assert result.reviewer_note == "Checked against source evidence."


def test_edit_uses_human_version_and_preserves_original_ai_draft() -> None:
    review_bundle = _reviewed_bundle("edit")
    result = resolve_effective_insights(review_bundle)[0]
    assert result.verification_status == "edited"
    assert result.final_observation == "Human-reviewed adoption evidence is mixed."
    assert result.provenance == "human_edited"
    assert (
        review_bundle.source_bundle.insight_drafts[0].observation
        == "Several evidence records describe the adoption experience."
    )


def test_rejected_is_excluded_but_can_be_inspected() -> None:
    review_bundle = _reviewed_bundle("reject")
    assert resolve_effective_insights(review_bundle) == []
    rejected = resolve_effective_insights(review_bundle, include_rejected=True)[0]
    assert rejected.verification_status == "rejected"
    assert not rejected.downstream_eligible
    assert rejected.rejection_reason


def test_review_recommended_does_not_block_default_downstream() -> None:
    result = resolve_effective_insights(_bundle(evidence_count=1, confidence="low"))[0]
    assert result.review_recommended
    assert result.downstream_eligible
    assert "Low confidence" in result.review_reasons
    assert "Only 1 supporting evidence item" in result.review_reasons


@pytest.mark.parametrize(
    ("bundle", "expected_reason"),
    [
        (_bundle(confidence="low"), "Low confidence"),
        (_bundle(evidence_count=2), "Only 2 supporting evidence items"),
        (
            _bundle(coverage="emergent"),
            "Emergent theme not fully covered by the predefined taxonomy",
        ),
    ],
)
def test_structural_signals_produce_explainable_recommendations(
    bundle: ThemeInsightBundle,
    expected_reason: str,
) -> None:
    result = resolve_effective_insights(bundle)[0]
    assert result.review_recommended
    assert expected_reason in result.review_reasons


def test_high_confidence_sufficient_evidence_has_no_unfounded_recommendation() -> None:
    result = resolve_effective_insights(_bundle())[0]
    assert not result.review_recommended
    assert result.review_reasons == []


def test_strict_mode_excludes_unreviewed_and_default_is_non_strict() -> None:
    bundle = _bundle()
    assert len(resolve_effective_insights(bundle)) == 1
    assert resolve_effective_insights(bundle, require_review=True) == []
    assert len(resolve_effective_insights(_reviewed_bundle("confirm"), require_review=True)) == 1


def test_evidence_inspection_uses_generic_original_source_without_review() -> None:
    insight = resolve_effective_insights(_bundle())[0]
    inspection = inspect_effective_insight_evidence(insight, load_rows(FIXTURE))
    assert inspection.evidence_count == 3
    assert inspection.evidence[0].evidence_id == "gf-001"
    assert inspection.evidence[0].source_type == "product_review"
    assert inspection.evidence[0].text.startswith("The setup guide")
    assert inspection.evidence[0].metadata == {"engagement_count": 7}


def test_effective_insight_file_defaults_to_unreviewed_and_checks_source_binding(
    tmp_path: Path,
) -> None:
    source = tmp_path / "bundle.json"
    source.write_text(
        json.dumps(_bundle().model_dump(mode="json"), ensure_ascii=False),
        encoding="utf-8",
    )
    output = tmp_path / "effective.jsonl"
    records = resolve_effective_insight_file(source, output)
    assert records[0].verification_status == "unreviewed"
    assert json.loads(output.read_text(encoding="utf-8"))["downstream_eligible"] is True


def test_effective_insight_file_rejects_mismatched_review_source_hash(
    tmp_path: Path,
) -> None:
    source = tmp_path / "bundle.json"
    source.write_text(
        json.dumps(_bundle().model_dump(mode="json"), ensure_ascii=False),
        encoding="utf-8",
    )
    review = tmp_path / "review.json"
    review.write_text(
        json.dumps(_reviewed_bundle("confirm").model_dump(mode="json"), ensure_ascii=False),
        encoding="utf-8",
    )
    with pytest.raises(ClassificationWorkflowError, match="checksum"):
        resolve_effective_insight_file(
            source,
            tmp_path / "effective.jsonl",
            human_review_file=review,
        )


def test_workbook_shows_optional_verification_with_decision_sheet(tmp_path: Path) -> None:
    rows = load_rows(FIXTURE)
    output = create_readable_workbook(
        rows,
        tmp_path / "optional-verification.xlsx",
        profile_id="general_feedback",
        theme_bundle=_bundle(),
    )
    workbook = load_workbook(output)
    assert len(workbook.sheetnames) == 11
    assert workbook.sheetnames[-1] == "11_决策支持"
    draft_headers = [cell.value for cell in workbook["09_洞察草稿"][1]]
    assert {"Verification Status", "Review Recommended", "Review Reasons"} <= set(draft_headers)
    assert workbook["09_洞察草稿"]["F2"].value == "unreviewed"
    assert workbook["10_可选核验"]["J2"].value == "Unreviewed"


def test_resolve_insights_cli_exposes_optional_review_and_non_strict_default() -> None:
    result = CliRunner().invoke(app, ["resolve-insights", "--help"])
    assert result.exit_code == 0
    assert "--human-review" in result.output
    assert "--require-review" in result.output
    assert "--no-require-review" in result.output
