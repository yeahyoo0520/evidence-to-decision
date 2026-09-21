from __future__ import annotations

import hashlib
import csv
import json
from copy import deepcopy
from pathlib import Path

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from youtube_comment_research.cli import app
from youtube_comment_research.errors import ClassificationWorkflowError
from youtube_comment_research.human_review import (
    HumanReviewBundle,
    HumanReviewDecision,
    HumanReviewPacket,
    HumanReviewResponse,
    ReviewItem,
    apply_review_response,
    build_review_packet,
    prepare_review_packet,
)
from youtube_comment_research.readable_workbook import SHEET_NAMES, create_readable_workbook
from youtube_comment_research.theme_discovery import InsightDraft, ThemeInsightBundle, ThemeRecord

FIXED_TIME = "2026-08-23T00:00:00+00:00"


def _theme(index: int) -> ThemeRecord:
    first = index * 2 - 1
    evidence_ids = [f"c{first}", f"c{first + 1}"]
    return ThemeRecord(
        theme_id=f"theme-{index}",
        theme_name=f"Theme {index}",
        theme_description=f"Theme {index} description.",
        related_categories=["Product Praise"],
        evidence_count=2,
        evidence_comment_ids=evidence_ids,
        representative_comments=[f"Representative comment {first}."],
        representative_evidence_ids=[evidence_ids[0]],
        discovery_method="taxonomy_based",
        taxonomy_coverage="covered",
        coverage_reason=(
            "This theme was produced directly from the predefined taxonomy and is "
            "therefore treated as covered."
        ),
        confidence="medium",
        review_status="preliminary",
    )


def _insight(index: int, theme: ThemeRecord) -> InsightDraft:
    return InsightDraft(
        insight_id=f"insight-{index}",
        theme_id=theme.theme_id,
        observation=f"Two comments form Theme {index}.",
        interpretation=f"This pattern may indicate Theme {index} attention.",
        potential_opportunity="Review the evidence before confirming this potential insight.",
        evidence_ids=list(theme.evidence_comment_ids),
        confidence="medium",
        review_status="preliminary",
        draft_label="AI generated draft",
    )


def _bundle() -> ThemeInsightBundle:
    themes = [_theme(index) for index in range(1, 4)]
    return ThemeInsightBundle(
        profile_id="consumer_product",
        generation_method="classification-semantic-aggregation-v1",
        draft_label="AI generated draft",
        generated_at=FIXED_TIME,
        source_comment_count=6,
        themes=themes,
        insight_drafts=[
            _insight(index, theme) for index, theme in enumerate(themes, start=1)
        ],
    )


def _write_bundle(path: Path) -> tuple[Path, ThemeInsightBundle, str]:
    bundle = _bundle()
    path.write_text(
        json.dumps(bundle.model_dump(), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    checksum = hashlib.sha256(path.read_bytes()).hexdigest()
    return path, bundle, checksum


def _packet(tmp_path: Path) -> tuple[HumanReviewPacket, ThemeInsightBundle, str, Path]:
    source, bundle, checksum = _write_bundle(tmp_path / "source_bundle.json")
    return prepare_review_packet(source, tmp_path / "review_packet.json"), bundle, checksum, source


def _decision(
    review_id: str,
    action: str,
    **overrides: object,
) -> HumanReviewDecision:
    return HumanReviewDecision(
        review_id=review_id,
        action=action,
        reviewed_at=FIXED_TIME,
        **overrides,
    )


def _response(checksum: str, *decisions: HumanReviewDecision) -> HumanReviewResponse:
    return HumanReviewResponse(
        source_bundle_sha256=checksum,
        decisions=list(decisions),
    )


def _comments() -> list[dict[str, object]]:
    return [
        {
            "project_id": "phase3-lite2",
            "comment_id": f"c{index}",
            "comment_text": (
                f"Representative comment {index}."
                if index % 2
                else f"Supporting comment {index}."
            ),
            "video_id": "video-1",
            "video_title": "Fixture Video",
            "channel_title": "Fixture Channel",
            "is_reply": False,
            "parent_comment_id": "",
            "comment_like_count": index,
            "reply_count": 0,
            "collection_order": index,
            "collected_at": FIXED_TIME,
            "codex_primary_category": "Product Praise",
            "codex_secondary_categories": "",
            "codex_subtopics": "comfort",
            "codex_lifecycle_stage": "Using",
            "codex_confidence": 0.82,
            "research_profile": "consumer_product",
        }
        for index in range(1, 7)
    ]


def test_prepare_review_creates_pending_items(tmp_path: Path) -> None:
    packet, bundle, _, _ = _packet(tmp_path)
    assert len(packet.items) == len(bundle.insight_drafts)
    assert {item.review_status for item in packet.items} == {"pending_review"}


def test_review_item_references_source_theme(tmp_path: Path) -> None:
    packet, _, _, _ = _packet(tmp_path)
    assert all(item.source_theme_id == item.theme.theme_id for item in packet.items)


def test_review_item_references_source_insight(tmp_path: Path) -> None:
    packet, _, _, _ = _packet(tmp_path)
    assert all(
        item.source_insight_id == item.insight_draft.insight_id
        for item in packet.items
    )


def test_representative_text_matches_comment_id(tmp_path: Path) -> None:
    packet, _, _, _ = _packet(tmp_path)
    first = packet.items[0]
    assert first.representative_evidence[0].comment_id == "c1"
    assert first.representative_evidence[0].comment_text == "Representative comment 1."


def test_confirm_creates_confirmed_insight(tmp_path: Path) -> None:
    packet, bundle, checksum, _ = _packet(tmp_path)
    result = apply_review_response(
        bundle,
        _response(checksum, _decision(packet.items[0].review_id, "confirm")),
        source_bundle_sha256=checksum,
        generated_at=FIXED_TIME,
    )
    assert result.confirmed_insights[0].status == "human_confirmed"


def test_confirm_does_not_modify_original_ai_draft(tmp_path: Path) -> None:
    packet, bundle, checksum, _ = _packet(tmp_path)
    before = deepcopy(bundle.model_dump())
    apply_review_response(
        bundle,
        _response(checksum, _decision(packet.items[0].review_id, "confirm")),
        source_bundle_sha256=checksum,
    )
    assert bundle.model_dump() == before


def test_edit_creates_human_edited_confirmed_insight(tmp_path: Path) -> None:
    packet, bundle, checksum, _ = _packet(tmp_path)
    result = apply_review_response(
        bundle,
        _response(
            checksum,
            _decision(
                packet.items[0].review_id,
                "edit",
                observation_override="Human-edited observation.",
                reviewer_note="Clarified the observed pattern.",
            ),
        ),
        source_bundle_sha256=checksum,
    )
    assert result.confirmed_insights[0].status == "human_edited"
    assert result.confirmed_insights[0].observation == "Human-edited observation."


def test_edit_preserves_original_ai_draft_in_review_bundle(tmp_path: Path) -> None:
    packet, bundle, checksum, _ = _packet(tmp_path)
    original = bundle.insight_drafts[0].observation
    result = apply_review_response(
        bundle,
        _response(
            checksum,
            _decision(
                packet.items[0].review_id,
                "edit",
                observation_override="Human-edited observation.",
            ),
        ),
        source_bundle_sha256=checksum,
    )
    assert result.source_bundle.insight_drafts[0].observation == original
    assert result.review_items[0].insight_draft.observation == original


def test_edit_requires_a_modification_field() -> None:
    with pytest.raises(Exception, match="modification field"):
        HumanReviewDecision(review_id="review-1", action="edit")


def test_edit_requires_an_actual_change(tmp_path: Path) -> None:
    packet, bundle, checksum, _ = _packet(tmp_path)
    with pytest.raises(ClassificationWorkflowError, match="actual change"):
        apply_review_response(
            bundle,
            _response(
                checksum,
                _decision(
                    packet.items[0].review_id,
                    "edit",
                    observation_override=bundle.insight_drafts[0].observation,
                ),
            ),
            source_bundle_sha256=checksum,
        )


def test_edit_evidence_must_be_source_theme_subset(tmp_path: Path) -> None:
    packet, bundle, checksum, _ = _packet(tmp_path)
    with pytest.raises(ClassificationWorkflowError, match="outside the source Theme"):
        apply_review_response(
            bundle,
            _response(
                checksum,
                _decision(
                    packet.items[0].review_id,
                    "edit",
                    selected_evidence_ids=["c1", "outside"],
                    representative_evidence_ids=["c1"],
                ),
            ),
            source_bundle_sha256=checksum,
        )


def test_representative_evidence_must_be_selected_subset(tmp_path: Path) -> None:
    packet, bundle, checksum, _ = _packet(tmp_path)
    with pytest.raises(ClassificationWorkflowError, match="outside selected evidence"):
        apply_review_response(
            bundle,
            _response(
                checksum,
                _decision(
                    packet.items[0].review_id,
                    "edit",
                    selected_evidence_ids=["c2"],
                    representative_evidence_ids=["c1"],
                ),
            ),
            source_bundle_sha256=checksum,
        )


def test_reject_does_not_create_confirmed_insight(tmp_path: Path) -> None:
    packet, bundle, checksum, _ = _packet(tmp_path)
    result = apply_review_response(
        bundle,
        _response(
            checksum,
            _decision(
                packet.items[0].review_id,
                "reject",
                rejection_reason="The evidence does not support the interpretation.",
            ),
        ),
        source_bundle_sha256=checksum,
    )
    assert result.confirmed_insights == []
    assert result.rejected_review_ids == [packet.items[0].review_id]


def test_reject_requires_reason_or_note() -> None:
    with pytest.raises(Exception, match="requires rejection_reason"):
        HumanReviewDecision(review_id="review-1", action="reject")


def test_pending_item_does_not_create_confirmed_insight(tmp_path: Path) -> None:
    packet, bundle, checksum, _ = _packet(tmp_path)
    result = apply_review_response(
        bundle,
        _response(checksum),
        source_bundle_sha256=checksum,
    )
    assert result.confirmed_insights == []
    assert result.pending_review_ids == [item.review_id for item in packet.items]


def test_conflicting_decisions_for_same_review_item_are_rejected(tmp_path: Path) -> None:
    packet, _, checksum, _ = _packet(tmp_path)
    review_id = packet.items[0].review_id
    with pytest.raises(Exception, match="conflicting"):
        _response(
            checksum,
            _decision(review_id, "confirm"),
            _decision(review_id, "reject", rejection_reason="Reject it."),
        )


def test_unknown_review_id_is_rejected(tmp_path: Path) -> None:
    _, bundle, checksum, _ = _packet(tmp_path)
    with pytest.raises(ClassificationWorkflowError, match="unknown review_id"):
        apply_review_response(
            bundle,
            _response(checksum, _decision("review-missing", "confirm")),
            source_bundle_sha256=checksum,
        )


def test_nonexistent_source_theme_id_is_rejected(tmp_path: Path) -> None:
    packet, _, _, _ = _packet(tmp_path)
    payload = packet.items[0].model_dump()
    payload["source_theme_id"] = "theme-missing"
    with pytest.raises(Exception, match="source_theme_id"):
        ReviewItem.model_validate(payload)


def test_nonexistent_source_insight_id_is_rejected(tmp_path: Path) -> None:
    packet, _, _, _ = _packet(tmp_path)
    payload = packet.items[0].model_dump()
    payload["source_insight_id"] = "insight-missing"
    with pytest.raises(Exception, match="source_insight_id"):
        ReviewItem.model_validate(payload)


def test_confirmed_insight_keeps_complete_evidence_chain(tmp_path: Path) -> None:
    packet, bundle, checksum, _ = _packet(tmp_path)
    result = apply_review_response(
        bundle,
        _response(checksum, _decision(packet.items[0].review_id, "confirm")),
        source_bundle_sha256=checksum,
    )
    confirmed = result.confirmed_insights[0]
    assert confirmed.source_theme_id == packet.items[0].source_theme_id
    assert confirmed.source_insight_id == packet.items[0].source_insight_id
    assert set(confirmed.evidence_ids) <= set(packet.items[0].theme.evidence_comment_ids)


def test_partial_review_progress_can_be_saved(tmp_path: Path) -> None:
    packet, bundle, checksum, _ = _packet(tmp_path)
    result = apply_review_response(
        bundle,
        _response(checksum, _decision(packet.items[0].review_id, "confirm")),
        source_bundle_sha256=checksum,
    )
    assert len(result.confirmed_insights) == 1
    assert len(result.pending_review_ids) == 2


def test_unreviewed_items_remain_pending(tmp_path: Path) -> None:
    packet, bundle, checksum, _ = _packet(tmp_path)
    result = apply_review_response(
        bundle,
        _response(
            checksum,
            _decision(
                packet.items[1].review_id,
                "reject",
                reviewer_note="Not useful for this research question.",
            ),
        ),
        source_bundle_sha256=checksum,
    )
    assert set(result.pending_review_ids) == {
        packet.items[0].review_id,
        packet.items[2].review_id,
    }


def test_review_json_models_round_trip(tmp_path: Path) -> None:
    packet, bundle, checksum, _ = _packet(tmp_path)
    response = _response(checksum, _decision(packet.items[0].review_id, "confirm"))
    review_bundle = apply_review_response(
        bundle,
        response,
        source_bundle_sha256=checksum,
        generated_at=FIXED_TIME,
    )
    assert HumanReviewPacket.model_validate_json(packet.model_dump_json()) == packet
    assert HumanReviewResponse.model_validate_json(response.model_dump_json()) == response
    assert HumanReviewBundle.model_validate_json(review_bundle.model_dump_json()) == review_bundle
    schema_root = Path(__file__).parents[1] / "schemas"
    for name in (
        "human_review_packet.schema.json",
        "human_review_response.schema.json",
        "confirmed_insight.schema.json",
        "human_review_bundle.schema.json",
    ):
        assert json.loads((schema_root / name).read_text(encoding="utf-8"))["type"] == "object"


def _mixed_review_bundle(tmp_path: Path) -> HumanReviewBundle:
    packet, bundle, checksum, _ = _packet(tmp_path)
    return apply_review_response(
        bundle,
        _response(
            checksum,
            _decision(packet.items[0].review_id, "confirm"),
            _decision(
                packet.items[1].review_id,
                "reject",
                rejection_reason="Insufficient evidence.",
            ),
        ),
        source_bundle_sha256=checksum,
        generated_at=FIXED_TIME,
    )


def test_workbook_adds_tenth_human_review_sheet_and_reopens(tmp_path: Path) -> None:
    review_bundle = _mixed_review_bundle(tmp_path)
    output = create_readable_workbook(
        _comments(),
        tmp_path / "review.xlsx",
        human_review_bundle=review_bundle,
    )
    workbook = load_workbook(output)
    assert workbook.sheetnames == SHEET_NAMES
    assert workbook.sheetnames[-2:] == ["10_人工审核", "11_决策支持"]
    statuses = {cell.value for cell in workbook["10_人工审核"]["J"][1:]}
    assert {"Reviewed", "Rejected", "Unreviewed"} <= statuses


def test_human_review_workbook_has_no_formula_error_literals(tmp_path: Path) -> None:
    review_bundle = _mixed_review_bundle(tmp_path)
    output = create_readable_workbook(
        _comments(),
        tmp_path / "review.xlsx",
        human_review_bundle=review_bundle,
    )
    workbook = load_workbook(output, data_only=False)
    errors = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"}
    assert not [
        cell.coordinate
        for sheet in workbook
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value in errors
    ]


def test_prepare_review_cli_writes_packet(tmp_path: Path) -> None:
    source, _, _ = _write_bundle(tmp_path / "source.json")
    output = tmp_path / "packet.json"
    result = CliRunner().invoke(
        app,
        ["prepare-review", str(source), "--output", str(output)],
    )
    assert result.exit_code == 0, result.output
    assert len(json.loads(output.read_text(encoding="utf-8"))["items"]) == 3


def test_apply_review_cli_writes_partial_bundle(tmp_path: Path) -> None:
    source, _, checksum = _write_bundle(tmp_path / "source.json")
    packet = build_review_packet(
        _bundle(),
        source_bundle_sha256=checksum,
        prepared_at=FIXED_TIME,
    )
    response_path = tmp_path / "response.json"
    response_path.write_text(
        json.dumps(
            _response(
                checksum,
                _decision(packet.items[0].review_id, "confirm"),
            ).model_dump(),
            indent=2,
        ),
        encoding="utf-8",
    )
    output = tmp_path / "review_bundle.json"
    result = CliRunner().invoke(
        app,
        [
            "apply-review",
            str(source),
            "--review",
            str(response_path),
            "--output",
            str(output),
        ],
    )
    assert result.exit_code == 0, result.output
    payload = json.loads(output.read_text(encoding="utf-8"))
    assert len(payload["confirmed_insights"]) == 1
    assert len(payload["pending_review_ids"]) == 2


def test_report_cli_loads_human_review_bundle(tmp_path: Path) -> None:
    review_bundle = _mixed_review_bundle(tmp_path)
    review_path = tmp_path / "human_review_bundle.json"
    review_path.write_text(
        json.dumps(review_bundle.model_dump(), ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    rows = _comments()
    comments_path = tmp_path / "comments.csv"
    with comments_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    output = tmp_path / "review_report.xlsx"
    result = CliRunner().invoke(
        app,
        [
            "report",
            str(comments_path),
            "--output",
            str(output),
            "--human-review",
            str(review_path),
        ],
    )
    assert result.exit_code == 0, result.output
    assert load_workbook(output).sheetnames[-2:] == ["10_人工审核", "11_决策支持"]
