from __future__ import annotations

import csv
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from youtube_comment_research.cli import app
from youtube_comment_research.errors import ClassificationWorkflowError
from youtube_comment_research.readable_workbook import SHEET_NAMES, create_readable_workbook
from youtube_comment_research.semantic_workflow import (
    BATCH_INPUT_FIELDS,
    RUBRIC_VERSION,
    build_review_sections,
    build_rule_candidate,
    create_blind_input_copy,
    effective_classification,
    merge_classifications,
    prepare_classification_batches,
    prepare_report_rows,
    validate_classification_payload,
)


def _source_rows(count: int = 3) -> list[dict[str, object]]:
    texts = [
        "I love these headphones.",
        "Does the microphone work on calls?",
        "Great video and editing.",
    ]
    return [
        {
            "project_id": "phase2a2-fixture",
            "comment_id": f"comment-{index + 1}",
            "comment_text": texts[index],
            "video_id": "video-one",
            "video_title": "Fixture video",
            "channel_title": "Fixture channel",
            "is_reply": False,
            "parent_comment_id": "",
            "comment_like_count": index + 1,
            "reply_count": 0,
            "collection_order": index + 1,
        }
        for index in range(count)
    ]


def _write_csv(path: Path, rows: list[dict[str, object]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields: list[str] = []
    for row in rows:
        for field in row:
            if field not in fields:
                fields.append(field)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
    return path


def _payload(item: dict[str, object], **overrides: object) -> dict[str, object]:
    text = str(item["comment_text"])
    if text.startswith("Does"):
        subject = "product"
        category = "Product Question"
        subtopics = ["microphone/calls"]
        evidence = ["microphone work"]
    elif "video" in text:
        subject = "video_or_creator"
        category = "Other"
        subtopics = []
        evidence = ["Great video"]
    else:
        subject = "product"
        category = "Product Praise"
        subtopics = []
        evidence = ["love these"]
    value: dict[str, object] = {
        "comment_id": item["comment_id"],
        "subject": subject,
        "primary_category": category,
        "secondary_categories": [],
        "subtopics": subtopics,
        "lifecycle_stage": "Unknown",
        "mentioned_brands": [],
        "mentioned_products": [],
        "confidence": 0.91,
        "evidence": evidence,
        "review_risk": "low",
        "review_triggers": [],
        "manual_review": False,
        "review_reason": "",
        "classification_source": "codex-assisted",
        "classifier_version": "offline-test-v1",
        "rubric_version": RUBRIC_VERSION,
    }
    value.update(overrides)
    return value


def _prepare_with_results(
    tmp_path: Path,
    rows: list[dict[str, object]] | None = None,
    *,
    batch_size: int = 2,
) -> tuple[Path, Path]:
    source_rows = rows or _source_rows()
    source = _write_csv(tmp_path / "source.csv", source_rows)
    batch_dir = tmp_path / "batches"
    summary = prepare_classification_batches(source, batch_dir, batch_size=batch_size)
    for batch_path in summary.batch_files:
        items = [
            json.loads(line)
            for line in batch_path.read_text(encoding="utf-8").splitlines()
        ]
        result_path = batch_dir / f"classified_{batch_path.name}"
        result_path.write_text(
            "".join(
                json.dumps(_payload(item), ensure_ascii=False) + "\n"
                for item in items
            ),
            encoding="utf-8",
        )
    return source, batch_dir


def test_phase2a2_reference_and_schema_files_exist() -> None:
    root = Path(__file__).resolve().parents[1]
    assert (root / "references" / "classification_rubric.md").exists()
    assert (root / "references" / "classification_examples.jsonl").exists()
    assert (root / "schemas" / "comment_classification.schema.json").exists()


def test_prepare_batches_enforces_maximum_twenty(tmp_path: Path) -> None:
    rows = [
        {
            "comment_id": f"id-{index}",
            "comment_text": f"Comment {index}",
            "video_id": "v",
            "is_reply": False,
        }
        for index in range(41)
    ]
    source = _write_csv(tmp_path / "source.csv", rows)
    summary = prepare_classification_batches(source, tmp_path / "batches", batch_size=20)
    assert summary.batch_count == 3
    assert [
        len(path.read_text(encoding="utf-8").splitlines())
        for path in summary.batch_files
    ] == [20, 20, 1]


def test_prepare_batches_has_no_omitted_or_duplicate_ids(tmp_path: Path) -> None:
    source = _write_csv(tmp_path / "source.csv", _source_rows())
    summary = prepare_classification_batches(source, tmp_path / "batches", batch_size=2)
    manifest = json.loads(summary.manifest_path.read_text(encoding="utf-8"))
    ids = [
        item["comment_id"]
        for path in summary.batch_files
        for item in (
            json.loads(line)
            for line in path.read_text(encoding="utf-8").splitlines()
        )
    ]
    assert ids == manifest["comment_ids"]
    assert len(ids) == len(set(ids)) == 3


def test_prepare_batches_excludes_predictions_and_expected_answers(tmp_path: Path) -> None:
    rows = _source_rows(1)
    rows[0].update(
        {
            "predicted_primary_category": "Other",
            "expected_primary_category": "Product Praise",
            "rule_candidate_categories": "Other",
            "annotation_note": "secret",
        }
    )
    source = _write_csv(tmp_path / "source.csv", rows)
    summary = prepare_classification_batches(source, tmp_path / "batches")
    item = json.loads(summary.batch_files[0].read_text(encoding="utf-8"))
    assert set(item) == set(BATCH_INPUT_FIELDS)
    assert not any(
        key.startswith(("predicted_", "expected_", "rule_")) for key in item
    )
    assert "annotation_note" not in item


def test_blind_copy_excludes_expected_and_prediction_fields(tmp_path: Path) -> None:
    rows = _source_rows(1)
    rows[0]["expected_primary_category"] = "Product Praise"
    rows[0]["predicted_primary_category"] = "Other"
    source = _write_csv(tmp_path / "source.csv", rows)
    output = tmp_path / "blind.csv"
    assert create_blind_input_copy(source, output) == 1
    with output.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        assert reader.fieldnames == list(BATCH_INPUT_FIELDS)


def test_prepare_manifest_and_checksums_are_created(tmp_path: Path) -> None:
    source = _write_csv(tmp_path / "source.csv", _source_rows())
    summary = prepare_classification_batches(source, tmp_path / "batches")
    checksums = json.loads(summary.checksums_path.read_text(encoding="utf-8"))
    assert summary.manifest_path.name in checksums["files"]
    assert summary.batch_files[0].name in checksums["files"]


def test_valid_codex_payload_passes_schema_and_evidence_check() -> None:
    item = {"comment_id": "one", "comment_text": "I love these headphones."}
    result = validate_classification_payload(
        _payload(item),
        input_text=item["comment_text"],
    )
    assert result.primary_category == "Product Praise"


@pytest.mark.parametrize(
    ("overrides", "error_fragment"),
    [
        ({"primary_category": "Not a category"}, "Schema validation failed"),
        ({"confidence": 1.1}, "Schema validation failed"),
        ({"classification_source": "rule-based"}, "Schema validation failed"),
    ],
)
def test_invalid_codex_payload_is_rejected(
    overrides: dict[str, object],
    error_fragment: str,
) -> None:
    item = {"comment_id": "one", "comment_text": "I love these headphones."}
    with pytest.raises(ClassificationWorkflowError, match=error_fragment):
        validate_classification_payload(
            _payload(item, **overrides),
            input_text=item["comment_text"],
        )


def test_non_exact_evidence_is_rejected() -> None:
    item = {"comment_id": "one", "comment_text": "I love these headphones."}
    with pytest.raises(ClassificationWorkflowError, match="exact excerpt"):
        validate_classification_payload(
            _payload(item, evidence=["excellent sound"]),
            input_text=item["comment_text"],
        )


def test_merge_rejects_malformed_json(tmp_path: Path) -> None:
    source, batch_dir = _prepare_with_results(tmp_path)
    (batch_dir / "classified_batch_001.jsonl").write_text("{bad json\n", encoding="utf-8")
    with pytest.raises(ClassificationWorkflowError, match="invalid JSON"):
        merge_classifications(source, batch_dir, tmp_path / "merged.csv")


def test_merge_rejects_unknown_comment_id(tmp_path: Path) -> None:
    source, batch_dir = _prepare_with_results(tmp_path, _source_rows(1))
    result = batch_dir / "classified_batch_001.jsonl"
    item = json.loads(result.read_text(encoding="utf-8"))
    item["comment_id"] = "unknown"
    result.write_text(json.dumps(item) + "\n", encoding="utf-8")
    with pytest.raises(ClassificationWorkflowError, match="Unknown comment_id"):
        merge_classifications(source, batch_dir, tmp_path / "merged.csv")


def test_merge_rejects_duplicate_comment_id(tmp_path: Path) -> None:
    source, batch_dir = _prepare_with_results(tmp_path, _source_rows(1))
    result = batch_dir / "classified_batch_001.jsonl"
    line = result.read_text(encoding="utf-8")
    result.write_text(line + line, encoding="utf-8")
    with pytest.raises(ClassificationWorkflowError, match="Duplicate classification"):
        merge_classifications(source, batch_dir, tmp_path / "merged.csv")


def test_merge_rejects_missing_comment_id(tmp_path: Path) -> None:
    source, batch_dir = _prepare_with_results(tmp_path)
    result = batch_dir / "classified_batch_002.jsonl"
    result.write_text("", encoding="utf-8")
    with pytest.raises(ClassificationWorkflowError, match="Missing classification comment_id"):
        merge_classifications(source, batch_dir, tmp_path / "merged.csv")


def test_merge_rejects_missing_batch(tmp_path: Path) -> None:
    source, batch_dir = _prepare_with_results(tmp_path)
    (batch_dir / "classified_batch_002.jsonl").unlink()
    with pytest.raises(ClassificationWorkflowError, match="Missing classification batch"):
        merge_classifications(source, batch_dir, tmp_path / "merged.csv")


def test_merge_preserves_human_review_fields(tmp_path: Path) -> None:
    rows = _source_rows(1)
    rows[0].update(
        {
            "human_confirmed_primary_category": "Product Concern",
            "human_confirmed_secondary_categories": "",
            "human_confirmed_subtopics": "durability",
            "human_confirmed_lifecycle_stage": "Using",
            "review_status": "corrected",
            "reviewed_at": "2026-07-27T00:00:00Z",
            "reviewer_note": "Human correction must survive.",
        }
    )
    source, batch_dir = _prepare_with_results(tmp_path, rows)
    output = tmp_path / "merged.csv"
    merge_classifications(source, batch_dir, output)
    with output.open("r", encoding="utf-8-sig", newline="") as handle:
        merged = next(csv.DictReader(handle))
    assert merged["human_confirmed_primary_category"] == "Product Concern"
    assert merged["review_status"] == "corrected"
    assert merged["reviewer_note"] == "Human correction must survive."


def test_reclassification_preserves_prior_codex_version(tmp_path: Path) -> None:
    source, first_dir = _prepare_with_results(tmp_path / "first", _source_rows(1))
    first_output = tmp_path / "first-merged.csv"
    merge_classifications(source, first_dir, first_output)
    second_dir = tmp_path / "second-batches"
    summary = prepare_classification_batches(first_output, second_dir)
    item = json.loads(summary.batch_files[0].read_text(encoding="utf-8"))
    payload = _payload(item, classifier_version="offline-test-v2")
    (second_dir / "classified_batch_001.jsonl").write_text(
        json.dumps(payload) + "\n",
        encoding="utf-8",
    )
    second_output = tmp_path / "second-merged.csv"
    merge_classifications(first_output, second_dir, second_output)
    with second_output.open("r", encoding="utf-8-sig", newline="") as handle:
        merged = next(csv.DictReader(handle))
    history = json.loads(merged["classification_history"])
    assert history[-1]["classifier_version"] == "offline-test-v1"
    assert merged["classifier_version"] == "offline-test-v2"


def test_rule_output_is_explicitly_labeled_candidate() -> None:
    candidate = build_rule_candidate(
        {"comment_id": "one", "comment_text": "I love these headphones."}
    )
    assert candidate["rule_classification_label"] == "Rule-based candidate / 规则候选"
    assert candidate["rule_candidate_categories"]


def test_human_confirmation_has_precedence_over_codex_and_rule() -> None:
    effective = effective_classification(
        {
            "rule_candidate_categories": "Product Praise",
            "codex_primary_category": "Product Question",
            "human_confirmed_primary_category": "Product Concern",
            "review_status": "corrected",
        }
    )
    assert effective["source"] == "human-reviewed"
    assert effective["primary_category"] == "Product Concern"


def test_review_sections_do_not_mix_rule_candidates_with_findings() -> None:
    rows = prepare_report_rows(_source_rows())
    sections = build_review_sections(rows)
    assert {row["report_section"] for row in sections} == {
        "Confirmed Findings",
        "Preliminary Codex-assisted Findings",
        "Review Queue",
    }
    assert all(
        row["count"] == 0
        for row in sections
        if row["report_section"] != "Review Queue"
    )


def test_workbook_shows_rule_codex_and_human_layers(tmp_path: Path) -> None:
    rows = _source_rows(3)
    rows[0].update(
        {
            "codex_subject": "product",
            "codex_primary_category": "Product Praise",
            "codex_confidence": 0.91,
            "codex_evidence": '["love these"]',
            "classification_source": "codex-assisted",
            "classifier_version": "offline-test-v1",
            "rubric_version": RUBRIC_VERSION,
            "review_status": "pending",
        }
    )
    rows[1].update(
        {
            "codex_subject": "product",
            "codex_primary_category": "Product Question",
            "codex_manual_review": True,
            "codex_review_reason": "Pronoun target needs confirmation.",
            "classification_source": "codex-assisted",
            "review_status": "pending",
        }
    )
    rows[2].update(
        {
            "codex_subject": "video_or_creator",
            "codex_primary_category": "Other",
            "human_confirmed_primary_category": "Other",
            "review_status": "confirmed",
            "reviewer_note": "Confirmed creator praise.",
        }
    )
    path = create_readable_workbook(rows, tmp_path / "review.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == SHEET_NAMES
    headers = [cell.value for cell in workbook[SHEET_NAMES[2]][1]]
    assert "规则候选分类 Rule-based candidate" in headers
    assert "Codex 主分类 Primary category" in headers
    assert "人工确认主分类 Human primary" in headers
    assert "Review Status" in headers
    sections = {
        cell.value
        for cell in workbook[SHEET_NAMES[3]]["A"][1:]
    }
    assert {
        "Confirmed Findings",
        "Preliminary Insights",
        "Insight Review Queue",
    }.issubset(sections)
    assert workbook[SHEET_NAMES[5]].max_row == len(rows) + 1


def test_cli_help_exposes_phase2a2_commands() -> None:
    result = CliRunner().invoke(app, ["--help"])
    assert result.exit_code == 0
    assert "prepare-classification" in result.stdout
    assert "merge-classification" in result.stdout
    assert "sanitize-classification-input" in result.stdout
