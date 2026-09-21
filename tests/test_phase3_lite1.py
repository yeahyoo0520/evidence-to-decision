from __future__ import annotations

import csv
from copy import deepcopy
from pathlib import Path

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from youtube_comment_research.cli import app
from youtube_comment_research.errors import ClassificationWorkflowError
from youtube_comment_research.readable_workbook import SHEET_NAMES, create_readable_workbook
from youtube_comment_research.theme_discovery import (
    InsightDraft,
    ThemeRecord,
    discover_themes,
    generate_theme_insights,
    load_theme_insight_bundle,
    validate_theme_insight_bundle,
)


def _row(
    comment_id: str,
    *,
    category: str = "Product Praise",
    subtopic: str = "comfort",
    profile: str = "consumer_product",
) -> dict[str, object]:
    return {
        "project_id": "phase3-lite-fixture",
        "comment_id": comment_id,
        "comment_text": f"Fixture statement {comment_id}",
        "video_id": "video-1",
        "video_title": "Fixture Video",
        "channel_title": "Fixture Channel",
        "is_reply": False,
        "parent_comment_id": "",
        "comment_like_count": 3,
        "reply_count": 0,
        "collection_order": 1,
        "collected_at": "2026-08-21T00:00:00Z",
        "codex_primary_category": category,
        "codex_secondary_categories": "",
        "codex_subtopics": subtopic,
        "codex_lifecycle_stage": "Unknown" if profile == "consumer_product" else "Pre-launch",
        "codex_confidence": 0.88,
        "research_profile": profile,
    }


def _write_csv(path: Path, rows: list[dict[str, object]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = list(rows[0]) if rows else ["comment_id", "comment_text"]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
    return path


def test_theme_output_fields_are_complete() -> None:
    expected = {
        "theme_id",
        "theme_name",
        "theme_description",
        "related_categories",
        "evidence_count",
        "evidence_comment_ids",
        "representative_comments",
        "representative_evidence_ids",
        "discovery_method",
        "taxonomy_coverage",
        "coverage_reason",
        "confidence",
        "review_status",
    }
    assert set(ThemeRecord.model_fields) == expected


def test_theme_evidence_ids_exist_in_classification_input() -> None:
    rows = [_row("c-1"), _row("c-2")]
    themes = discover_themes(rows, profile="consumer_product")
    source_ids = {str(row["comment_id"]) for row in rows}
    assert themes
    assert all(set(theme.evidence_comment_ids) <= source_ids for theme in themes)


def test_insight_must_reference_an_existing_theme(tmp_path: Path) -> None:
    rows = [_row("c-1")]
    output = tmp_path / "bundle.json"
    bundle = generate_theme_insights(
        _write_csv(tmp_path / "input.csv", rows),
        output,
        profile_id="consumer_product",
    )
    payload = bundle.model_dump()
    payload["insight_drafts"][0]["theme_id"] = "theme-missing"
    with pytest.raises(ClassificationWorkflowError, match="unknown themes"):
        validate_theme_insight_bundle(payload, rows, profile="consumer_product")


def test_empty_comment_input_generates_empty_readable_bundle(tmp_path: Path) -> None:
    bundle = generate_theme_insights(
        _write_csv(tmp_path / "empty.csv", []),
        tmp_path / "empty-bundle.json",
        profile_id="consumer_product",
    )
    assert bundle.source_comment_count == 0
    assert bundle.themes == []
    assert bundle.insight_drafts == []


def test_different_profiles_generate_profile_specific_themes() -> None:
    consumer = discover_themes(
        [_row("consumer", category="Product Praise", subtopic="comfort")],
        profile="consumer_product",
    )
    game = discover_themes(
        [
            _row(
                "game",
                category="Version Expectation",
                subtopic="release timing",
                profile="game_publishing",
            )
        ],
        profile="game_publishing",
    )
    assert consumer[0].related_categories == ["Product Praise"]
    assert game[0].related_categories == ["Version Expectation"]
    assert consumer[0].theme_id != game[0].theme_id


def test_generated_output_can_be_reloaded(tmp_path: Path) -> None:
    rows = [_row("c-1"), _row("c-2")]
    output = tmp_path / "bundle.json"
    generated = generate_theme_insights(
        _write_csv(tmp_path / "input.csv", rows),
        output,
        profile_id="consumer_product",
    )
    reloaded = load_theme_insight_bundle(
        output,
        source_rows=rows,
        profile_id="consumer_product",
    )
    assert reloaded == generated


def test_theme_discovery_does_not_modify_classification_rows() -> None:
    rows = [_row("c-1"), _row("c-2")]
    before = deepcopy(rows)
    discover_themes(rows, profile="consumer_product")
    assert rows == before


def test_insight_draft_language_is_cautious() -> None:
    with pytest.raises(ValueError, match="certainty language"):
        InsightDraft(
            insight_id="i-1",
            theme_id="t-1",
            observation="One classified comment forms a theme.",
            interpretation="This definitely proves demand.",
            potential_opportunity="A potential direction is further review.",
            evidence_ids=["c-1"],
            confidence="low",
            review_status="preliminary",
            draft_label="AI generated draft",
        )


def test_workbook_appends_theme_and_insight_sheets(tmp_path: Path) -> None:
    rows = [_row("c-1"), _row("c-2")]
    bundle = generate_theme_insights(
        _write_csv(tmp_path / "input.csv", rows),
        tmp_path / "bundle.json",
        profile_id="consumer_product",
    )
    workbook_path = create_readable_workbook(
        rows,
        tmp_path / "research.xlsx",
        theme_bundle=bundle,
    )
    workbook = load_workbook(workbook_path)
    assert workbook.sheetnames == SHEET_NAMES
    assert workbook.sheetnames[-4:] == [
        "08_主题发现",
        "09_洞察草稿",
        "10_人工审核",
        "11_决策支持",
    ]
    assert [cell.value for cell in workbook["08_主题发现"][1]] == [
        "Theme",
        "Description",
        "Evidence Count",
        "Related Categories",
        "Discovery Method",
        "Taxonomy Coverage",
        "Coverage Reason",
        "Confidence",
        "Review Status",
        "Representative Comments",
        "Representative Evidence IDs",
        "Evidence IDs",
    ]


def test_generate_insights_cli_help_is_available() -> None:
    result = CliRunner().invoke(app, ["generate-insights", "--help"])
    assert result.exit_code == 0
    assert "--profile" in result.output
    assert "--output" in result.output
