from datetime import datetime, timezone

from openpyxl import load_workbook

from youtube_comment_research.exporters import EXPORT_FIELDS, export_csv, export_jsonl, export_xlsx
from youtube_comment_research.models import CommentRecord


def sample_record() -> CommentRecord:
    return CommentRecord(
        project_id="test",
        video_id="25_IVHMGTVU",
        video_url="https://www.youtube.com/watch?v=25_IVHMGTVU",
        video_title="测试 Video",
        comment_id="comment-1",
        comment_text="Great sound 🎧 很舒服",
        author_display_name="Commenter 001",
        collected_at=datetime.now(timezone.utc),
        collection_order=1,
    )


def test_csv_is_utf8_sig_and_has_headers(tmp_path) -> None:
    path = export_csv([sample_record()], tmp_path / "comments.csv")
    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    text = raw.decode("utf-8-sig")
    assert "comment_text" in text
    assert "🎧" in text


def test_jsonl_preserves_unicode(tmp_path) -> None:
    path = export_jsonl([sample_record()], tmp_path / "comments.jsonl")
    assert "很舒服" in path.read_text(encoding="utf-8")


def test_xlsx_has_expected_header(tmp_path) -> None:
    path = export_xlsx([sample_record()], tmp_path / "comments.xlsx")
    workbook = load_workbook(path, read_only=True)
    header = [cell.value for cell in next(workbook["06_原始数据"].iter_rows())]
    assert header == EXPORT_FIELDS
