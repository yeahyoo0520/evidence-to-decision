from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, Callable

import pytest
from openpyxl import load_workbook
from typer._click.utils import strip_ansi
from typer.testing import CliRunner

import youtube_comment_research.cli as cli_module
from youtube_comment_research.cli import app
from youtube_comment_research.errors import ClassificationWorkflowError
from youtube_comment_research.readable_workbook import SHEET_NAMES, create_readable_workbook
from youtube_comment_research.theme_discovery import (
    ConsolidatedThemeBatch,
    ThemeInsightBundle,
    discover_inductive_themes,
    discover_themes,
    draft_insights,
    generate_theme_insights,
    load_theme_insight_bundle,
    validate_theme_insight_bundle,
)
from youtube_comment_research.theme_provider import CodexFileThemeDiscoveryProvider


def _row(
    comment_id: str,
    text: str,
    category: str,
    *,
    secondary: str = "",
) -> dict[str, object]:
    return {
        "project_id": "phase3-lite15-fixture",
        "comment_id": comment_id,
        "comment_text": text,
        "video_id": "game-video",
        "video_title": "Game Trailer",
        "channel_title": "Publisher",
        "is_reply": False,
        "parent_comment_id": "",
        "comment_like_count": 2,
        "reply_count": 0,
        "collection_order": int(comment_id.removeprefix("c")),
        "collected_at": "2026-08-21T00:00:00Z",
        "codex_primary_category": category,
        "codex_secondary_categories": secondary,
        "codex_subtopics": "",
        "codex_lifecycle_stage": "Pre-launch",
        "codex_confidence": 0.86,
        "research_profile": "game_publishing",
    }


def _rows() -> list[dict[str, object]]:
    return [
        _row("c1", "I still do not know when Europe launches.", "Version Expectation"),
        _row(
            "c2",
            "Japan can play already, but Europe has no release date.",
            "Version Expectation",
            secondary="Community / Social",
        ),
        _row(
            "c3",
            "The characters look amazing, but Europe still has no date.",
            "Character & Story",
            secondary="Version Expectation",
        ),
        _row("c4", "The character designs are gorgeous.", "Character & Story"),
        _row("c5", "Ranked matches feel unfair because teams are unbalanced.", "Gameplay Feedback"),
        _row("c6", "Competitive matchmaking keeps creating unfair teams.", "Gameplay Feedback"),
    ]


def _write_csv(path: Path, rows: list[dict[str, object]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)
    return path


class FakeThemeProvider:
    provider_id = "fake"

    def __init__(
        self,
        candidate_batches: list[dict[str, object]],
        consolidate: Callable[[list[dict[str, Any]]], dict[str, object]],
    ) -> None:
        self.candidate_batches = list(candidate_batches)
        self.consolidate = consolidate
        self.candidate_calls: list[list[dict[str, Any]]] = []

    def discover_candidate_themes(
        self, comments: list[dict[str, Any]], *, profile: object
    ) -> object:
        self.candidate_calls.append(comments)
        return self.candidate_batches.pop(0)

    def consolidate_themes(
        self,
        candidates: list[dict[str, Any]],
        comments: list[dict[str, Any]],
        *,
        profile: object,
    ) -> object:
        return self.consolidate(candidates)


def _three_theme_provider() -> FakeThemeProvider:
    candidates = {
        "themes": [
            {
                "theme_name": "Regional Release Questions",
                "description": "Players lack a clear Europe release date.",
                "evidence_ids": ["c1", "c2", "c3"],
                "representative_evidence_ids": ["c2"],
            },
            {
                "theme_name": "Character Visual Appeal",
                "description": "Players praise character appearance.",
                "evidence_ids": ["c3", "c4"],
                "representative_evidence_ids": ["c3"],
            },
            {
                "theme_name": "Competitive Fairness Concerns",
                "description": "Players describe unfair competitive team matching.",
                "evidence_ids": ["c5", "c6"],
                "representative_evidence_ids": ["c5"],
            },
        ]
    }

    def consolidate(items: list[dict[str, Any]]) -> dict[str, object]:
        coverage = {
            "Regional Release Questions": "covered",
            "Character Visual Appeal": "covered",
            "Competitive Fairness Concerns": "partially_covered",
        }
        coverage_reasons = {
            "Regional Release Questions": (
                "Version Expectation explicitly covers release and region questions, "
                "which fully expresses the Europe launch-date uncertainty in the evidence."
            ),
            "Character Visual Appeal": (
                "Character & Story already covers discussion of character design and "
                "appearance, which is the Theme's complete core meaning."
            ),
            "Competitive Fairness Concerns": (
                "Gameplay Feedback covers balance and play experience, but it does not "
                "fully express the recurring concern about fairness in team formation."
            ),
        }
        return {
            "themes": [
                {
                    "theme_name": item["theme_name"],
                    "description": item["description"],
                    "source_candidate_ids": [item["candidate_id"]],
                    "representative_evidence_ids": item[
                        "representative_evidence_ids"
                    ],
                    "taxonomy_coverage": coverage[item["theme_name"]],
                    "coverage_reason": coverage_reasons[item["theme_name"]],
                }
                for item in items
            ]
        }

    return FakeThemeProvider([candidates], consolidate)


def _emergent_rows() -> list[dict[str, object]]:
    return [
        _row(
            "c7",
            "I will not install this game while its anti-cheat requires kernel-level access to my computer.",
            "Other",
        ),
        _row(
            "c8",
            "Kernel-level anti-cheat feels invasive; the publisher should explain what data it can access.",
            "Other",
        ),
    ]


def _emergent_provider() -> FakeThemeProvider:
    candidates = {
        "themes": [
            {
                "theme_name": "Kernel-level Anti-cheat Privacy Concerns",
                "description": (
                    "Players repeatedly question the privacy and trust implications of "
                    "kernel-level anti-cheat access."
                ),
                "evidence_ids": ["c7", "c8"],
                "representative_evidence_ids": ["c7", "c8"],
            }
        ]
    }

    def consolidate(items: list[dict[str, Any]]) -> dict[str, object]:
        return {
            "themes": [
                {
                    "theme_name": "Kernel-level Anti-cheat Privacy Concerns",
                    "description": (
                        "Players repeatedly question the privacy and trust implications "
                        "of kernel-level anti-cheat access."
                    ),
                    "source_candidate_ids": [items[0]["candidate_id"]],
                    "representative_evidence_ids": ["c7", "c8"],
                    "taxonomy_coverage": "emergent",
                    "coverage_reason": (
                        "The taxonomy covers technical failures and installation problems, "
                        "but it has no category for privacy or trust concerns about "
                        "kernel-level monitoring, which is the core issue in both comments."
                    ),
                }
            ]
        }

    return FakeThemeProvider([candidates], consolidate)


def _inductive_themes() -> list[Any]:
    return discover_inductive_themes(
        _rows(),
        provider=_three_theme_provider(),
        profile="game_publishing",
        min_evidence_count=2,
    )


def _emergent_themes() -> list[Any]:
    return discover_inductive_themes(
        _emergent_rows(),
        provider=_emergent_provider(),
        profile="game_publishing",
        min_evidence_count=2,
    )


def test_taxonomy_based_discovery_still_runs() -> None:
    themes = discover_themes(_rows(), profile="game_publishing")
    assert themes and {theme.discovery_method for theme in themes} == {"taxonomy_based"}
    assert {theme.taxonomy_coverage for theme in themes} == {"covered"}
    assert all("predefined taxonomy" in theme.coverage_reason for theme in themes)


def test_cli_theme_method_taxonomy(tmp_path: Path) -> None:
    output = tmp_path / "taxonomy.json"
    result = CliRunner().invoke(
        app,
        [
            "generate-insights",
            str(_write_csv(tmp_path / "input.csv", _rows())),
            "--output",
            str(output),
            "--profile",
            "game_publishing",
            "--theme-method",
            "taxonomy",
        ],
    )
    assert result.exit_code == 0
    assert json.loads(output.read_text(encoding="utf-8"))["generation_method"] == "classification-semantic-aggregation-v1"


def test_cli_theme_method_llm_uses_provider(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    provider = _three_theme_provider()
    monkeypatch.setattr(cli_module, "CodexFileThemeDiscoveryProvider", lambda _: provider)
    output = tmp_path / "llm.json"
    result = CliRunner().invoke(
        app,
        [
            "generate-insights",
            str(_write_csv(tmp_path / "input.csv", _rows())),
            "--output",
            str(output),
            "--profile",
            "game_publishing",
            "--theme-method",
            "llm",
        ],
    )
    assert result.exit_code == 0
    assert json.loads(output.read_text(encoding="utf-8"))["generation_method"] == "llm-inductive-two-stage-v1"


def test_cli_default_remains_taxonomy(tmp_path: Path) -> None:
    output = tmp_path / "default.json"
    result = CliRunner().invoke(
        app,
        ["generate-insights", str(_write_csv(tmp_path / "input.csv", _rows())), "-o", str(output), "--profile", "game_publishing"],
    )
    assert result.exit_code == 0
    assert json.loads(output.read_text(encoding="utf-8"))["generation_method"] == "classification-semantic-aggregation-v1"


def test_llm_theme_can_be_outside_taxonomy() -> None:
    assert "Kernel-level Anti-cheat Privacy Concerns" in {
        theme.theme_name for theme in _emergent_themes()
    }


def test_emergent_theme_is_legal() -> None:
    theme = _emergent_themes()[0]
    assert theme.taxonomy_coverage == "emergent"
    assert "privacy or trust" in theme.coverage_reason


def test_regional_release_questions_are_covered_by_version_expectation() -> None:
    theme = next(
        item
        for item in _inductive_themes()
        if item.theme_name == "Regional Release Questions"
    )
    assert theme.taxonomy_coverage == "covered"
    assert "Version Expectation" in theme.coverage_reason


def test_one_comment_can_support_multiple_themes() -> None:
    themes = _inductive_themes()
    assert sum("c3" in theme.evidence_comment_ids for theme in themes) == 2


def test_theme_can_map_to_multiple_categories() -> None:
    theme = next(item for item in _inductive_themes() if item.theme_name == "Regional Release Questions")
    assert {"Version Expectation", "Community / Social", "Character & Story"} <= set(theme.related_categories)


def test_unknown_candidate_comment_id_is_rejected() -> None:
    provider = FakeThemeProvider(
        [{"themes": [{"theme_name": "Bad", "description": "Bad ID", "evidence_ids": ["missing"], "representative_evidence_ids": ["missing"]}]}],
        lambda _: {"themes": []},
    )
    with pytest.raises(ClassificationWorkflowError, match="unknown comment_id"):
        discover_inductive_themes(_rows(), provider=provider, profile="game_publishing")


def test_representative_ids_must_belong_to_evidence() -> None:
    provider = FakeThemeProvider(
        [{"themes": [{"theme_name": "Bad", "description": "Bad representative", "evidence_ids": ["c1", "c2"], "representative_evidence_ids": ["c3"]}]}],
        lambda _: {"themes": []},
    )
    with pytest.raises(ClassificationWorkflowError, match="candidate theme validation"):
        discover_inductive_themes(_rows(), provider=provider, profile="game_publishing")


def test_consolidation_preserves_all_candidate_evidence() -> None:
    provider = FakeThemeProvider(
        [
            {"themes": [{"theme_name": "Release A", "description": "Release uncertainty", "evidence_ids": ["c1", "c2"], "representative_evidence_ids": ["c1"]}]},
            {"themes": [{"theme_name": "Release B", "description": "Regional timing", "evidence_ids": ["c3", "c4"], "representative_evidence_ids": ["c3"]}]},
        ],
        lambda items: {"themes": [{"theme_name": "Regional Release Uncertainty", "description": "Players lack release clarity.", "source_candidate_ids": [item["candidate_id"] for item in items], "representative_evidence_ids": ["c1", "c3"], "taxonomy_coverage": "partially_covered", "coverage_reason": "Version Expectation covers release timing, while cross-region communication clarity adds meaning that it does not fully express."}]},
    )
    themes = discover_inductive_themes(_rows()[:4], provider=provider, profile="game_publishing", batch_size=2)
    assert themes[0].evidence_comment_ids == ["c1", "c2", "c3", "c4"]


def test_consolidation_deduplicates_evidence_ids() -> None:
    provider = FakeThemeProvider(
        [{"themes": [{"theme_name": "Release", "description": "Release timing", "evidence_ids": ["c1", "c1", "c2"], "representative_evidence_ids": ["c1"]}]}],
        lambda items: {"themes": [{"theme_name": "Release", "description": "Release timing", "source_candidate_ids": [items[0]["candidate_id"]], "representative_evidence_ids": ["c1"], "taxonomy_coverage": "covered", "coverage_reason": "Version Expectation fully expresses the release-timing question in the supporting comments."}]},
    )
    theme = discover_inductive_themes(_rows(), provider=provider, profile="game_publishing")[0]
    assert theme.evidence_comment_ids == ["c1", "c2"]


def test_evidence_count_matches_unique_ids() -> None:
    assert all(theme.evidence_count == len(theme.evidence_comment_ids) for theme in _inductive_themes())


def test_minimum_evidence_threshold_filters_noise() -> None:
    provider = FakeThemeProvider(
        [{"themes": [{"theme_name": "One-off", "description": "One comment only", "evidence_ids": ["c1"], "representative_evidence_ids": ["c1"]}]}],
        lambda items: {"themes": [{"theme_name": "One-off", "description": "One comment only", "source_candidate_ids": [items[0]["candidate_id"]], "representative_evidence_ids": ["c1"], "taxonomy_coverage": "covered", "coverage_reason": "Version Expectation expresses this release-timing comment, which is filtered only because it lacks recurring evidence."}]},
    )
    assert discover_inductive_themes(_rows(), provider=provider, profile="game_publishing", min_evidence_count=2) == []


def test_related_but_distinct_themes_are_not_forced_to_merge() -> None:
    assert {theme.theme_name for theme in _inductive_themes()} == {
        "Regional Release Questions",
        "Character Visual Appeal",
        "Competitive Fairness Concerns",
    }


def test_discovery_method_is_recorded() -> None:
    assert {theme.discovery_method for theme in _inductive_themes()} == {"llm_inductive"}


def test_taxonomy_coverage_accepts_all_three_values() -> None:
    themes = [*_inductive_themes(), *_emergent_themes()]
    assert {theme.taxonomy_coverage for theme in themes} == {
        "covered",
        "partially_covered",
        "emergent",
    }
    assert all(
        type(theme).model_validate_json(theme.model_dump_json()) == theme
        for theme in themes
    )
    with pytest.raises(Exception):
        ConsolidatedThemeBatch.model_validate({"themes": [{"theme_name": "x", "description": "x", "source_candidate_ids": ["c"], "representative_evidence_ids": ["e"], "taxonomy_coverage": "unknown", "coverage_reason": "The proposed value is outside the supported preliminary coverage states."}]})


def test_llm_theme_requires_nonempty_coverage_reason() -> None:
    proposal = {
        "theme_name": "Release",
        "description": "Release timing",
        "source_candidate_ids": ["candidate-1"],
        "representative_evidence_ids": ["c1"],
        "taxonomy_coverage": "covered",
    }
    for reason in (None, "   "):
        payload = {"themes": [dict(proposal)]}
        if reason is not None:
            payload["themes"][0]["coverage_reason"] = reason
        with pytest.raises(Exception, match="coverage_reason"):
            ConsolidatedThemeBatch.model_validate(payload)


def test_insight_evidence_is_within_theme() -> None:
    themes = _inductive_themes()
    bundle = ThemeInsightBundle(
        profile_id="game_publishing",
        generation_method="llm-inductive-two-stage-v1",
        draft_label="AI generated draft",
        generated_at="2026-08-21T00:00:00Z",
        source_comment_count=len(_rows()),
        themes=themes,
        insight_drafts=draft_insights(themes, profile="game_publishing"),
    )
    validated = validate_theme_insight_bundle(bundle, _rows(), profile="game_publishing")
    themes_by_id = {theme.theme_id: theme for theme in validated.themes}
    assert all(set(item.evidence_ids) <= set(themes_by_id[item.theme_id].evidence_comment_ids) for item in validated.insight_drafts)


def test_json_schemas_parse() -> None:
    root = Path(__file__).parents[1] / "schemas"
    for name in ("theme.schema.json", "insight_draft.schema.json", "theme_insight_bundle.schema.json"):
        assert json.loads((root / name).read_text(encoding="utf-8"))["type"] == "object"
    theme_schema = json.loads((root / "theme.schema.json").read_text(encoding="utf-8"))
    assert theme_schema["properties"]["coverage_reason"]["minLength"] == 1
    assert "coverage_reason" in theme_schema["allOf"][0]["then"]["required"]


def test_llm_output_can_be_reloaded(tmp_path: Path) -> None:
    source = _write_csv(tmp_path / "input.csv", _rows())
    output = tmp_path / "bundle.json"
    generated = generate_theme_insights(source, output, profile_id="game_publishing", theme_method="llm", provider=_three_theme_provider())
    assert load_theme_insight_bundle(output, source_rows=_rows(), profile_id="game_publishing") == generated


def test_workbook_keeps_discovery_columns_after_review_sheet_is_added(tmp_path: Path) -> None:
    source = _write_csv(tmp_path / "input.csv", _rows())
    bundle = generate_theme_insights(source, tmp_path / "bundle.json", profile_id="game_publishing", theme_method="llm", provider=_three_theme_provider())
    path = create_readable_workbook(_rows(), tmp_path / "report.xlsx", profile_id="game_publishing", theme_bundle=bundle)
    workbook = load_workbook(path, data_only=False)
    assert workbook.sheetnames == SHEET_NAMES
    headers = [cell.value for cell in workbook["08_主题发现"][1]]
    assert {"Discovery Method", "Taxonomy Coverage", "Coverage Reason", "Representative Evidence IDs"} <= set(headers)
    reason_column = headers.index("Coverage Reason") + 1
    assert all(
        workbook["08_主题发现"].cell(row, reason_column).value
        for row in range(2, workbook["08_主题发现"].max_row + 1)
    )


def test_workbook_has_no_formula_error_literals(tmp_path: Path) -> None:
    source = _write_csv(tmp_path / "input.csv", _rows())
    bundle = generate_theme_insights(source, tmp_path / "bundle.json", profile_id="game_publishing", theme_method="llm", provider=_three_theme_provider())
    path = create_readable_workbook(_rows(), tmp_path / "report.xlsx", profile_id="game_publishing", theme_bundle=bundle)
    workbook = load_workbook(path, data_only=False)
    errors = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"}
    assert not [cell.coordinate for sheet in workbook for row in sheet.iter_rows() for cell in row if isinstance(cell.value, str) and cell.value in errors]


def test_generate_insights_help_lists_new_options() -> None:
    result = CliRunner().invoke(app, ["generate-insights", "--help"])
    assert result.exit_code == 0
    help_text = " ".join(strip_ansi(result.output).split())
    assert "--theme-method" in help_text
    assert "--theme-batch-size" in help_text
    assert "--min-evidence-count" in help_text
    assert "--theme-provider" in help_text
    assert "taxonomy-based" in help_text
    assert "baseline" in help_text
    assert "file-based" in help_text
    assert "LLM-assisted" in help_text


def test_codex_file_provider_runs_candidate_then_consolidation_stages(
    tmp_path: Path,
) -> None:
    work_dir = tmp_path / "llm-work"
    with pytest.raises(ClassificationWorkflowError, match="responses are pending"):
        discover_inductive_themes(
            _rows(),
            provider=CodexFileThemeDiscoveryProvider(work_dir),
            profile="game_publishing",
            batch_size=2,
        )
    request_paths = sorted(work_dir.glob("candidate_request_*.json"))
    assert len(request_paths) == 3
    for request_path in request_paths:
        request = json.loads(request_path.read_text(encoding="utf-8"))
        evidence_ids = [item["comment_id"] for item in request["comments"]]
        response_path = work_dir / request["response_file"]
        response_path.write_text(
            json.dumps(
                {
                    "themes": [
                        {
                            "theme_name": f"Batch theme {request_path.stem[-3:]}",
                            "description": "A recurring batch-local player pattern.",
                            "evidence_ids": evidence_ids,
                            "representative_evidence_ids": [evidence_ids[0]],
                        }
                    ]
                }
            ),
            encoding="utf-8",
        )
    with pytest.raises(ClassificationWorkflowError, match="consolidation response"):
        discover_inductive_themes(
            _rows(),
            provider=CodexFileThemeDiscoveryProvider(work_dir),
            profile="game_publishing",
            batch_size=2,
        )
    consolidation_request = json.loads(
        (work_dir / "consolidation_request.json").read_text(encoding="utf-8")
    )
    (work_dir / "consolidation_response.json").write_text(
        json.dumps(
            {
                "themes": [
                    {
                        "theme_name": "Cross-batch player pattern",
                        "description": "The same research pattern recurs across batches.",
                        "source_candidate_ids": [
                            item["candidate_id"]
                            for item in consolidation_request["candidates"]
                        ],
                        "representative_evidence_ids": ["c1"],
                        "taxonomy_coverage": "partially_covered",
                        "coverage_reason": (
                            "The taxonomy captures part of the recurring player pattern "
                            "but does not fully express its cross-batch meaning."
                        ),
                    }
                ]
            }
        ),
        encoding="utf-8",
    )
    themes = discover_inductive_themes(
        _rows(),
        provider=CodexFileThemeDiscoveryProvider(work_dir),
        profile="game_publishing",
        batch_size=2,
    )
    assert themes[0].evidence_count == 6
    assert themes[0].coverage_reason.startswith("The taxonomy captures")
