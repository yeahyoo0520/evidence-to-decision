from __future__ import annotations

import csv
import json
from datetime import datetime, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook

from youtube_comment_research.errors import ClassificationWorkflowError
from youtube_comment_research.insight_synthesis import (
    InsightRecord,
    insight_section,
    mark_insight_evidence,
    validate_insight_payload,
)
from youtube_comment_research.profiles import (
    DEFAULT_PROFILE_ID,
    available_profiles,
    load_profile,
    profile_sha256,
)
from youtube_comment_research.readable_workbook import (
    MAIN_COMMENT_FILL,
    REPLY_FILL,
    SHEET_NAMES,
    create_readable_workbook,
)
from youtube_comment_research.review_sampling import (
    apply_review_sampling,
    mandatory_review,
    review_queue,
    summarize_review_workload,
)
from youtube_comment_research.semantic_workflow import (
    merge_classifications,
    prepare_classification_batches,
    validate_classification_payload,
)

PROFILE_IDS = (
    "consumer_product",
    "game_publishing",
    "gtm_research",
    "overseas_content",
)


def _source_rows(count: int = 4) -> list[dict[str, object]]:
    rows = []
    for index in range(count):
        rows.append(
            {
                "project_id": "phase2b-fixture",
                "comment_id": f"c-{index:03d}",
                "comment_text": f"Fixture comment {index}",
                "video_id": f"video-{index % 2}",
                "video_title": f"Video {index % 2}",
                "channel_title": "Fixture Channel",
                "is_reply": False,
                "parent_comment_id": "",
                "comment_like_count": index,
                "reply_count": 0,
                "collection_order": index + 1,
                "collected_at": "2026-08-16T00:00:00Z",
            }
        )
    return rows


def _write_csv(path: Path, rows: list[dict[str, object]]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = list(rows[0])
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)
    return path


def _classification_payload(
    comment_id: str,
    *,
    profile_id: str = DEFAULT_PROFILE_ID,
) -> dict[str, object]:
    profile = load_profile(profile_id)
    category = profile.primary_categories[0]
    subject = profile.subject_options[0]
    return {
        "comment_id": comment_id,
        "subject": subject,
        "primary_category": category,
        "secondary_categories": [],
        "subtopics": [profile.subtopic_guidance[0]],
        "lifecycle_stage": profile.stage_options[0],
        "mentioned_brands": [],
        "mentioned_products": [],
        "confidence": 0.9,
        "evidence": ["Fixture comment"],
        "review_risk": "low",
        "review_triggers": [],
        "manual_review": False,
        "review_reason": "",
        "classification_source": "codex-assisted",
        "classifier_version": "codex-semantic-v2",
        "rubric_version": profile.rubric_version,
    }


def _insight_payload(**overrides: object) -> dict[str, object]:
    payload: dict[str, object] = {
        "insight_id": "insight-1",
        "profile_id": "game_publishing",
        "theme": "Localization wording",
        "theme_description": "Players discuss translation wording.",
        "evidence_count": 2,
        "evidence_comment_ids": ["c-000", "c-001"],
        "representative_comment_ids": ["c-000"],
        "source_video_ids": ["video-0", "video-1"],
        "source_count": 2,
        "insight": "Repeated wording questions indicate an information gap in the localized release notes.",
        "action_opportunity": "Clarify localized terminology in release notes and pinned comments.",
        "priority": "P1",
        "priority_reason": "The theme recurs across two videos, is relevant to the research goal, and has direct evidence.",
        "confidence_level": "medium",
        "status": "preliminary",
        "evidence_review_status": "pending",
        "needs_human_confirmation": True,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "rubric_version": load_profile("game_publishing").rubric_version,
    }
    payload.update(overrides)
    return payload


# Profile coverage (1-9)
@pytest.mark.parametrize("profile_id", PROFILE_IDS)
def test_all_four_profiles_load(profile_id: str) -> None:
    assert load_profile(profile_id).profile_id == profile_id


@pytest.mark.parametrize("profile_id", PROFILE_IDS)
def test_all_four_profiles_validate_complete_taxonomy(profile_id: str) -> None:
    profile = load_profile(profile_id)
    assert set(profile.primary_categories) == set(profile.category_definitions)


def test_consumer_profile_keeps_phase2a3_taxonomy() -> None:
    profile = load_profile("consumer_product")
    assert "Product Praise" in profile.primary_categories
    assert "Purchase Intent" in profile.primary_categories
    assert "Using" in profile.stage_options


def test_game_profile_rejects_consumer_only_category() -> None:
    payload = _classification_payload("c-000", profile_id="game_publishing")
    payload["primary_category"] = "Product Praise"
    with pytest.raises(ClassificationWorkflowError, match="Schema validation failed"):
        validate_classification_payload(
            payload, input_text="Fixture comment 0", profile="game_publishing"
        )


def test_gtm_profile_uses_own_taxonomy() -> None:
    assert "Adoption Barrier" in load_profile("gtm_research").primary_categories
    assert "Product Praise" not in load_profile("gtm_research").primary_categories


def test_overseas_content_profile_uses_own_taxonomy() -> None:
    categories = load_profile("overseas_content").primary_categories
    assert "Positive Content Reaction" in categories
    assert "Localization / Cultural Feedback" in categories


def test_missing_profile_has_clear_error() -> None:
    with pytest.raises(ClassificationWorkflowError, match="Research profile not found"):
        load_profile("missing_profile")


def test_different_profiles_cannot_merge(tmp_path: Path) -> None:
    source = _write_csv(tmp_path / "source.csv", _source_rows(1))
    batch_dir = tmp_path / "batches"
    summary = prepare_classification_batches(
        source, batch_dir, profile_id="game_publishing"
    )
    item = json.loads(summary.batch_files[0].read_text(encoding="utf-8"))
    payload = _classification_payload(item["comment_id"], profile_id="game_publishing")
    (batch_dir / "classified_batch_001.jsonl").write_text(
        json.dumps(payload) + "\n", encoding="utf-8"
    )
    with pytest.raises(ClassificationWorkflowError, match="profile mismatch"):
        merge_classifications(
            source,
            batch_dir,
            tmp_path / "merged.csv",
            profile_id="consumer_product",
        )


def test_manifest_records_profile_and_hash(tmp_path: Path) -> None:
    source = _write_csv(tmp_path / "source.csv", _source_rows(1))
    summary = prepare_classification_batches(
        source, tmp_path / "batches", profile_id="game_publishing"
    )
    manifest = json.loads(summary.manifest_path.read_text(encoding="utf-8"))
    assert manifest["profile_id"] == "game_publishing"
    assert manifest["profile_sha256"] == profile_sha256("game_publishing")
    assert manifest["classifier_version"] == "codex-semantic-v2"
    assert manifest["rubric_sha256"]


# Risk review coverage (10-20)
def test_low_risk_does_not_force_manual_review() -> None:
    assert not mandatory_review({"comment_id": "a", "codex_review_risk": "low"})


def test_medium_risk_does_not_force_manual_review() -> None:
    assert not mandatory_review({"comment_id": "a", "codex_review_risk": "medium"})


def test_high_risk_forces_manual_review() -> None:
    rows = apply_review_sampling(
        [{"comment_id": "a", "codex_review_risk": "high"}],
        load_profile("consumer_product"),
        seed=7,
    )
    assert rows[0]["codex_manual_review"] is True


def test_low_risk_qa_is_about_ten_percent() -> None:
    rows = [{"comment_id": f"l-{i}", "codex_review_risk": "low"} for i in range(100)]
    sampled = apply_review_sampling(rows, load_profile("consumer_product"), seed=5)
    assert sum(row["qa_sample"] for row in sampled) == 10


def test_medium_risk_qa_is_about_twenty_five_percent() -> None:
    rows = [{"comment_id": f"m-{i}", "codex_review_risk": "medium"} for i in range(100)]
    sampled = apply_review_sampling(rows, load_profile("consumer_product"), seed=5)
    assert sum(row["qa_sample"] for row in sampled) == 25


def test_fixed_seed_repeats_same_sample() -> None:
    rows = [{"comment_id": f"c-{i}", "codex_review_risk": "low"} for i in range(60)]
    first = apply_review_sampling(rows, load_profile("consumer_product"), seed=99)
    second = apply_review_sampling(rows, load_profile("consumer_product"), seed=99)
    assert [row["comment_id"] for row in first if row["qa_sample"]] == [
        row["comment_id"] for row in second if row["qa_sample"]
    ]


def test_high_risk_all_enters_review_queue() -> None:
    rows = apply_review_sampling(
        [{"comment_id": f"h-{i}", "codex_review_risk": "high"} for i in range(5)],
        load_profile("consumer_product"),
        seed=1,
    )
    assert len(review_queue(rows)) == 5


def test_evidence_can_require_review_at_any_risk() -> None:
    rows = mark_insight_evidence(
        [{"comment_id": "c-000", "codex_review_risk": "low"}],
        [_insight_payload(evidence_count=1, evidence_comment_ids=["c-000"], representative_comment_ids=["c-000"], source_video_ids=["video-0"], source_count=1)],
    )
    assert rows[0]["evidence_review_required"] is True


def test_unconfirmed_core_evidence_not_confirmed_finding() -> None:
    assert insight_section(_insight_payload()) == "Insight Review Queue"


def test_confirmed_evidence_can_enter_confirmed_findings() -> None:
    payload = _insight_payload(
        status="confirmed",
        evidence_review_status="confirmed",
        needs_human_confirmation=False,
    )
    assert insight_section(payload) == "Confirmed Findings"


def test_review_workload_deduplicates_overlapping_reasons() -> None:
    rows = [
        {
            "comment_id": "same",
            "codex_review_risk": "high",
            "qa_sample": True,
            "evidence_review_required": True,
        }
    ]
    assert summarize_review_workload(rows).unique_human_review_workload == 1


# Insight coverage (21-30)
def test_insight_requires_evidence_ids() -> None:
    with pytest.raises(ClassificationWorkflowError):
        validate_insight_payload(
            _insight_payload(evidence_count=0, evidence_comment_ids=[]),
            _source_rows(),
            profile="game_publishing",
        )


def test_unknown_evidence_comment_id_is_rejected() -> None:
    with pytest.raises(ClassificationWorkflowError, match="unknown comment_id"):
        validate_insight_payload(
            _insight_payload(evidence_count=1, evidence_comment_ids=["missing"], representative_comment_ids=[] , source_video_ids=["video-0"], source_count=1),
            _source_rows(),
            profile="game_publishing",
        )


def test_evidence_count_must_match_ids() -> None:
    with pytest.raises(ClassificationWorkflowError, match="evidence_count"):
        validate_insight_payload(
            _insight_payload(evidence_count=3),
            _source_rows(),
            profile="game_publishing",
        )


def test_preliminary_insight_not_confirmed_section() -> None:
    assert insight_section(_insight_payload(needs_human_confirmation=False)) == "Preliminary Insights"


def test_rejected_insight_not_confirmed_section() -> None:
    assert insight_section(_insight_payload(status="rejected")) != "Confirmed Findings"


def test_priority_only_allows_p0_p1_p2() -> None:
    with pytest.raises(ClassificationWorkflowError):
        validate_insight_payload(
            _insight_payload(priority="P3"), _source_rows(), profile="game_publishing"
        )


def test_priority_requires_reason() -> None:
    with pytest.raises(ClassificationWorkflowError):
        validate_insight_payload(
            _insight_payload(priority_reason=""), _source_rows(), profile="game_publishing"
        )


def test_action_opportunity_requires_insight() -> None:
    with pytest.raises(ClassificationWorkflowError):
        validate_insight_payload(
            _insight_payload(insight=""), _source_rows(), profile="game_publishing"
        )


def test_insight_requires_theme() -> None:
    with pytest.raises(ClassificationWorkflowError):
        validate_insight_payload(
            _insight_payload(theme=""), _source_rows(), profile="game_publishing"
        )


def test_insight_source_videos_must_match_current_project() -> None:
    with pytest.raises(ClassificationWorkflowError, match="source_video_ids"):
        validate_insight_payload(
            _insight_payload(source_video_ids=["other-video"], source_count=1),
            _source_rows(),
            profile="game_publishing",
        )


# Workbook coverage (31-38)
def _workbook_rows() -> list[dict[str, object]]:
    rows = _source_rows(3)
    rows[0].update(
        {
            "codex_subject": "game",
            "codex_primary_category": "Game / Content Praise",
            "codex_lifecycle_stage": "Playing",
            "codex_confidence": 0.9,
            "codex_review_risk": "low",
            "codex_manual_review": False,
            "review_status": "confirmed",
            "human_confirmed_primary_category": "Game / Content Praise",
            "qa_status": "confirmed",
            "evidence_review_status": "confirmed",
            "reviewer_note": "Keep this human note.",
        }
    )
    rows[1].update(
        {
            "is_reply": True,
            "video_id": "video-0",
            "parent_comment_id": "c-000",
            "codex_subject": "game",
            "codex_primary_category": "Localization Feedback",
            "codex_lifecycle_stage": "Playing",
            "codex_confidence": 0.78,
            "codex_review_risk": "medium",
        }
    )
    rows[2].update(
        {
            "codex_subject": "game",
            "codex_primary_category": "Technical Issue",
            "codex_lifecycle_stage": "Playing",
            "codex_confidence": 0.7,
            "codex_review_risk": "high",
            "codex_manual_review": True,
        }
    )
    return rows


def _game_workbook(path: Path) -> Path:
    rows = _workbook_rows()
    insight = InsightRecord.model_validate(
        _insight_payload(
            status="confirmed",
            evidence_review_status="confirmed",
            needs_human_confirmation=False,
        )
    )
    return create_readable_workbook(
        rows,
        path,
        profile_id="game_publishing",
        insights=[insight],
    )


def test_workbook_has_seven_sheets(tmp_path: Path) -> None:
    assert len(load_workbook(_game_workbook(tmp_path / "game.xlsx")).sheetnames) == len(
        SHEET_NAMES
    )


def test_workbook_sheet_order_is_correct(tmp_path: Path) -> None:
    assert load_workbook(_game_workbook(tmp_path / "game.xlsx")).sheetnames == SHEET_NAMES


def test_conversation_keeps_yellow_and_blue_rows(tmp_path: Path) -> None:
    workbook = load_workbook(_game_workbook(tmp_path / "game.xlsx"))
    sheet = workbook[SHEET_NAMES[1]]
    assert sheet.cell(2, 1).fill.fgColor.rgb == MAIN_COMMENT_FILL.fgColor.rgb
    assert sheet.cell(3, 1).fill.fgColor.rgb == REPLY_FILL.fgColor.rgb


def test_semantic_sheet_shows_risk_qa_and_evidence_review(tmp_path: Path) -> None:
    workbook = load_workbook(_game_workbook(tmp_path / "game.xlsx"))
    headers = [cell.value for cell in workbook[SHEET_NAMES[2]][1]]
    assert "Review Risk" in headers
    assert "QA Sample" in headers
    assert "Evidence Review Required" in headers


def test_insight_sheet_separates_preliminary_and_confirmed(tmp_path: Path) -> None:
    workbook = load_workbook(_game_workbook(tmp_path / "game.xlsx"))
    sections = {cell.value for cell in workbook[SHEET_NAMES[3]]["A"][1:]}
    assert {"Confirmed Findings", "Preliminary Insights", "Insight Review Queue"}.issubset(sections)


def test_raw_data_row_count_is_not_reduced(tmp_path: Path) -> None:
    workbook = load_workbook(_game_workbook(tmp_path / "game.xlsx"))
    assert workbook[SHEET_NAMES[5]].max_row == len(_workbook_rows()) + 1


def test_human_confirmation_is_not_overwritten(tmp_path: Path) -> None:
    workbook = load_workbook(_game_workbook(tmp_path / "game.xlsx"))
    sheet = workbook[SHEET_NAMES[2]]
    headers = [cell.value for cell in sheet[1]]
    note_column = headers.index("Reviewer Note") + 1
    assert sheet.cell(2, note_column).value == "Keep this human note."


def test_generated_xlsx_reopens_with_openpyxl(tmp_path: Path) -> None:
    path = _game_workbook(tmp_path / "game.xlsx")
    assert load_workbook(path).sheetnames == SHEET_NAMES


def test_profile_demo_fixture_covers_all_available_taxonomies() -> None:
    path = Path(__file__).parent / "fixtures" / "phase2b_profile_demo.json"
    payload = json.loads(path.read_text(encoding="utf-8"))
    assert set(payload["demonstrations"]) == set(available_profiles())
    assert len({item["primary_category"] for item in payload["demonstrations"].values()}) == len(available_profiles())
